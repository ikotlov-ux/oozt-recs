# -*- coding: utf-8 -*-
"""
oozt_recs.py — подготовка природоохранных рекомендаций (мероприятий) для ООЗТ Москвы
на основе ГИС-слоёв и таблицы «Мероприятия по Приказу».

Единицы назначения мероприятий (режим выбирается в окне или ключом unit_mode в конфигурации)
--------------------------------------------------------------------------------------------
    kv_vyd — кварталы и выделы: поля ООЗТ, квартала, выдела, типа; отчёт и схемы — по кварталам;
    eco    — слой экосистем с уникальным ID, без кварталов: поля ID и типа (ООЗТ и название — по желанию);
             отчёт и схемы — по листам регулярной сетки (sheet_size_m, по умолчанию 1000 м)
             или по значениям поля группировки (report_group_field).

Запуск
------
    python oozt_recs.py                   # окно выбора данных: режим, слои, живая сводка мероприятий
    python oozt_recs.py --config cfg.json # повторный запуск с сохранённой конфигурацией, без окна
    python oozt_recs.py --config cfg.json --gui        # окно, предзаполненное из конфигурации
    python oozt_recs.py --config cfg.json --kv 1,2,5   # отчёт только по выбранным кварталам / листам
    python oozt_recs.py --config cfg.json --report-only --formats docx,pdf,doc   # пересобрать отчёт

Выходные данные (папка out_dir)
-------------------------------
    measures_long.xlsx / .csv   — длинная таблица ООЗТ–Квартал(Лист)–Выдел(ID)–Мероприятие (+ основание)
    evidence.gpkg               — полигоны с мероприятиями (vydel_measures / eco_measures), measures (полигон × мероприятие),
                                  слои-основания evidence_*
    oozt_measures.qgs           — проект QGIS с этими слоями и стилями (относительные пути)
    report/report.md            — отчёт по кварталам или листам (схема + таблицы), рисунки report/fig_*.png
    report/report.docx, .pdf    — тот же отчёт в Word и PDF (report_formats / --formats; .doc — по запросу)
    config_used.json            — конфигурация запуска (можно передать в --config)
    run.log                     — журнал, в начале — статус каждого мероприятия при выбранных данных

Правила (ключ Метод выявления в Excel → расчёт) — classify_measure() и раздел ПРАВИЛА;
связь правил со слоями — RULES (data / opt). Обязательны только таблица мероприятий и слой территории.
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import tempfile
import warnings
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Iterable

# ---- локальная папка pylib рядом со скриптом (чистые Python-пакеты: openpyxl, contextily, python-docx) ----
_pylib = Path(__file__).resolve().parent / "pylib"
if _pylib.is_dir() and str(_pylib) not in sys.path:
    sys.path.append(str(_pylib))          # append: пакеты окружения имеют приоритет

import faulthandler
faulthandler.enable()                     # аварийный стек в stderr при падении C-расширений

# ---- окружение GDAL/matplotlib (conda на Windows часто не выставляет GDAL_DATA) ----
_env_root = Path(sys.prefix)
# Windows: при запуске python.exe без activate папки с DLL окружения не попадают в PATH,
# и C-расширения matplotlib падают с 0xc06d007f (delay-load DLL). Добавляем их сами.
def enable_win_dll_path():
    """Windows: python.exe без `conda activate` не видит DLL окружения, и C-расширения matplotlib
    падают с 0xc06d007f (delay-load DLL). Включаем PATH окружения только перед отрисовкой:
    делать это на старте нельзя — меняется порядок загрузки DLL для GDAL."""
    if os.name != "nt" or os.environ.get("_OOZT_DLLPATH"):
        return
    dirs = [str(p) for p in (_env_root, _env_root / "Library" / "bin",
                             _env_root / "Library" / "mingw-w64" / "bin",
                             _env_root / "Library" / "usr" / "bin",
                             _env_root / "DLLs", _env_root / "Scripts") if p.is_dir()]
    for d in dirs:
        try:
            os.add_dll_directory(d)
        except OSError:
            pass
    os.environ["PATH"] = os.pathsep.join(dirs + [os.environ.get("PATH", "")])
    os.environ["_OOZT_DLLPATH"] = "1"
for _cand in (_env_root / "Library" / "share" / "gdal", _env_root / "share" / "gdal"):
    if _cand.exists() and not os.environ.get("GDAL_DATA"):
        os.environ["GDAL_DATA"] = str(_cand)
for _cand in (_env_root / "Library" / "share" / "proj", _env_root / "share" / "proj"):
    if _cand.exists() and not os.environ.get("PROJ_LIB"):
        os.environ["PROJ_LIB"] = str(_cand)
# Windows/Python ≥ 3.12.4: tempfile.mkdtemp создаёт папку с ограниченным ACL; под сервисными учётными записями
# это ломает matplotlib/joblib/contextily. Заменяем на обычный os.makedirs.
if os.name == "nt":
    import tempfile as _tf
    import uuid as _uuid

    def _mkdtemp(suffix=None, prefix=None, dir=None):
        base = dir or _tf.gettempdir()
        for _ in range(20):
            p = os.path.join(base, (prefix or "tmp") + _uuid.uuid4().hex[:10] + (suffix or ""))
            if not os.path.exists(p):
                os.makedirs(p)
                return p
        raise FileExistsError("не удалось создать временную папку")
    _tf.mkdtemp = _mkdtemp

os.environ.setdefault("MPLCONFIGDIR", os.path.join(os.environ.get("TEMP", "."), "mpl_cache"))
os.environ.setdefault("PYTHONIOENCODING", "utf-8")
os.environ.setdefault("SHAPE_RESTORE_SHX", "YES")  # восстанавливать .shx, если файл отсутствует/повреждён
try:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
except Exception:
    pass

warnings.filterwarnings("ignore")

import numpy as np
import pandas as pd
import geopandas as gpd
import pyogrio
import rasterio
from rasterio import features as rfeatures
from rasterio.warp import reproject, Resampling, calculate_default_transform
from rasterio.transform import from_origin
from shapely.geometry import shape, box, Point
from shapely.ops import unary_union
from scipy import ndimage

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.lines import Line2D

log = logging.getLogger("oozt")
logging.getLogger("pyogrio").setLevel(logging.WARNING)
logging.getLogger("rasterio").setLevel(logging.WARNING)

# =====================================================================================
#                                    ПАРАМЕТРЫ
# =====================================================================================
WORK_EPSG = 32637          # рабочая проекция (WGS84 / UTM 37N, метры)
CELL = 10.0                # шаг сетки для растров плотности/уклона, м
TRAIL_WINDOW_M = 50.0      # радиус скользящего окна плотности троп, м
CORVUS_WINDOW_M = 500.0    # радиус окна плотности врановых, м
RARE_BUFFER_M = 50.0       # буфер от точек редких видов, м
SHORE_BUFFER_M = 50.0      # буфер вдоль водных объектов для «укрепления берегов», м
SDM_THRESHOLD = 0.75       # порог вероятности SDM (по умолчанию, редкие виды)
SDM_THRESHOLD_INVASIVE = 0.95   # порог вероятности SDM для инвазионных видов
SDM_THRESHOLD_RARE = SDM_THRESHOLD  # порог вероятности SDM для редких видов
MIN_PATCH_HA = 0.05        # минимальная площадь пятна (уклон/плотность/SDM) в выделе для назначения, га
QUANTILE = 0.75            # «верхний квартиль»

# Типы дорог OSM, считающиеся дорожно-тропиночной сетью (если в слое есть поле HIGHWAY/fclass)
TRAIL_HIGHWAY_TYPES = {"path", "footway", "track", "cycleway", "steps", "bridleway",
                       "pedestrian", "living_street", "service", "unclassified"}

TYPE_FOREST = ("лес",)
TYPE_MEADOW = ("луг",)
TYPE_WATER = ("водн",)          # «Водные»
TYPE_NEARWATER = ("околовод", "болот")

# =====================================================================================
#                                   КОНФИГУРАЦИЯ
# =====================================================================================
UNIT_MODES = {
    "kv_vyd": "Кварталы и выделы",
    "eco": "Слой экосистем с уникальным ID (без кварталов)",
}


@dataclass
class Config:
    measures_xlsx: str = ""
    unit_mode: str = "kv_vyd"            # kv_vyd — кварталы/выделы; eco — полигоны экосистем с уникальным ID
    oozt_path: str = ""                  # слой единиц территории (выделы или экосистемы)
    oozt_layer: str | None = None
    oozt_name: str = ""                  # eco: название ООЗТ, если в слое нет поля ООЗТ
    report_group_field: str | None = None  # eco: поле группировки отчёта; пусто — листы регулярной сетки
    sheet_size_m: float = 1000.0         # eco: размер листа сетки для схем, м
    dem_path: str | None = None
    trails: list[str] = field(default_factory=list)
    hydro_lines: list[str] = field(default_factory=list)
    hydro_polys: list[str] = field(default_factory=list)
    # слои видов и редких экосистем: [{path, category, group, kind, field}], см. SPECIES_*
    species_layers: list = field(default_factory=list)
    out_dir: str = ""
    report_kv: list = field(default_factory=list)   # кварталы / листы / группы для отчёта; пусто = все
    basemap: bool = True
    min_share_pct: float = 25.0          # мин. доля полигона под участком-основанием (тропы, уклон, вороны, SDM), %
    split_multipart: bool = True         # многочастный объект → отдельные контуры (ID-1, ID-2, …)
    report_formats: list = field(default_factory=lambda: ["docx", "pdf"])  # docx, pdf, doc; report.md — всегда
    fields: dict = field(default_factory=dict)      # kv_vyd: OOZT, KV, VYD, TYPE, NAME; eco: ID, TYPE, NAME, OOZT

    def save(self, path: str | Path):
        Path(path).write_text(json.dumps(asdict(self), ensure_ascii=False, indent=2), encoding="utf-8")

    @staticmethod
    def load(path: str | Path) -> "Config":
        d = json.loads(Path(path).read_text(encoding="utf-8"))
        known = set(Config.__dataclass_fields__)
        if "shore_min_share_pct" in d and "min_share_pct" not in d:   # прежнее имя параметра
            d["min_share_pct"] = d["shore_min_share_pct"]
        layers = list(d.get("species_layers") or [])
        for key, (cat, grp, kind) in LEGACY_SPECIES_KEYS.items():  # прежние поля конфигурации
            for pth in d.get(key) or []:
                if not any(l.get("path") == pth for l in layers):
                    layers.append(dict(path=pth, category=cat, group=grp, kind=kind, field=""))
        d["species_layers"] = layers
        return Config(**{k: v for k, v in d.items() if k in known})


# прежние поля конфигурации (до версии со списком слоёв видов) → категория, группа, тип данных
LEGACY_SPECIES_KEYS = {
    "invasive_points": ("инвазивные", "не указана", "points"),
    "invasive_polys": ("инвазивные", "не указана", "polygons"),
    "invasive_rasters": ("инвазивные", "не указана", "sdm"),
    "rare_points": ("охраняемые", "не указана", "points"),
    "rare_polys": ("охраняемые", "не указана", "polygons"),
    "rare_rasters": ("охраняемые", "не указана", "sdm"),
    "fauna_points": ("прочие", "птицы", "points"),
}


def unit_labels(cfg: Config) -> dict:
    """Подписи территориальных единиц в таблицах, отчёте и на схемах (зависят от режима)."""
    if cfg.unit_mode == "eco":
        grp = cfg.report_group_field
        return {
            "group": grp if grp else "Лист", "group_pl": f"группам «{grp}»" if grp else "листам",
            "groups_n": "Групп" if grp else "Листов",
            "unit": "ID экосистемы", "units_n": "Экосистем", "units_gen": "экосистем",
            "unit_type": "Тип полигона", "unit_name": "Экосистема",
        }
    return {"group": "Квартал", "group_pl": "кварталам", "groups_n": "Кварталов",
            "unit": "Выдел", "units_n": "Выделов", "units_gen": "выделов",
            "unit_type": "Тип выдела", "unit_name": "Экосистема выдела"}


def display_columns(cfg: Config) -> dict:
    """Внутренние имена колонок (Квартал/Выдел/Тип выдела) → подписи текущего режима."""
    L = unit_labels(cfg)
    return {"Квартал": L["group"], "Выдел": L["unit"], "Тип выдела": L["unit_type"]}


# =====================================================================================
#                 ВХОДНЫЕ СЛОИ И ИХ СВЯЗЬ С МЕРОПРИЯТИЯМИ (для интерфейса и журнала)
# =====================================================================================
# key — поле Config; multi — можно выбрать несколько файлов
LAYER_SPECS = [
    dict(key="dem_path", title="Цифровая модель рельефа", multi=False, kind="ras",
         what="растр высот, 1 канал (tif, asc); уклон считается на сетке 10 м"),
    dict(key="trails", title="Дорожно-тропиночная сеть", multi=True, kind="vec",
         what="линии; OSM highway (path, footway, track…) или полевая съёмка"),
    dict(key="hydro_lines", title="Гидрография — линии", multi=True, kind="vec",
         what="реки, ручьи, каналы"),
    dict(key="hydro_polys", title="Гидрография — полигоны", multi=True, kind="vec",
         what="пруды, озёра, водохранилища"),
]
LAYER_TITLE = {s["key"]: s["title"] for s in LAYER_SPECS}
LAYER_TITLE["species"] = "Слои видов и редких экосистем"

# ---- слои видов и редких экосистем ---------------------------------------------------
# Каждый слой описывается категорией, группой организмов и типом данных. Мероприятия
# выбирают слои по этим признакам (столбцы «Виды: …» таблицы мероприятий).
SPECIES_CATEGORIES = ["инвазивные", "охраняемые", "прочие"]
SPECIES_CATEGORY_HELP = {
    "инвазивные": "чужеродные инвазивные виды",
    "охраняемые": "виды Красной книги и другие охраняемые или редкие виды, редкие экосистемы",
    "прочие": "прочие наблюдения (например, все птицы — для плотности врановых)",
}
SPECIES_GROUPS = ["сосудистые растения", "грибы", "лишайники", "млекопитающие", "птицы",
                  "низшие позвоночные", "насекомые", "редкие экосистемы", "не указана"]
SPECIES_KINDS = {"points": "точки", "polygons": "полигоны", "sdm": "SDM (растр 0–1)"}
RASTER_EXT = (".tif", ".tiff", ".asc", ".img", ".vrt")


def norm_category(v) -> str:
    t = str(v or "").strip().lower().replace("ё", "е")
    if not t:
        return ""
    if t.startswith(("инваз", "чужерод")):
        return "инвазивные"
    if t.startswith(("охран", "редк", "краснокниж")):
        return "охраняемые"
    if t.startswith(("проч", "друг", "фаун", "наблюд")):
        return "прочие"
    return t


def norm_group(v) -> str:
    t = re.sub(r"\s+", " ", str(v or "").strip().lower().replace("ё", "е"))
    return t or "не указана"


def group_matches(layer_group: str, groups) -> bool:
    """Группа слоя входит в список групп мероприятия (сравнение без учёта окончаний)."""
    if not groups:
        return True
    lg = norm_group(layer_group)
    for g in groups:
        g = norm_group(g)
        n = max(4, min(len(g), len(lg)) - 2)
        if lg == g or lg[:n] == g[:n]:
            return True
    return False


def detect_species_kind(path: str) -> str:
    """Тип данных слоя видов: SDM — растр, иначе по геометрии векторного слоя."""
    p = str(path).split("|")[0]
    if p.lower().endswith(RASTER_EXT):
        return "sdm"
    layer = str(path).split("|")[1] if "|" in str(path) else None
    for enc in (None, "latin1"):            # latin1 — если имена полей в другой кодировке
        try:
            info = pyogrio.read_info(p, layer=layer, **({"encoding": enc} if enc else {}))
        except Exception:
            continue
        gt = str(info.get("geometry_type") or "")
        if "Point" in gt:
            return "points"
        if "Polygon" in gt:
            return "polygons"
        return ""
    return ""


def species_layers(cfg) -> list[dict]:
    """Нормализованный список слоев видов из конфигурации."""
    out = []
    for l in getattr(cfg, "species_layers", None) or []:
        if isinstance(l, str):
            l = dict(path=l)
        pth = l.get("path")
        if not pth:
            continue
        out.append(dict(path=pth, category=norm_category(l.get("category")) or "прочие",
                        group=norm_group(l.get("group")),
                        kind=l.get("kind") or detect_species_kind(pth), field=l.get("field") or ""))
    return out


def species_layer_label(l: dict) -> str:
    return (f"{Path(str(l['path']).split('|')[0]).name} ({l['category']}, {l['group']}, "
            f"{SPECIES_KINDS.get(l['kind'], l['kind'] or '?')})")


# ---- правила назначения ---------------------------------------------------------------
REQ, SHARED = "__обязательно__", "__окно__"      # параметр обязателен / берётся из окна скрипта
SPECIES_FILTERS = ("category", "groups", "taxon")
RULES = {
    "type": dict(
        title="по типу полигона", data=[], opt=[],
        params=dict(types=REQ),
        help="Назначается всем полигонам указанных типов. Пространственных расчётов нет."),
    "trail_density": dict(
        title="плотность троп в полигоне", data=["trails"], opt=[],
        params=dict(types=None, quantile=0.75),
        help="Длина троп внутри полигона / площадь полигона (м/га); назначается полигонам с плотностью "
             "не ниже квантиля среди полигонов этих типов."),
    "trail_hotspot": dict(
        title="участки повышенной плотности троп", data=["trails"], opt=[],
        params=dict(types=None, window=50.0, quantile=0.75, min_ha=0.05, min_share=SHARED),
        help="Растр плотности троп в скользящем окне; участки не ниже квантиля; назначается полигонам, "
             "где такие участки занимают заданную площадь и долю."),
    "shore": dict(
        title="тропы у водных объектов", data=["trails", "hydro"], opt=[],
        params=dict(types=None, dist=50.0, window=50.0, quantile=0.75, min_ha=0.01, min_share=SHARED),
        help="Участки повышенной плотности троп в полосе заданной ширины вокруг линий и полигонов "
             "гидрографии."),
    "slope": dict(
        title="уклон", data=["dem"], opt=[],
        params=dict(types=None, slope=REQ, min_ha=0.05, min_share=SHARED),
        help="Участки с уклоном больше порога по ЦМР."),
    "species": dict(
        title="виды: находки", data=["species"], opt=[],
        params=dict(types=None, category=None, groups=None, taxon=None, sources=None, sdm=0.75,
                    min_ha=0.05, min_share=SHARED),
        help="Полигоны, где есть точки или полигоны выбранных видов; по моделям SDM — отдельная строка "
             "«мониторинг»."),
    "species_trails": dict(
        title="виды: тропы у находок", data=["species"], opt=["trails"],
        params=dict(types=None, category=None, groups=None, taxon=None, sources=None, dist=50.0, sdm=0.75,
                    min_ha=0.05, min_share=SHARED),
        help="Тропы в заданном расстоянии от точек и внутри полигонов выбранных видов; по моделям SDM — "
             "отдельная строка «мониторинг»."),
    "species_density": dict(
        title="виды: плотность находок", data=["species"], opt=[],
        params=dict(types=None, category=None, groups=None, taxon=None, window=500.0, quantile=0.75,
                    min_ha=0.05, min_share=SHARED),
        help="Растр плотности точек выбранных видов в скользящем окне; участки не ниже квантиля."),
}
RULE_TITLES = [v["title"] for v in RULES.values()]

# столбцы таблицы мероприятий, задающие правило (ключ, заголовок, начало заголовка для поиска, пояснение)
SPEC_COLUMNS = [
    ("rule", "Правило назначения", "правило", "одно из правил (список); пусто — правило по тексту «Метода выявления»"),
    ("types", "Типы полигонов", "типы полиг",
     "типы полигонов слоя территории через запятую: лесные, луговые, водные, околоводные, болотные…; "
     "«любые» или пусто — все полигоны"),
    ("category", "Виды: категория", "виды: катег", "инвазивные, охраняемые, прочие (через запятую); «любые» или пусто — все"),
    ("groups", "Виды: группы организмов", "виды: групп",
     "сосудистые растения, грибы, лишайники, млекопитающие, птицы, низшие позвоночные, насекомые, "
     "редкие экосистемы…; «любые» или пусто — все"),
    ("taxon", "Виды: таксоны (лат.)", "виды: такс",
     "латинские названия видов или родов через запятую (например, Corvus); пусто — все виды"),
    ("sources", "Виды: данные", "виды: данн", "точки, полигоны, SDM (через запятую); «любые» или пусто — все"),
    ("dist", "Расстояние, м", "расстоян", "ширина полосы у воды или радиус вокруг точек видов"),
    ("window", "Окно плотности, м", "окно", "радиус скользящего окна для растра плотности"),
    ("quantile", "Квантиль", "квантил", "порог плотности; 0,75 — верхний квартиль"),
    ("slope", "Уклон более, %", "уклон", "порог уклона, %"),
    ("sdm", "Порог SDM (0–1)", "порог sdm", "вероятность модели, начиная с которой участок учитывается"),
    ("min_ha", "Мин. площадь участка, га", "мин. площ", "минимальная площадь участка внутри полигона"),
    ("min_share", "Мин. доля полигона, %", "мин. дол",
     "минимальная доля площади полигона под участками; пусто — общий порог из окна скрипта"),
]
SPEC_TITLE = {k: t for k, t, _, _ in SPEC_COLUMNS}
DESC_COLUMN = "Процедура назначения (как работает скрипт)"

TYPE_WORDS = [("околовод", "околовод"), ("болот", "болот"), ("водн", "водн"), ("вод", "водн"),
              ("лес", "лес"), ("луг", "луг")]


def type_matches(type_value, keys) -> bool:
    """Тип полигона относится к одному из ключей. «водн» не совпадает с «околоводные»:
    иначе мероприятия только для водных экосистем (аэрация) попадают и на околоводные."""
    t = str(type_value).lower()
    t_water = t.replace("околовод", "")          # «водн» ищем вне слова «околоводные»
    for k in keys:
        if k in (t_water if k in TYPE_WATER else t):
            return True
    return False


def classify_measure(measure: dict) -> dict:
    """Метод выявления → правило. Возвращает {'rule', 'params', 'reason'} (rule=None — не рассчитывается).
    Правило определяется по тексту метода выявления; название мероприятия используется только там,
    где метод сам по себе неоднозначен (общий текст «точечные / полигональные / растровые данные»,
    поиск троп без указания объекта)."""
    m = re.sub(r"\s+", " ", (measure.get("method") or "").lower())
    nm = (measure.get("name") or "").lower()
    if not m:
        return dict(rule=None, params={}, reason="метод выявления не задан")
    if "нет данных" in m:
        return dict(rule=None, params={}, reason="по таблице: нет данных")
    water_words = ("водн", "водо", "берег", "акватор", "реки", "ручь", "пруд")
    species_data = any(w in m for w in ("точечн", "полигональн", "растров"))
    trails = "троп" in m or "дорог" in m
    if "corvus" in m or "врано" in m or ("птиц" in m and "плотн" in m):
        return dict(rule="corvus", params={}, reason="")
    if "уклон" in m or "цмр" in m:
        mt = re.search(r"(\d+(?:[.,]\d+)?)\s*%", m)
        thr = float(mt.group(1).replace(",", ".")) if mt else 7.0
        keys = TYPE_MEADOW if "лугов" in m else None
        if keys is None and "люб" in str(measure.get("eco", "")).lower():
            thr = float(mt.group(1).replace(",", ".")) if mt else 7.0
        return dict(rule="slope", params=dict(threshold_pct=thr, type_keys=keys), reason="")
    if "точек видов" in m or ("троп" in m and "точ" in m and "50" in m and "вид" in m):
        return dict(rule="rare", params={}, reason="")
    if ("инвази" in m or "чужерод" in m or "чежерод" in m) or (
            species_data and not trails and any(w in nm for w in ("инвази", "чужерод", "чежерод"))):
        return dict(rule="invasive", params={}, reason="")
    if trails and ("периметр" in m or any(w in m for w in water_words) or "берег" in nm):
        return dict(rule="shore", params={}, reason="")
    if "скольз" in m and "троп" in m:
        return dict(rule="trail_hotspot", params={}, reason="")
    if "плотност" in m and trails:
        return dict(rule="trail_density", params={}, reason="")
    if "лугов" in m:
        return dict(rule="type", params=dict(keys=TYPE_MEADOW, label="луговые"), reason="")
    if "околовод" in m or "болот" in m:
        return dict(rule="type", params=dict(keys=TYPE_WATER + TYPE_NEARWATER, label="водные и околоводные"),
                    reason="")
    if any(w in m for w in ("водн", "водо", "акватор")):
        return dict(rule="type", params=dict(keys=TYPE_WATER, label="водные"), reason="")
    if "лесн" in m:
        return dict(rule="type", params=dict(keys=TYPE_FOREST, label="лесные"), reason="")
    if "ворон" in nm and "плотн" in m:
        return dict(rule="corvus", params={}, reason="")
    return dict(rule=None, params={}, reason="метод не распознан")


_MIN_SHARE = [25.0]     # текущий порог доли полигона (обновляется из конфигурации)


def _fmt(x, nd=2) -> str:
    """Число для описаний: 0,75; 50; 0,05."""
    if x is None:
        return "—"
    x = float(x)
    s = f"{x:.{nd}f}".rstrip("0").rstrip(".") if abs(x - round(x)) > 1e-9 else str(int(round(x)))
    return s.replace(".", ",")


def _is_any(t: str) -> bool:
    t = t.strip().lower()
    return t in ("", "-", "—", "любые", "любой", "любая", "все", "всё", "any")


def _split_list(t) -> list[str]:
    return [x.strip() for x in re.split(r"[,;/]|\s+и\s+", str(t)) if x.strip()]


def parse_types(text):
    """«лесные, луговые» → (('лес', 'луг'), 'лесные, луговые'); «любые» → None."""
    t = str(text or "").strip()
    if _is_any(t):
        return None
    keys = []
    for tok in _split_list(t.lower().replace("ё", "е")):
        k = next((key for stem, key in TYPE_WORDS if stem in tok), None)
        if k is None:                      # другой тип: сравнение по началу слова
            k = tok[:max(4, len(tok) - 2)]
        if k not in keys:
            keys.append(k)
    return (tuple(keys), t) if keys else None


def _parse_value(key: str, raw):
    """Значение ячейки → параметр правила. Возвращает (значение, ошибка)."""
    if raw is None or (isinstance(raw, float) and np.isnan(raw)) or str(raw).strip() == "":
        return None, None
    t = str(raw).strip()
    if key == "types":
        return parse_types(t), None
    if key in ("category", "groups", "sources"):
        if _is_any(t):
            return None, None
        items = _split_list(t)
        if key == "category":
            vals = [norm_category(x) for x in items]
            bad = [x for x, v in zip(items, vals) if v not in SPECIES_CATEGORIES]
            return (vals, None) if not bad else (None, f"неизвестная категория: {', '.join(bad)}")
        if key == "groups":
            return [norm_group(x) for x in items], None
        vals = []
        for x in items:
            xl = x.lower()
            v = ("points" if xl.startswith("точ") else "polygons" if xl.startswith("полиг") else
                 "sdm" if xl.startswith(("sdm", "растр", "модел")) else None)
            if v is None:
                return None, f"неизвестный тип данных: {x}"
            vals.append(v)
        return vals, None
    if key == "taxon":
        return (None, None) if _is_any(t) else (_split_list(t), None)
    try:
        v = float(t.replace(",", ".").replace("%", "").strip())
    except ValueError:
        return None, f"не число: «{t}»"
    if key == "quantile" and not 0 < v < 1:
        return None, "квантиль должен быть между 0 и 1"
    if key == "sdm" and not 0 <= v <= 1:
        return None, "порог SDM должен быть от 0 до 1"
    if v < 0:
        return None, "отрицательное значение"
    return v, None


def rule_from_title(t) -> str | None:
    s = re.sub(r"\s+", " ", str(t or "").strip().lower().replace("ё", "е"))
    if not s:
        return None
    for k, v in RULES.items():
        if s == v["title"] or s == k:
            return k
    for k, v in RULES.items():                       # без учёта двоеточий и окончаний
        if re.sub(r"[^а-яa-z ]", "", s) == re.sub(r"[^а-яa-z ]", "", v["title"]):
            return k
    return "?"


TYPE_KEY_WORD = {"лес": "лесные", "луг": "луговые", "водн": "водные", "околовод": "околоводные", "болот": "болотные"}


def types_label(keys) -> str:
    return ", ".join(TYPE_KEY_WORD.get(k, k) for k in keys)


def classify_method_text(measure: dict) -> dict:
    """Прежний способ: правило по ключевым словам «Метода выявления» (если столбец «Правило назначения» пуст)."""
    c = classify_measure(measure)
    r, pr = c["rule"], c["params"]
    if r is None:
        return dict(rule=None, p={}, reason=c["reason"])
    forest = (TYPE_FOREST, types_label(TYPE_FOREST))
    if r == "corvus":
        return dict(rule="species_density", p=dict(types=forest, taxon=["Corvus"], window=CORVUS_WINDOW_M))
    if r == "slope":
        return dict(rule="slope", p=dict(types=(TYPE_MEADOW, types_label(TYPE_MEADOW)) if pr.get("type_keys") else None,
                                         slope=pr["threshold_pct"]))
    if r == "invasive":
        return dict(rule="species", p=dict(category=["инвазивные"], sdm=SDM_THRESHOLD_INVASIVE))
    if r == "rare":
        return dict(rule="species_trails", p=dict(category=["охраняемые"], dist=RARE_BUFFER_M,
                                                  sdm=SDM_THRESHOLD_RARE))
    if r == "shore":
        return dict(rule="shore", p=dict(dist=SHORE_BUFFER_M, window=TRAIL_WINDOW_M, min_ha=MIN_PATCH_HA / 5))
    if r in ("trail_hotspot", "trail_density"):
        return dict(rule=r, p=dict(types=forest))
    return dict(rule="type", p=dict(types=(tuple(pr["keys"]), types_label(pr["keys"]))))


def measure_spec(measure: dict) -> dict:
    """Полное описание правила мероприятия: rule, p (все параметры), src (откуда взят параметр),
    origin («таблица» / «текст метода»), errors, reason (почему не рассчитывается)."""
    raw = measure.get("spec") or {}
    errors = []
    rule_raw = raw.get("rule")
    if rule_raw is not None and str(rule_raw).strip() and str(rule_raw).strip().lower() not in ("nan",):
        rule = rule_from_title(rule_raw)
        if rule == "?":
            return dict(rule=None, p={}, src={}, origin="таблица", errors=[],
                        reason=f"правило «{rule_raw}» не распознано; допустимы: " + "; ".join(RULE_TITLES))
        origin, given = "таблица", {}
        for k, v in raw.items():
            if k == "rule" or k not in RULES[rule]["params"]:
                continue
            val, err = _parse_value(k, v)
            if err:
                errors.append(f"{SPEC_TITLE[k]}: {err}")
            elif val is not None:
                given[k] = val
    else:
        lc = classify_method_text(measure)
        if lc["rule"] is None:
            return dict(rule=None, p={}, src={}, origin="текст метода", errors=[], reason=lc["reason"])
        rule, given, origin = lc["rule"], lc["p"], "текст метода"
    p, src = {}, {}
    for k, d in RULES[rule]["params"].items():
        if k in given:
            p[k], src[k] = given[k], origin
        elif d == REQ:
            p[k], src[k] = None, "не задано"
            errors.append(f"не задан обязательный параметр «{SPEC_TITLE[k]}»")
        elif d == SHARED:
            p[k], src[k] = None, "окно скрипта"
        else:
            p[k], src[k] = d, "по умолчанию"
    return dict(rule=rule, p=p, src=src, origin=origin, errors=errors, reason="")


def spec_share(p: dict) -> float:
    v = p.get("min_share")
    return float(_MIN_SHARE[0] if v is None else v)


def _types_txt(p, any_txt="полигоны любого типа") -> str:
    t = p.get("types")
    return any_txt if not t else f"полигоны типа «{t[1]}»"


def _q_txt(q) -> str:
    return "верхний квартиль (0,75)" if abs(float(q) - 0.75) < 1e-9 else f"квантиль {_fmt(q)}"


def _sel_txt(p) -> str:
    """Какие слои видов выбирает мероприятие."""
    cat = ", ".join(p.get("category") or []) or "любые"
    grp = ", ".join(p.get("groups") or []) or "любые"
    parts = [f"категория — {cat}", f"группы — {grp}"]
    if p.get("taxon"):
        parts.append("таксоны — " + ", ".join(p["taxon"]))
    if "sources" in p:
        parts.append("данные — " + (", ".join(SPECIES_KINDS[s] for s in p["sources"]) if p.get("sources")
                                    else "точки, полигоны, SDM"))
    return "; ".join(parts)


def _patch_txt(p) -> str:
    sh = spec_share(p)
    return (f"участок ≥ {_fmt(p['min_ha'])} га" + (f" и ≥ {_fmt(sh)}% площади полигона" if sh > 0 else ""))


def _uses(p, kind) -> bool:
    return not p.get("sources") or kind in p["sources"]


def describe_spec(spec: dict, full: bool = False) -> str:
    """Описание правила: короткое (журнал, отчёт, окно) или полное (столбец таблицы, окно)."""
    r, p = spec.get("rule"), spec.get("p", {})
    if r is None:
        return "не рассчитывается: " + spec.get("reason", "")
    if spec.get("errors"):
        return "не рассчитывается — ошибка в параметрах: " + "; ".join(spec["errors"])
    if not full:
        if r == "type":
            return f"все {_types_txt(p)}"
        if r == "trail_density":
            return f"плотность троп на полигон (м/га) ≥ {_q_txt(p['quantile'])}; {_types_txt(p)}"
        if r == "trail_hotspot":
            return (f"участки повышенной плотности троп (окно {_fmt(p['window'])} м, {_q_txt(p['quantile'])}); "
                    f"{_types_txt(p)}; {_patch_txt(p)}")
        if r == "shore":
            return (f"участки повышенной плотности троп (окно {_fmt(p['window'])} м, {_q_txt(p['quantile'])}) "
                    f"в {_fmt(p['dist'])} м от водных объектов; {_types_txt(p)}; {_patch_txt(p)}")
        if r == "slope":
            return f"уклон по ЦМР более {_fmt(p['slope'])}%; {_types_txt(p)}; {_patch_txt(p)}"
        sel = _sel_txt(p)
        sdm = (f"; SDM ≥ {_fmt(p['sdm'])} — отдельная строка «мониторинг» ({_patch_txt(p)})"
               if "sdm" in p and _uses(p, "sdm") else "")
        if r == "species":
            return f"находки видов ({sel}): полигоны с точками или полигонами находок{sdm}; {_types_txt(p)}"
        if r == "species_trails":
            return (f"тропы в {_fmt(p['dist'])} м от точек и в полигонах видов ({sel}){sdm}; {_types_txt(p)}")
        if r == "species_density":
            return (f"плотность точек видов ({sel}) в окне {_fmt(p['window'])} м ≥ {_q_txt(p['quantile'])}; "
                    f"{_types_txt(p)}; {_patch_txt(p)}")
    # ---- полное описание
    trails = ("слой троп (если в слое есть поле HIGHWAY или fclass, учитываются типы "
              + ", ".join(sorted(TRAIL_HIGHWAY_TYPES)) + ")")
    sh = spec_share(p)
    share = (f"не меньше {_fmt(p.get('min_ha'))} га и не меньше {_fmt(sh)}% площади полигона" if sh > 0
             else f"не меньше {_fmt(p.get('min_ha'))} га (порог доли полигона отключён)")
    tgt_each = "каждого полигона любого типа" if not p.get("types") else f"каждого полигона типа «{p['types'][1]}»"
    q = _q_txt(p.get("quantile", 0.75))
    hot = (f"По слою троп строится растр плотности: сетка {_fmt(CELL)} м, круговое скользящее окно радиусом "
           f"{_fmt(p.get('window'))} м, значение — длина троп в окне на гектар (м/га). По ячейкам внутри ООЗТ с "
           f"ненулевой плотностью вычисляется {q}; ячейки не ниже него объединяются в участки повышенной "
           f"плотности троп.")
    fin = (f"Мероприятие назначается, если такие участки занимают в полигоне {share}. В основании указываются "
           f"площадь участков и доля полигона; участки работ записываются в слой measure_zones.")
    sdm = (f"Модели SDM: выделяются участки с вероятностью не ниже {_fmt(p.get('sdm'))}; полигону назначается "
           f"отдельная строка «… — мониторинг (SDM ≥ {_fmt(p.get('sdm'))})», если участки одного вида занимают в нём "
           f"не меньше {_fmt(p.get('min_ha'))} га, а все вместе — "
           + (f"не меньше {_fmt(sh)}% площади полигона." if sh > 0 else "любую долю полигона.")
           ) if "sdm" in p and _uses(p, "sdm") else ""
    sp = (f"Слои видов выбираются по признакам, заданным при загрузке: {_sel_txt(p)}. Название вида берётся из "
          "поля с латинскими названиями (если его нет — из имени файла; у растров SDM — всегда из имени файла).")
    if r == "type":
        return (f"Назначается всем {('полигонам типа «' + p['types'][1] + '»') if p.get('types') else 'полигонам'} "
                "слоя территории. Тип определяется по началу слова (лес…, луг…, болот…); «водные» не включают "
                "«околоводные». Других данных не требуется.")
    if r == "trail_density":
        return (f"Данные: {trails}. Для {tgt_each} считается плотность троп — длина троп внутри полигона, "
                f"делённая на его площадь (м/га). Среди этих полигонов с ненулевой плотностью вычисляется {q}. "
                "Мероприятие назначается полигонам с плотностью не ниже порога; полигон оценивается целиком, "
                "порог доли не применяется. На схеме красным показаны тропы внутри зон повышенной плотности.")
    if r == "trail_hotspot":
        return f"Данные: {trails}. 1) {hot} 2) Для {tgt_each} считается площадь таких участков внутри него. {fin}"
    if r == "shore":
        return (f"Данные: {trails}; гидрография — линии и полигоны. 1) {hot} 2) Водные объекты объединяются, "
                f"вокруг них строится полоса шириной {_fmt(p['dist'])} м; сама акватория исключается. 3) Участки "
                f"повышенной плотности троп обрезаются этой полосой. 4) Для {tgt_each} считается площадь таких "
                f"участков. {fin}")
    if r == "slope":
        return (f"Данные: цифровая модель рельефа, приведённая к сетке {_fmt(CELL)} м. 1) Уклон считается в "
                "процентах после лёгкого сглаживания (окно 3 × 3 ячейки). 2) Отбираются ячейки внутри ООЗТ с "
                f"уклоном более {_fmt(p['slope'])}%; одиночные ячейки убираются (морфологическое открытие 2 × 2), "
                f"остальные объединяются в участки. 3) Для {tgt_each} считается площадь участков. {fin}")
    if r == "species":
        parts = [sp]
        if _uses(p, "points"):
            parts.append("Точки: мероприятие назначается каждому полигону, внутри которого есть хотя бы одна "
                         "точка, без порогов площади; в основании — виды, число точек и средние координаты "
                         "WGS 84 по каждому виду.")
        if _uses(p, "polygons"):
            parts.append("Полигоны находок: мероприятие назначается каждому полигону, который с ними "
                         "пересекается; указывается площадь пересечения по видам.")
        if sdm:
            parts.append(sdm)
        parts.append("Учитываются " + _types_txt(p) + ".")
        return " ".join(parts)
    if r == "species_trails":
        z = []
        if _uses(p, "points"):
            z.append(f"круг радиусом {_fmt(p['dist'])} м вокруг каждой точки")
        if _uses(p, "polygons"):
            z.append("сами полигоны находок")
        parts = [sp]
        if z:
            parts.append(f"1) Зоны видов — {' и '.join(z)}. 2) Тропы ({trails}) обрезаются по этим зонам. "
                         "3) Мероприятие назначается каждому полигону, в котором есть хотя бы фрагмент таких "
                         "троп, без порогов длины и доли; в основании — виды и длина троп в зоне, м. Если слой "
                         "троп не загружен, мероприятие назначается полигонам, пересекающим сами зоны.")
        if sdm:
            parts.append(sdm)
        parts.append("Учитываются " + _types_txt(p) + ".")
        return " ".join(parts)
    if r == "species_density":
        return (f"{sp} Используются только точки. 1) Строится растр плотности: сетка {_fmt(CELL)} м, круговое окно "
                f"радиусом {_fmt(p['window'])} м, значение — число точек на гектар. 2) По ячейкам внутри ООЗТ с "
                f"ненулевой плотностью вычисляется {q}; ячейки не ниже него объединяются в участки. 3) Для "
                f"{tgt_each} считается площадь участков. {fin}")
    return ""


def describe_rule(measure: dict, full: bool = False) -> str:
    return describe_spec(measure_spec(measure), full)


def spec_params_lines(spec: dict) -> list[str]:
    """Параметры правила с источником значения — для окна скрипта."""
    out = []
    for k, v in spec.get("p", {}).items():
        src = spec["src"].get(k, "")
        if k == "types":
            val = v[1] if v else "любые"
        elif k in ("category", "groups", "taxon"):
            val = ", ".join(v) if v else "любые"
        elif k == "sources":
            val = ", ".join(SPECIES_KINDS[s] for s in v) if v else "точки, полигоны, SDM"
        elif k == "min_share":
            val = f"{_fmt(spec_share(spec['p']))}"
        else:
            val = _fmt(v)
        out.append(f"{SPEC_TITLE[k]}: {val}  ({src})")
    return out


def _cfg_has(cfg: Config, key: str) -> bool:
    v = getattr(cfg, key, None)
    return bool(v) if not isinstance(v, list) else len(v) > 0


def matching_species_layers(cfg, p: dict, kinds=None) -> list[dict]:
    """Загруженные слои видов, подходящие мероприятию по категории, группе и типу данных."""
    out = []
    for l in species_layers(cfg):
        if p.get("category") and l["category"] not in p["category"]:
            continue
        if not group_matches(l["group"], p.get("groups")):
            continue
        if kinds is not None and l["kind"] not in kinds:
            continue
        if p.get("sources") and l["kind"] not in p["sources"]:
            continue
        out.append(l)
    return out


DATA_KEYS = {"dem": ["dem_path"], "trails": ["trails"], "hydro": ["hydro_lines", "hydro_polys"],
             "species": ["species"]}


def measure_status(measure: dict, cfg: Config, type_counts: dict | None = None) -> dict:
    """Статус мероприятия при текущем наборе данных: ok / missing / off / na (+ текст, нужные слои)."""
    spec = measure_spec(measure)
    rule, p = spec["rule"], spec.get("p", {})
    need, layers = [], []
    if rule:
        for d in RULES[rule]["data"]:
            if d == "species":
                need.append("слои видов (" + _sel_txt(p) + ")")
            else:
                need.append(" или ".join(LAYER_TITLE[k] for k in DATA_KEYS[d]))
        if p.get("types"):
            need.append("поле типа в слое территории")
    need_txt = "; ".join(need) if need else "только слой территории"
    base = dict(need=need_txt, rule=rule, spec=spec, layers=layers)
    if not measure.get("on"):
        return dict(base, state="off", text="выключено в таблице")
    if rule is None:
        return dict(base, state="na", text=spec["reason"], need="—")
    if spec["errors"]:
        return dict(base, state="missing", text="ошибка в таблице: " + "; ".join(spec["errors"]))
    missing = []
    for d in RULES[rule]["data"]:
        if d == "species":
            kinds = ["points"] if rule == "species_density" else None
            layers[:] = matching_species_layers(cfg, p, kinds)
            if not layers:
                missing.append("подходящие слои видов (" + _sel_txt(p) + ")")
        elif not any(_cfg_has(cfg, k) for k in DATA_KEYS[d]):
            missing.append(" или ".join(LAYER_TITLE[k] for k in DATA_KEYS[d]))
    if missing:
        return dict(base, state="missing", text="не хватает: " + "; ".join(missing))
    txt = "будет рассчитано"
    keys = (p.get("types") or (None,))[0]
    if keys and type_counts is not None:
        n = sum(v for t, v in type_counts.items() if type_matches(t, keys))
        if n == 0:
            return dict(base, state="missing", text="в слое нет полигонов типа «%s»" % p["types"][1])
        txt += f" ({n} полигонов типа «{p['types'][1]}»)"
    if layers:
        txt += f"; слоёв видов: {len(layers)}"
    if rule == "species_trails" and not cfg.trails:
        txt += " — без троп: по самим зонам видов"
    return dict(base, state="ok", text=txt)


# =====================================================================================
#                       ОКНО ВЫБОРА ИСХОДНЫХ ДАННЫХ (tkinter, одно окно)
# =====================================================================================
VEC_TYPES = [("Векторные данные", "*.shp *.gpkg *.geojson *.json *.kml *.dxf *.gml"), ("Все файлы", "*.*")]
RAS_TYPES = [("Растры", "*.tif *.tiff *.asc *.img *.vrt"), ("Все файлы", "*.*")]
XLS_TYPES = [("Excel", "*.xlsx *.xlsm *.xls"), ("Все файлы", "*.*")]
JSON_TYPES = [("Конфигурация", "*.json"), ("Все файлы", "*.*")]
NONE_FIELD = "— нет —"

# поля слоя территории: ключ, подпись, обязательность, кандидаты для автоподбора
FIELD_SPECS = {
    "kv_vyd": [
        ("OOZT", "ООЗТ", True, ["OOZT", "ООЗТ", "Name_OOZT"]),
        ("KV", "Квартал", True, ["KV", "FID_kv", "kv", "Квартал", "kvartal"]),
        ("VYD", "Выдел", True, ["VYD", "FID_ecosys", "vyd", "Выдел", "vydel"]),
        ("TYPE", "Тип (лес / луг / водные / околоводные)", True, ["Type", "TYPE", "Тип"]),
        ("NAME", "Название экосистемы", False, ["Name", "NAME", "Название", "Ecosystem"]),
    ],
    "eco": [
        ("ID", "Уникальный ID экосистемы", True, ["ID", "Id", "id", "eco_id", "ECO_ID", "uid", "UID",
                                                  "FID_ecosys", "fid_1"]),
        ("TYPE", "Тип (лес / луг / водные / околоводные)", True, ["Type", "TYPE", "Тип"]),
        ("NAME", "Название экосистемы", False, ["Name", "NAME", "Название", "Ecosystem"]),
        ("OOZT", "ООЗТ (если поля нет — название ниже)", False, ["OOZT", "ООЗТ", "Name_OOZT"]),
    ],
}


def _vector_fields(path: str, layer: str | None):
    """Имена полей и таблица атрибутов (без геометрии) — для подбора полей в интерфейсе."""
    df = pyogrio.read_dataframe(path, layer=layer, read_geometry=False)
    ren = {c: _remojibake(c) for c in df.columns}
    ren = {k: v for k, v in ren.items() if v != k and v not in df.columns}
    if ren:
        df = df.rename(columns=ren)
    return list(df.columns), df


class SetupWindow:
    """Одно окно: режим территории, обязательные и тематические слои, живая сводка мероприятий."""

    C_OK, C_WARN, C_OFF, C_MUTED = "#2E7D32", "#B35C00", "#8A8A8A", "#5F6368"

    def __init__(self, cfg: Config | None = None):
        import tkinter as tk
        from tkinter import ttk, font as tkfont
        self.tk, self.ttk = tk, ttk
        if os.name == "nt":
            try:
                import ctypes
                ctypes.windll.shcore.SetProcessDpiAwareness(1)
            except Exception:
                pass
        self.root = tk.Tk()
        self.root.title("ООЗТ: природоохранные мероприятия — исходные данные")
        self.root.geometry("1480x900")
        self.root.minsize(1180, 700)
        self.cfg = cfg or Config()
        self.result: Config | None = None
        self.measures: list[dict] = []
        self.type_counts: dict | None = None
        self.attr_df = None
        self.field_vars: dict = {}
        self.layer_rows: dict = {}

        fam = "Segoe UI" if os.name == "nt" else ("Noto Sans" if "Noto Sans" in tkfont.families() else None)
        for fn in ("TkDefaultFont", "TkTextFont", "TkHeadingFont", "TkMenuFont", "TkCaptionFont"):
            try:
                f_ = tkfont.nametofont(fn)
                f_.configure(size=10, **({"family": fam} if fam else {}))
            except Exception:
                pass
        base = tkfont.nametofont("TkDefaultFont")
        self.f_bold = base.copy(); self.f_bold.configure(weight="bold")
        self.f_small = base.copy(); self.f_small.configure(size=9)
        self.f_head = base.copy(); self.f_head.configure(size=13, weight="bold")

        st = ttk.Style(self.root)
        st.configure("Treeview", rowheight=24)
        # обход ошибки Tk 8.6.9: цвета тегов Treeview не применяются
        def _fixed(opt):
            return [e for e in st.map("Treeview", query_opt=opt) if e[:2] != ("!disabled", "!selected")]
        st.map("Treeview", foreground=_fixed("foreground"), background=_fixed("background"))

        self._build()
        self._load_from_cfg(self.cfg)

    # ------------------------------------------------------------------ разметка
    def _build(self):
        tk, ttk = self.tk, self.ttk
        top = ttk.Frame(self.root, padding=(12, 10, 12, 4))
        top.pack(fill="x")
        ttk.Label(top, text="Исходные данные для расчёта мероприятий", font=self.f_head).pack(anchor="w")
        ttk.Label(top, foreground=self.C_MUTED, text=(
            "Обязательны таблица мероприятий и слой территории. Остальные слои необязательны: справа видно, "
            "какие мероприятия будут рассчитаны при выбранном наборе данных и чего не хватает остальным.")
        ).pack(anchor="w")

        pw = ttk.PanedWindow(self.root, orient="horizontal")
        pw.pack(fill="both", expand=True, padx=12, pady=6)

        # ---- левая колонка (прокручиваемая)
        left_outer = ttk.Frame(pw)
        canvas = tk.Canvas(left_outer, highlightthickness=0, borderwidth=0)
        vsb = ttk.Scrollbar(left_outer, orient="vertical", command=canvas.yview)
        self.left = ttk.Frame(canvas)
        self.left.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        win = canvas.create_window((0, 0), window=self.left, anchor="nw")
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(win, width=e.width))
        canvas.configure(yscrollcommand=vsb.set)
        canvas.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        canvas.bind_all("<MouseWheel>", lambda e: canvas.yview_scroll(int(-e.delta / 120), "units"))
        canvas.bind_all("<Button-4>", lambda e: canvas.yview_scroll(-2, "units"))
        canvas.bind_all("<Button-5>", lambda e: canvas.yview_scroll(2, "units"))
        pw.add(left_outer, weight=5)

        # 1. таблица мероприятий
        s1 = ttk.LabelFrame(self.left, text=" 1. Таблица мероприятий (обязательно) ", padding=8)
        s1.pack(fill="x", pady=(0, 8))
        self.v_xlsx = tk.StringVar()
        r = ttk.Frame(s1); r.pack(fill="x")
        ttk.Entry(r, textvariable=self.v_xlsx, state="readonly").pack(side="left", fill="x", expand=True)
        ttk.Button(r, text="Выбрать…", command=self._pick_xlsx).pack(side="left", padx=(6, 0))
        self.l_xlsx = ttk.Label(s1, foreground=self.C_MUTED, font=self.f_small,
                                text="Excel «Мероприятия по Приказу»: колонки «Мероприятия», «Вкл/выкл», «Метод выявления»")
        self.l_xlsx.pack(anchor="w", pady=(4, 0))
        ttk.Label(s1, foreground=self.C_MUTED, font=self.f_small, wraplength=700, justify="left",
                  text="Способ назначения задаётся столбцами таблицы: «Правило назначения», «Типы полигонов», "
                       "«Виды: категория», «Виды: группы организмов», «Виды: таксоны (лат.)», «Виды: данные», "
                       "«Расстояние, м», «Окно плотности, м», «Квантиль», «Уклон более, %», «Порог SDM (0–1)», "
                       "«Мин. площадь участка, га», «Мин. доля полигона, %». Если «Правило назначения» пусто, "
                       "правило определяется по тексту «Метода выявления». Справка — на листе «Справка по правилам»."
                  ).pack(anchor="w", pady=(2, 0))
        r = ttk.Frame(s1); r.pack(fill="x", pady=(4, 0))
        ttk.Button(r, text="Записать правила и описания в таблицу…", command=self._update_table).pack(side="left")
        ttk.Label(r, foreground=self.C_MUTED, font=self.f_small,
                  text="  столбцы правил, справка, описания расчёта").pack(side="left")

        # 2. слой территории
        s2 = ttk.LabelFrame(self.left, text=" 2. Слой территории — единицы назначения мероприятий (обязательно) ",
                            padding=8)
        s2.pack(fill="x", pady=(0, 8))
        self.v_mode = tk.StringVar(value="kv_vyd")
        mf = ttk.Frame(s2); mf.pack(fill="x")
        ttk.Radiobutton(mf, text="Кварталы и выделы", value="kv_vyd", variable=self.v_mode,
                        command=self._on_mode).grid(row=0, column=0, sticky="w")
        ttk.Label(mf, foreground=self.C_MUTED, font=self.f_small,
                  text="полигоны выделов с полями ООЗТ, квартала и выдела; отчёт и схемы — по кварталам"
                  ).grid(row=1, column=0, sticky="w", padx=(24, 0))
        ttk.Radiobutton(mf, text="Слой экосистем с уникальным ID (без кварталов)", value="eco",
                        variable=self.v_mode, command=self._on_mode).grid(row=2, column=0, sticky="w", pady=(4, 0))
        ttk.Label(mf, foreground=self.C_MUTED, font=self.f_small,
                  text="полигоны экосистем; отчёт и схемы — по листам регулярной сетки или по полю группировки"
                  ).grid(row=3, column=0, sticky="w", padx=(24, 0))

        self.v_oozt = tk.StringVar()
        r = ttk.Frame(s2); r.pack(fill="x", pady=(8, 0))
        ttk.Label(r, text="Файл:", width=8).pack(side="left")
        ttk.Entry(r, textvariable=self.v_oozt, state="readonly").pack(side="left", fill="x", expand=True)
        ttk.Button(r, text="Выбрать…", command=self._pick_oozt).pack(side="left", padx=(6, 0))
        r = ttk.Frame(s2); r.pack(fill="x", pady=(4, 0))
        ttk.Label(r, text="Слой GPKG:", width=11).pack(side="left")
        self.v_layer = tk.StringVar()
        self.cb_layer = ttk.Combobox(r, textvariable=self.v_layer, state="disabled", width=40)
        self.cb_layer.pack(side="left")
        self.cb_layer.bind("<<ComboboxSelected>>", lambda e: self._read_oozt_fields())

        self.fields_frame = ttk.Frame(s2)
        self.fields_frame.pack(fill="x", pady=(8, 0))
        self.eco_frame = ttk.Frame(s2)
        self.v_oozt_name = tk.StringVar()
        self.v_group = tk.StringVar(value=NONE_FIELD)
        self.v_sheet = tk.StringVar(value="1000")
        ttk.Label(self.eco_frame, text="Название ООЗТ:").grid(row=0, column=0, sticky="w")
        ttk.Entry(self.eco_frame, textvariable=self.v_oozt_name, width=30).grid(row=0, column=1, sticky="w", padx=6)
        ttk.Label(self.eco_frame, foreground=self.C_MUTED, font=self.f_small,
                  wraplength=300, justify="left", text="используется, если поле ООЗТ не выбрано").grid(row=0, column=2, sticky="w")
        ttk.Label(self.eco_frame, text="Группировка отчёта:").grid(row=1, column=0, sticky="w", pady=(4, 0))
        self.cb_group = ttk.Combobox(self.eco_frame, textvariable=self.v_group, state="readonly", width=28)
        self.cb_group.grid(row=1, column=1, sticky="w", padx=6, pady=(4, 0))
        self.cb_group.bind("<<ComboboxSelected>>", lambda e: self._refresh())
        ttk.Label(self.eco_frame, foreground=self.C_MUTED, font=self.f_small,
                  wraplength=300, justify="left", text="поле (например, участок или район); «нет» — листы регулярной сетки"
                  ).grid(row=1, column=2, sticky="w")
        ttk.Label(self.eco_frame, text="Размер листа, м:").grid(row=2, column=0, sticky="w", pady=(4, 0))
        ttk.Spinbox(self.eco_frame, from_=200, to=5000, increment=100, textvariable=self.v_sheet, width=8,
                    command=self._refresh).grid(row=2, column=1, sticky="w", padx=6, pady=(4, 0))
        ttk.Label(self.eco_frame, foreground=self.C_MUTED, font=self.f_small,
                  wraplength=300, justify="left", text="лист = квадрат сетки; полигон относится к листу по своей внутренней точке"
                  ).grid(row=2, column=2, sticky="w")
        self.l_types = ttk.Label(s2, foreground=self.C_MUTED, font=self.f_small, wraplength=700, justify="left")
        self.l_types.pack(anchor="w", pady=(6, 0))

        # 3. тематические слои
        s3 = ttk.LabelFrame(self.left, text=" 3. Тематические слои (необязательно) ", padding=8)
        s3.pack(fill="x", pady=(0, 8))
        ttk.Label(s3, foreground=self.C_MUTED, font=self.f_small, wraplength=700, justify="left",
                  text="✔ — слой выбран;  ! — нужен включённым мероприятиям, но не выбран;  ○ — не используется "
                       "включёнными мероприятиями. Коды мероприятий — по таблице (колонка с номерами I.1, II.1…).",
                  ).pack(anchor="w", pady=(0, 4))
        for spec in LAYER_SPECS:
            self._layer_row(s3, spec)
        self._species_block(s3)

        # ---- правая колонка
        right = ttk.Frame(pw, padding=(8, 0, 0, 0))
        pw.add(right, weight=4)
        self._pw = pw
        self.root.bind("<Map>", self._place_sash, add="+")
        s4 = ttk.LabelFrame(right, text=" 4. Что будет рассчитано ", padding=8)
        s4.pack(fill="both", expand=True)
        self.l_sum = ttk.Label(s4, font=self.f_bold, text="Выберите таблицу мероприятий")
        self.l_sum.pack(anchor="w", pady=(0, 6))
        tf = ttk.Frame(s4); tf.pack(fill="both", expand=True)
        cols = ("code", "name", "status")
        self.tree = ttk.Treeview(tf, columns=cols, show="headings", selectmode="browse")
        for c, t, w, st in (("code", "Код", 60, False), ("name", "Мероприятие", 300, True),
                            ("status", "Статус при выбранных данных", 300, True)):
            self.tree.heading(c, text=t, anchor="w")
            self.tree.column(c, width=w, stretch=st, anchor="w")
        self.tree.tag_configure("ok", foreground=self.C_OK)
        self.tree.tag_configure("missing", foreground=self.C_WARN)
        self.tree.tag_configure("off", foreground=self.C_OFF)
        self.tree.tag_configure("na", foreground=self.C_OFF)
        ys = ttk.Scrollbar(tf, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=ys.set)
        self.tree.pack(side="left", fill="both", expand=True)
        ys.pack(side="right", fill="y")
        self.tree.bind("<<TreeviewSelect>>", self._on_select)
        df_ = ttk.Frame(s4); df_.pack(fill="both", expand=False, pady=(6, 0))
        self.detail = tk.Text(df_, height=14, wrap="word", relief="flat", background="#F4F5F6",
                              font=self.f_small, padx=6, pady=4)
        dsb = ttk.Scrollbar(df_, orient="vertical", command=self.detail.yview)
        self.detail.configure(yscrollcommand=dsb.set)
        self.detail.pack(side="left", fill="both", expand=True)
        dsb.pack(side="right", fill="y")
        self.detail.insert("1.0", "Выберите мероприятие в списке, чтобы увидеть правило назначения, его параметры, "
                                  "полное описание расчёта и нужные слои.")
        self.detail.configure(state="disabled")

        s5 = ttk.LabelFrame(right, text=" 5. Результаты ", padding=8)
        s5.pack(fill="x", pady=(8, 0))
        self.v_out = tk.StringVar()
        r = ttk.Frame(s5); r.pack(fill="x")
        ttk.Label(r, text="Папка:", width=8).pack(side="left")
        ttk.Entry(r, textvariable=self.v_out).pack(side="left", fill="x", expand=True)
        ttk.Button(r, text="Выбрать…", command=self._pick_out).pack(side="left", padx=(6, 0))
        r = ttk.Frame(s5); r.pack(fill="x", pady=(6, 0))
        self.l_kv = ttk.Label(r, text="Кварталы для отчёта:", width=22)
        self.l_kv.pack(side="left")
        self.v_kv = tk.StringVar()
        ttk.Entry(r, textvariable=self.v_kv, width=30).pack(side="left")
        ttk.Label(r, foreground=self.C_MUTED, font=self.f_small, text="  через запятую; пусто — все").pack(side="left")
        self.v_basemap = tk.BooleanVar(value=True)
        ttk.Checkbutton(s5, text="Подложка OpenStreetMap на схемах (нужен интернет)",
                        variable=self.v_basemap).pack(anchor="w", pady=(6, 0))
        r = ttk.Frame(s5); r.pack(fill="x", pady=(4, 0))
        ttk.Label(r, text="Общий порог: участок-основание не меньше").pack(side="left")
        self.v_minshare = tk.StringVar(value="25")
        ttk.Entry(r, textvariable=self.v_minshare, width=5).pack(side="left", padx=4)
        ttk.Label(r, text="% площади полигона (0 — без порога; если в таблице не указан свой)").pack(side="left")
        self.v_split = tk.BooleanVar(value=True)
        ttk.Checkbutton(s5, text="Делить многочастные объекты на отдельные контуры (ID-1, ID-2, …) — "
                                 "мероприятие назначается каждому контуру",
                        variable=self.v_split).pack(anchor="w", pady=(4, 0))
        r = ttk.Frame(s5); r.pack(fill="x", pady=(6, 0))
        ttk.Label(r, text="Отчёт:", width=8).pack(side="left")
        self.v_formats = {k: tk.BooleanVar(value=k in ("docx", "pdf")) for k in REPORT_FORMATS}
        for k, t in REPORT_FORMATS.items():
            ttk.Checkbutton(r, text=t, variable=self.v_formats[k]).pack(side="left", padx=(0, 10))
        ttk.Label(s5, foreground=self.C_MUTED, font=self.f_small, wraplength=640, justify="left",
                  text="report.md создаётся всегда. PDF и DOC делаются через Microsoft Word (если установлен), "
                       "иначе через LibreOffice; без них PDF собирается встроенной вёрсткой, DOC — не создаётся."
                  ).pack(anchor="w", pady=(2, 0))

        # ---- кнопки
        bot = ttk.Frame(self.root, padding=(12, 4, 12, 12))
        bot.pack(fill="x")
        ttk.Button(bot, text="Загрузить конфигурацию…", command=self._load_cfg_file).pack(side="left")
        ttk.Button(bot, text="Сохранить конфигурацию…", command=self._save_cfg_file).pack(side="left", padx=6)
        ttk.Button(bot, text="Отмена", command=self._cancel).pack(side="right")
        ttk.Button(bot, text="Запустить расчёт", command=self._run).pack(side="right", padx=6)
        self.root.protocol("WM_DELETE_WINDOW", self._cancel)

    def _place_sash(self, _e=None):
        if getattr(self, "_sash_done", False):
            return
        self.root.update_idletasks()
        w = self._pw.winfo_width()
        if w > 100:
            self._pw.sashpos(0, int(w * 0.53))
            self._sash_done = True

    def _layer_row(self, parent, spec):
        ttk = self.ttk
        key = spec["key"]
        card = ttk.Frame(parent, padding=(0, 5, 0, 5))
        card.pack(fill="x")
        ttk.Separator(parent, orient="horizontal").pack(fill="x")
        card.columnconfigure(1, weight=1)
        st = ttk.Label(card, text="○", width=2, foreground=self.C_OFF, font=self.f_bold)
        st.grid(row=0, column=0, sticky="w")
        ttk.Label(card, text=spec["title"], font=self.f_bold).grid(row=0, column=1, sticky="w")
        bf = ttk.Frame(card)
        bf.grid(row=0, column=2, rowspan=2, sticky="ne")
        ttk.Button(bf, text="Добавить…" if spec["multi"] else "Выбрать…", width=11,
                   command=lambda s_=spec: self._pick_layer(s_)).pack(side="left")
        ttk.Button(bf, text="Очистить", width=9, command=lambda k=key: self._clear_layer(k)).pack(side="left", padx=(4, 0))
        ttk.Label(card, text=spec["what"] + ("; можно несколько файлов" if spec["multi"] else ""),
                  foreground=self.C_MUTED, font=self.f_small, wraplength=520, justify="left"
                  ).grid(row=1, column=1, sticky="w")
        used = ttk.Label(card, text="", font=self.f_small, wraplength=640, justify="left")
        used.grid(row=2, column=1, columnspan=2, sticky="w")
        files = ttk.Label(card, text="не выбран", foreground=self.C_OFF, font=self.f_small,
                          wraplength=640, justify="left")
        files.grid(row=3, column=1, columnspan=2, sticky="w")
        self.layer_rows[key] = dict(status=st, used=used, files=files, spec=spec)

    # ------------------------------------------------------------------ слои видов
    SP_TYPES = [("Слои видов (векторы и растры SDM)",
                 "*.shp *.gpkg *.geojson *.json *.kml *.gml *.tif *.tiff *.asc *.img *.vrt"), ("Все файлы", "*.*")]

    def _species_block(self, parent):
        tk, ttk = self.tk, self.ttk
        card = ttk.Frame(parent, padding=(0, 6, 0, 4))
        card.pack(fill="x")
        card.columnconfigure(1, weight=1)
        self.sp_status = ttk.Label(card, text="○", width=2, foreground=self.C_OFF, font=self.f_bold)
        self.sp_status.grid(row=0, column=0, sticky="w")
        ttk.Label(card, text="Виды и редкие экосистемы", font=self.f_bold).grid(row=0, column=1, sticky="w")
        bf = ttk.Frame(card)
        bf.grid(row=0, column=2, sticky="ne")
        ttk.Button(bf, text="Добавить…", width=11, command=self._sp_add).pack(side="left")
        ttk.Button(bf, text="Удалить", width=9, command=self._sp_remove).pack(side="left", padx=(4, 0))
        ttk.Label(card, foreground=self.C_MUTED, font=self.f_small, wraplength=640, justify="left",
                  text="Точки и полигоны находок (shp, gpkg…) или модели SDM (растры 0–1); можно выбрать сразу "
                       "несколько файлов. Для каждого слоя укажите категорию (инвазивные / охраняемые / прочие) и "
                       "группу организмов; тип данных определяется автоматически. Мероприятия выбирают слои по "
                       "этим признакам — столбцы «Виды: …» таблицы мероприятий."
                  ).grid(row=1, column=1, columnspan=2, sticky="w")
        tf = ttk.Frame(card)
        tf.grid(row=2, column=1, columnspan=2, sticky="we", pady=(4, 0))
        cols = ("file", "kind", "category", "group", "field")
        self.sp_tree = ttk.Treeview(tf, columns=cols, show="headings", selectmode="extended", height=5)
        for c, t, w in (("file", "Файл", 190), ("kind", "Тип данных", 130), ("category", "Категория", 100),
                        ("group", "Группа организмов", 150), ("field", "Поле названия", 90)):
            self.sp_tree.heading(c, text=t, anchor="w")
            self.sp_tree.column(c, width=w, stretch=(c == "file"), anchor="w")
        ys = ttk.Scrollbar(tf, orient="vertical", command=self.sp_tree.yview)
        self.sp_tree.configure(yscrollcommand=ys.set)
        self.sp_tree.pack(side="left", fill="x", expand=True)
        ys.pack(side="right", fill="y")
        self.sp_tree.bind("<<TreeviewSelect>>", self._sp_on_select)
        ed = ttk.Frame(card)
        ed.grid(row=3, column=1, columnspan=2, sticky="w", pady=(4, 0))
        self.v_sp_cat = tk.StringVar(value="охраняемые")
        self.v_sp_grp = tk.StringVar(value="сосудистые растения")
        self.v_sp_kind = tk.StringVar(value="")
        self.v_sp_field = tk.StringVar(value="авто")
        ttk.Label(ed, text="Категория:").grid(row=0, column=0, sticky="w")
        ttk.Combobox(ed, textvariable=self.v_sp_cat, values=SPECIES_CATEGORIES, state="readonly",
                     width=12).grid(row=0, column=1, sticky="w", padx=(4, 10))
        ttk.Label(ed, text="Группа:").grid(row=0, column=2, sticky="w")
        ttk.Combobox(ed, textvariable=self.v_sp_grp, values=SPECIES_GROUPS, width=20
                     ).grid(row=0, column=3, sticky="w", padx=(4, 10))
        ttk.Label(ed, text="Тип данных:").grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Combobox(ed, textvariable=self.v_sp_kind, values=list(SPECIES_KINDS.values()), state="readonly",
                     width=14).grid(row=1, column=1, sticky="w", padx=(4, 10), pady=(4, 0))
        ttk.Label(ed, text="Поле названия:").grid(row=1, column=2, sticky="w", pady=(4, 0))
        self.cb_sp_field = ttk.Combobox(ed, textvariable=self.v_sp_field, values=["авто"], width=20)
        self.cb_sp_field.grid(row=1, column=3, sticky="w", padx=(4, 10), pady=(4, 0))
        ttk.Button(ed, text="Применить к выбранным", command=self._sp_apply).grid(row=0, column=4, rowspan=2,
                                                                                 sticky="w")
        ttk.Label(card, foreground=self.C_MUTED, font=self.f_small, wraplength=640, justify="left",
                  text="Группу можно ввести вручную (новые группы организмов). «Поле названия: авто» — поле с "
                       "латинскими названиями ищется само; для редких экосистем выберите поле с названием. "
                       "У растров SDM название вида — имя файла."
                  ).grid(row=4, column=1, columnspan=2, sticky="w", pady=(2, 0))
        self.sp_used = ttk.Label(card, text="", font=self.f_small, wraplength=640, justify="left")
        self.sp_used.grid(row=5, column=1, columnspan=2, sticky="w", pady=(2, 0))
        ttk.Separator(parent, orient="horizontal").pack(fill="x")

    @staticmethod
    def _sp_guess(path: str, cat: str, grp: str):
        n = Path(path).stem.lower()
        if any(w in n for w in ("invas", "инваз", "alien", "чужерод")):
            cat = "инвазивные"
        elif any(w in n for w in ("rare", "redbook", "red_book", "редк", "красн", "охран", "protect")):
            cat = "охраняемые"
        for words, g in ((("vascular", "plant", "сосуд", "flora", "растен"), "сосудистые растения"),
                         (("fung", "гриб", "myco"), "грибы"), (("lichen", "лишай"), "лишайники"),
                         (("mamm", "млекоп", "theri"), "млекопитающие"), (("aves", "bird", "птиц", "орнит"), "птицы"),
                         (("amph", "rept", "herp", "fish", "земновод", "пресмык", "рыб", "низш"), "низшие позвоночные"),
                         (("insect", "насеком", "entom"), "насекомые"),
                         (("ecosys", "экосист", "habitat", "сообществ"), "редкие экосистемы")):
            if any(w in n for w in words):
                grp = g
                break
        return cat, grp

    def _sp_add(self):
        from tkinter import filedialog
        ps = list(filedialog.askopenfilenames(parent=self.root, title="Слои видов и редких экосистем — можно "
                                              "выбрать несколько", filetypes=self.SP_TYPES))
        if not ps:
            return
        cur = list(self.cfg.species_layers or [])
        have = {l.get("path") for l in cur}
        for pth in ps:
            if pth in have:
                continue
            cat, grp = self._sp_guess(pth, self.v_sp_cat.get(), norm_group(self.v_sp_grp.get()))
            cur.append(dict(path=pth, category=cat, group=grp, kind=detect_species_kind(pth), field=""))
        self.cfg.species_layers = cur
        self._sp_render()
        self._refresh()

    def _sp_remove(self):
        sel = {int(i) for i in self.sp_tree.selection()}
        if not sel:
            return
        self.cfg.species_layers = [l for i, l in enumerate(self.cfg.species_layers or []) if i not in sel]
        self._sp_render()
        self._refresh()

    def _sp_on_select(self, _e=None):
        sel = [int(i) for i in self.sp_tree.selection()]
        if not sel:
            return
        l = (self.cfg.species_layers or [])[sel[0]]
        self.v_sp_cat.set(norm_category(l.get("category")) or "прочие")
        self.v_sp_grp.set(norm_group(l.get("group")))
        self.v_sp_kind.set(SPECIES_KINDS.get(l.get("kind") or "", ""))
        self.v_sp_field.set(l.get("field") or "авто")
        vals = ["авто"]
        if len(sel) == 1 and l.get("kind") in ("points", "polygons"):
            try:
                pth = str(l["path"])
                cols, _ = _vector_fields(pth.split("|")[0], pth.split("|")[1] if "|" in pth else None)
                vals += [c for c in cols if c.lower() not in ("geometry",)]
            except Exception:
                pass
        self.cb_sp_field.configure(values=vals)

    def _sp_apply(self):
        sel = [int(i) for i in self.sp_tree.selection()]
        if not sel:
            from tkinter import messagebox
            messagebox.showinfo("Слои видов", "Выберите в списке один или несколько слоёв.", parent=self.root)
            return
        kind = {v: k for k, v in SPECIES_KINDS.items()}.get(self.v_sp_kind.get())
        fld = self.v_sp_field.get().strip()
        for i in sel:
            l = self.cfg.species_layers[i]
            l["category"] = norm_category(self.v_sp_cat.get()) or "прочие"
            l["group"] = norm_group(self.v_sp_grp.get())
            if kind:
                l["kind"] = kind
            if len(sel) == 1:
                l["field"] = "" if fld in ("", "авто") else fld
        self._sp_render()
        for i in sel:
            self.sp_tree.selection_add(str(i))
        self._refresh()

    def _sp_render(self):
        self.sp_tree.delete(*self.sp_tree.get_children())
        for i, l in enumerate(self.cfg.species_layers or []):
            kind = l.get("kind") or detect_species_kind(l.get("path", ""))
            l["kind"] = kind
            self.sp_tree.insert("", "end", iid=str(i), values=(
                Path(str(l.get("path", "")).split("|")[0]).name, SPECIES_KINDS.get(kind, "? не определён"),
                norm_category(l.get("category")) or "прочие", norm_group(l.get("group")), l.get("field") or "авто"))

    def _sp_refresh_status(self, statuses):
        n = len(self.cfg.species_layers or [])
        parts, bad = [], False
        for m, s in zip(self.measures, statuses):
            if not (s["rule"] and "species" in RULES[s["rule"]]["data"]) or not m["on"]:
                continue
            k = len(matching_species_layers(self.cfg, s["spec"].get("p", {}),
                                            ["points"] if s["rule"] == "species_density" else None))
            parts.append(f"{m.get('code') or '?'} ({_sel_txt(s['spec'].get('p', {}))}) — слоёв: {k}")
            bad |= k == 0
        txt = ("Нужны для включённых: " + ";  ".join(parts)) if parts else "Не нужны включённым мероприятиям"
        self.sp_used.configure(text=txt, foreground="#202124" if parts else self.C_MUTED)
        if n and not bad:
            self.sp_status.configure(text="✔", foreground=self.C_OK)
        elif parts and bad:
            self.sp_status.configure(text="!", foreground=self.C_WARN)
        else:
            self.sp_status.configure(text="✔" if n else "○", foreground=self.C_OK if n else self.C_OFF)

    def _update_table(self):
        from tkinter import messagebox
        if not self.cfg.measures_xlsx:
            messagebox.showinfo("Таблица мероприятий", "Сначала выберите таблицу мероприятий.", parent=self.root)
            return
        if not messagebox.askyesno(
                "Записать правила и описания в таблицу",
                "В таблицу мероприятий будут записаны:\n\n"
                "• столбцы правил («Правило назначения», «Типы полигонов», «Виды: …», «Расстояние, м» и др.) — "
                "только в строках, где «Правило назначения» пусто: правило и параметры, которые скрипт сейчас "
                "определяет по тексту «Метода выявления»;\n"
                "• выпадающие списки в этих столбцах;\n"
                "• лист «Справка по правилам»;\n"
                f"• столбец «{DESC_COLUMN}» — полное описание расчёта для каждой строки.\n\n"
                "Перед записью создаётся копия таблицы с окончанием «_до_правил». Файл должен быть закрыт в Excel. "
                "Продолжить?", parent=self.root):
            return
        try:
            info = update_measures_xlsx(self.cfg.measures_xlsx, self._collect())
        except PermissionError:
            messagebox.showerror("Таблица мероприятий", "Файл открыт в другой программе (Excel). Закройте его "
                                 "и повторите.", parent=self.root)
            return
        except Exception as e:
            messagebox.showerror("Таблица мероприятий", f"Не удалось записать: {e}", parent=self.root)
            return
        self._read_measures()
        self._refresh()
        messagebox.showinfo("Таблица мероприятий", info, parent=self.root)

    # ------------------------------------------------------------------ выбор файлов
    def _ask(self, kind, multi, title):
        from tkinter import filedialog
        types = {"vec": VEC_TYPES, "ras": RAS_TYPES, "xls": XLS_TYPES}[kind]
        if multi:
            return list(filedialog.askopenfilenames(parent=self.root, title=title, filetypes=types))
        p = filedialog.askopenfilename(parent=self.root, title=title, filetypes=types)
        return [p] if p else []

    def _pick_xlsx(self):
        ps = self._ask("xls", False, "Таблица мероприятий (Excel)")
        if ps:
            self.cfg.measures_xlsx = ps[0]
            self._read_measures()
            self._refresh()

    def _pick_oozt(self):
        ps = self._ask("vec", False, "Слой территории: " + UNIT_MODES[self.v_mode.get()])
        if not ps:
            return
        self.cfg.oozt_path = ps[0]
        self.cfg.oozt_layer = None
        if not self.v_out.get():
            self.v_out.set(str(Path(ps[0]).parent / "oozt_recs_out"))
        self._setup_layer_combo()
        self._read_oozt_fields()

    def _pick_layer(self, spec):
        ps = self._ask(spec["kind"], spec["multi"], spec["title"] + (" — можно выбрать несколько" if spec["multi"] else ""))
        if not ps:
            return
        if spec["multi"]:
            cur = list(getattr(self.cfg, spec["key"]) or [])
            setattr(self.cfg, spec["key"], cur + [p for p in ps if p not in cur])
        else:
            setattr(self.cfg, spec["key"], ps[0])
        self._refresh()

    def _clear_layer(self, key):
        setattr(self.cfg, key, [] if isinstance(getattr(self.cfg, key), list) else None)
        self._refresh()

    def _pick_out(self):
        from tkinter import filedialog
        p = filedialog.askdirectory(parent=self.root, title="Папка для результатов")
        if p:
            self.v_out.set(p)

    # ------------------------------------------------------------------ чтение входов
    def _read_measures(self):
        try:
            self.measures = read_measures(self.cfg.measures_xlsx)
            self.v_xlsx.set(self.cfg.measures_xlsx)
            n_on = sum(m["on"] for m in self.measures)
            self.l_xlsx.configure(foreground=self.C_OK,
                                  text=f"✔ {len(self.measures)} мероприятий, включено (Вкл/выкл = 1): {n_on}")
        except Exception as e:
            self.measures = []
            self.v_xlsx.set(self.cfg.measures_xlsx)
            self.l_xlsx.configure(foreground=self.C_WARN, text=f"! таблица не прочитана: {e}")

    def _setup_layer_combo(self):
        self.v_oozt.set(self.cfg.oozt_path)
        p = self.cfg.oozt_path or ""
        if p.lower().endswith(".gpkg") and Path(p).exists():
            layers = [l[0] for l in pyogrio.list_layers(p)]
            self.cb_layer.configure(values=layers, state="readonly")
            self.v_layer.set(self.cfg.oozt_layer if self.cfg.oozt_layer in layers else layers[0])
        else:
            self.cb_layer.configure(values=[], state="disabled")
            self.v_layer.set("")

    def _read_oozt_fields(self):
        self.attr_df = None
        cols = []
        if self.cfg.oozt_path and Path(self.cfg.oozt_path).exists():
            try:
                cols, self.attr_df = _vector_fields(self.cfg.oozt_path, self.v_layer.get() or None)
            except Exception as e:
                self.l_types.configure(foreground=self.C_WARN, text=f"! слой не прочитан: {e}")
        self.columns = cols
        self._build_fields(guess=True)
        self._refresh()

    def _build_fields(self, guess=False):
        ttk = self.ttk
        for w in self.fields_frame.winfo_children():
            w.destroy()
        mode = self.v_mode.get()
        cols = getattr(self, "columns", [])
        prev = {k: v.get() for k, v in self.field_vars.items()}
        self.field_vars = {}
        ttk.Label(self.fields_frame, text="Поля слоя:", font=self.f_bold).grid(row=0, column=0, sticky="w")
        for i, (key, label, req, cands) in enumerate(FIELD_SPECS[mode], start=1):
            v = self.tk.StringVar()
            values = cols if req else [NONE_FIELD] + cols
            cur = (self.cfg.fields or {}).get(key) or prev.get(key)
            if cur not in cols and guess:
                cur = next((c for c in cands if c in cols), None) or next(
                    (c for c in cols if c.lower() in [x.lower() for x in cands]), None)
            v.set(cur if cur in cols else ("" if req else NONE_FIELD))
            ttk.Label(self.fields_frame, text=label + (" *" if req else "")).grid(row=i, column=0, sticky="w", pady=1)
            cb = ttk.Combobox(self.fields_frame, textvariable=v, values=values, state="readonly", width=26)
            cb.grid(row=i, column=1, sticky="w", padx=6, pady=1)
            cb.bind("<<ComboboxSelected>>", lambda e: self._refresh())
            self.field_vars[key] = v
        if mode == "eco":
            self.eco_frame.pack(fill="x", pady=(6, 0), before=self.l_types)
            self.cb_group.configure(values=[NONE_FIELD] + cols)
            if self.v_group.get() not in [NONE_FIELD] + cols:
                self.v_group.set(NONE_FIELD)
        else:
            self.eco_frame.pack_forget()

    def _on_mode(self):
        self.cfg.fields = {}
        self._build_fields(guess=True)
        self._refresh()

    # ------------------------------------------------------------------ сводка
    def _collect(self) -> Config:
        c = self.cfg
        c.unit_mode = self.v_mode.get()
        c.oozt_layer = self.v_layer.get() or None
        c.fields = {k: v.get() for k, v in self.field_vars.items() if v.get() and v.get() != NONE_FIELD}
        c.oozt_name = self.v_oozt_name.get().strip()
        g = self.v_group.get()
        c.report_group_field = g if (c.unit_mode == "eco" and g and g != NONE_FIELD) else None
        try:
            c.sheet_size_m = max(50.0, float(str(self.v_sheet.get()).replace(",", ".")))
        except ValueError:
            c.sheet_size_m = 1000.0
        c.out_dir = self.v_out.get().strip()
        c.report_kv = [x.strip() for x in re.split(r"[,;\s]+", self.v_kv.get()) if x.strip()]
        c.report_kv = [int(x) if x.isdigit() else x for x in c.report_kv]
        c.basemap = bool(self.v_basemap.get())
        c.split_multipart = bool(self.v_split.get())
        try:
            c.min_share_pct = max(0.0, float(self.v_minshare.get().replace(",", ".") or 0))
        except ValueError:
            c.min_share_pct = 25.0
        c.report_formats = [k for k, v in self.v_formats.items() if v.get()]
        return c

    def _refresh(self):
        cfg = self._collect()
        L = unit_labels(cfg)
        self.l_kv.configure(text=("Кварталы для отчёта:" if cfg.unit_mode != "eco" else
                                  "Листы для отчёта:" if not cfg.report_group_field else f"{L['group']}: для отчёта"))
        # типы полигонов
        self.type_counts = None
        tf = cfg.fields.get("TYPE")
        if self.attr_df is not None and tf in self.attr_df.columns:
            vc = self.attr_df[tf].map(_remojibake).astype(str).str.strip().value_counts()
            self.type_counts = vc.to_dict()
            known = (TYPE_FOREST + TYPE_MEADOW + TYPE_WATER + TYPE_NEARWATER)
            parts, unk = [], []
            for t, n in vc.items():
                (parts if any(k in t.lower() for k in known) else unk).append(f"{t} — {n}")
            txt = f"Полигонов: {len(self.attr_df)}. Типы, распознаваемые правилами: " + (", ".join(parts) or "нет")
            if unk:
                txt += ". Не используются правилами по типу: " + ", ".join(unk[:8])
            idf = cfg.fields.get("ID")
            if cfg.unit_mode == "eco" and idf in self.attr_df.columns:
                dup = int(self.attr_df[idf].duplicated().sum())
                txt += (f". ID уникальны." if dup == 0 else f". ! В поле ID {dup} повторов — ID должен быть уникальным.")
            self.l_types.configure(foreground=self.C_WARN if (not parts or "!" in txt) else self.C_MUTED, text=txt)
        elif self.attr_df is not None:
            self.l_types.configure(foreground=self.C_WARN, text="! выберите поле типа полигона")
        # тематические слои
        need_by = {k: [] for k in self.layer_rows}
        statuses = []
        for m in self.measures:
            s = measure_status(m, cfg, self.type_counts)
            statuses.append(s)
            if s["rule"]:
                for d in RULES[s["rule"]]["data"] + RULES[s["rule"]]["opt"]:
                    for k in DATA_KEYS[d]:
                        if k in need_by:
                            need_by[k].append((m.get("code") or "?", m["on"]))
        self._sp_refresh_status(statuses)
        for key, row in self.layer_rows.items():
            val = getattr(cfg, key)
            files = val if isinstance(val, list) else ([val] if val else [])
            codes_on = sorted(dict.fromkeys(c for c, on in need_by[key] if on), key=_code_key)
            codes_off = sorted(dict.fromkeys(c for c, on in need_by[key] if not on and c not in codes_on), key=_code_key)
            if not self.measures:
                utxt = "Нужен для мероприятий: выберите таблицу мероприятий"
            else:
                utxt = "Нужен для включённых: " + (", ".join(codes_on) if codes_on else "—")
                if codes_off:
                    utxt += ";  выключенных: " + ", ".join(codes_off)
            row["used"].configure(text=utxt, foreground="#202124" if codes_on else self.C_MUTED)
            if files:
                names = [Path(f.split("|")[0]).name for f in files]
                row["files"].configure(foreground=self.C_OK,
                                       text="Выбрано" + (f" ({len(files)})" if len(files) > 1 else "") + ": " +
                                       ", ".join(names[:5]) + (" …" if len(names) > 5 else ""))
                row["status"].configure(text="✔", foreground=self.C_OK)
            else:
                row["files"].configure(foreground=self.C_WARN if codes_on else self.C_OFF, text="Не выбран")
                if codes_on:
                    row["status"].configure(text="!", foreground=self.C_WARN)
                else:
                    row["status"].configure(text="○", foreground=self.C_OFF)
        # таблица мероприятий
        self.tree.delete(*self.tree.get_children())
        icon = {"ok": "✔", "missing": "!", "off": "—", "na": "—"}
        for i, (m, s) in enumerate(zip(self.measures, statuses)):
            self.tree.insert("", "end", iid=str(i), tags=(s["state"],),
                             values=(m.get("code", ""), m["name"], f"{icon[s['state']]} {s['text']}"))
        self._statuses = statuses
        if self.measures:
            n_on = sum(m["on"] for m in self.measures)
            n_ok = sum(s["state"] == "ok" for s in statuses)
            n_miss = sum(s["state"] == "missing" for s in statuses)
            self.l_sum.configure(text=f"Будет рассчитано {n_ok} из {n_on} включённых мероприятий"
                                      + (f"; не хватает данных: {n_miss}" if n_miss else ""),
                                 foreground=self.C_OK if n_miss == 0 else self.C_WARN)

    def _on_select(self, _e=None):
        sel = self.tree.selection()
        if not sel:
            return
        i = int(sel[0])
        m, s = self.measures[i], self._statuses[i]
        spec = s["spec"]
        lines = [f"{m.get('code', '')} {m['name']}".strip(),
                 f"Пункт Приказа: {m.get('prikaz') or '—'};  тип экосистемы: {m.get('eco') or '—'};  "
                 f"Вкл/выкл: {'1' if m['on'] else '0'}",
                 f"Метод выявления (Excel): {m.get('method') or '—'}", ""]
        if spec["rule"]:
            how = ("из столбца «Правило назначения»" if spec["origin"] == "таблица" else
                   "определено по тексту «Метода выявления» (столбец «Правило назначения» пуст)")
            lines.append(f"Правило: {RULES[spec['rule']]['title']} — {how}")
            lines.append("Параметры (в скобках — откуда взято значение):")
            lines += ["  • " + x for x in spec_params_lines(spec)]
            lines += ["", "Как рассчитывается: " + describe_spec(spec, full=True), ""]
            lines.append(f"Нужные данные: {s['need']}")
            if RULES[spec["rule"]]["opt"]:
                lines.append("Дополнительно используется: " +
                             ", ".join(LAYER_TITLE[k] for d in RULES[spec["rule"]]["opt"] for k in DATA_KEYS[d]))
            if "species" in RULES[spec["rule"]]["data"]:
                ls = matching_species_layers(self.cfg, spec["p"],
                                             ["points"] if spec["rule"] == "species_density" else None)
                lines.append("Подходящие слои видов: " + ("; ".join(species_layer_label(l) for l in ls) or "нет"))
        else:
            lines.append("Правило: не задано — " + spec.get("reason", ""))
        if spec.get("errors"):
            lines.append("Ошибки в столбцах правил: " + "; ".join(spec["errors"]))
        lines.append(f"Статус: {s['text']}")
        self.detail.configure(state="normal")
        self.detail.delete("1.0", "end")
        self.detail.insert("1.0", "\n".join(lines))
        self.detail.configure(state="disabled")

    # ------------------------------------------------------------------ конфигурация
    def _load_from_cfg(self, cfg: Config):
        self.cfg = cfg
        self.v_mode.set(cfg.unit_mode if cfg.unit_mode in UNIT_MODES else "kv_vyd")
        self.v_oozt_name.set(cfg.oozt_name or "")
        self.v_group.set(cfg.report_group_field or NONE_FIELD)
        self.v_sheet.set(str(int(cfg.sheet_size_m)) if cfg.sheet_size_m else "1000")
        self.v_out.set(cfg.out_dir or "")
        self.v_kv.set(", ".join(map(str, cfg.report_kv or [])))
        self.v_basemap.set(bool(cfg.basemap))
        self.v_split.set(bool(getattr(cfg, "split_multipart", True)))
        self.v_minshare.set(f"{float(getattr(cfg, 'min_share_pct', 25.0)):g}")
        for k, v in self.v_formats.items():
            v.set(k in (cfg.report_formats or []))
        self.field_vars = {}
        self._sp_render()
        if cfg.measures_xlsx:
            self._read_measures()
        self._setup_layer_combo()
        if cfg.oozt_path and Path(cfg.oozt_path).exists():
            try:
                self.columns, self.attr_df = _vector_fields(cfg.oozt_path, self.v_layer.get() or None)
            except Exception:
                self.columns, self.attr_df = [], None
        self._build_fields(guess=True)
        self._refresh()

    def _load_cfg_file(self):
        from tkinter import filedialog, messagebox
        p = filedialog.askopenfilename(parent=self.root, title="Конфигурация (json)", filetypes=JSON_TYPES)
        if not p:
            return
        try:
            self._load_from_cfg(Config.load(p))
        except Exception as e:
            messagebox.showerror("Конфигурация", f"Не удалось прочитать {p}:\n{e}", parent=self.root)

    def _save_cfg_file(self):
        from tkinter import filedialog
        p = filedialog.asksaveasfilename(parent=self.root, title="Сохранить конфигурацию", defaultextension=".json",
                                         filetypes=JSON_TYPES, initialfile="config_local.json")
        if p:
            self._collect().save(p)

    def _validate(self) -> list[str]:
        c = self._collect()
        err = []
        if not c.measures_xlsx or not self.measures:
            err.append("не выбрана или не прочитана таблица мероприятий")
        if not c.oozt_path:
            err.append("не выбран слой территории")
        for key, label, req, _ in FIELD_SPECS[c.unit_mode]:
            if req and not c.fields.get(key):
                err.append(f"не задано поле «{label}»")
        if c.unit_mode == "eco" and not c.fields.get("OOZT") and not c.oozt_name:
            err.append("укажите поле ООЗТ или название ООЗТ")
        if c.unit_mode == "eco" and self.attr_df is not None and c.fields.get("ID") in self.attr_df.columns:
            if self.attr_df[c.fields["ID"]].duplicated().any():
                err.append("в поле ID экосистемы есть повторы")
        if not c.out_dir:
            err.append("не указана папка результатов")
        return err

    def _run(self):
        from tkinter import messagebox
        err = self._validate()
        if err:
            messagebox.showwarning("Не хватает данных", "• " + "\n• ".join(err), parent=self.root)
            return
        miss = [f"{m.get('code', '')} {m['name']}: {s['text']}" for m, s in zip(self.measures, self._statuses)
                if s["state"] == "missing"]
        if miss and not messagebox.askyesno(
                "Часть мероприятий не будет рассчитана",
                "Для этих включённых мероприятий не хватает данных:\n\n" + "\n".join(miss[:15]) +
                ("\n…" if len(miss) > 15 else "") + "\n\nЗапустить расчёт без них?", parent=self.root):
            return
        self.result = self._collect()
        self.root.destroy()

    def _cancel(self):
        self.result = None
        self.root.destroy()

    def show(self) -> Config | None:
        self.root.mainloop()
        return self.result


def _code_key(code: str):
    """Сортировка кодов вида «III.5»: римская часть, затем номер."""
    rom = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6}
    mt = re.match(r"^([IVX]+)\.(\d+)", str(code))
    return (rom.get(mt.group(1), 99), int(mt.group(2))) if mt else (99, 0)


def gui_collect_config(cfg: Config | None = None) -> Config:
    res = SetupWindow(cfg).show()
    if res is None:
        print("Отменено пользователем")
        sys.exit(0)
    return res


# =====================================================================================
#                                 ЗАГРУЗКА ДАННЫХ
# =====================================================================================
def guess_crs_from_bounds(b) -> int | None:
    """Эвристика CRS по координатам, если она не задана или задана неверно."""
    minx, miny, maxx, maxy = b
    if -180 <= minx <= 180 and -90 <= miny <= 90 and maxx <= 180 and maxy <= 90:
        return 4326
    if 1e5 < minx < 9e5 and 5.5e6 < miny < 7e6:
        return WORK_EPSG          # UTM 37N (Москва ≈ 350–470 тыс. E, 6.1–6.3 млн N)
    if 3e6 < minx < 5e6 and 7e6 < miny < 8e6:
        return 3857               # Web Mercator (Москва ≈ 4.1–4.3 млн E, 7.4–7.6 млн N)
    return None


def fix_crs(obj_crs, bounds, what: str):
    """Возвращает CRS, исправленную по границам (частая ошибка: tif «в UTM», а координаты Web Mercator)."""
    guess = guess_crs_from_bounds(bounds)
    declared = None
    try:
        declared = obj_crs.to_epsg() if obj_crs else None
    except Exception:
        pass
    if declared is None:
        if guess is None:
            log.warning("%s: CRS не задана и не распознана по границам — считаем EPSG:%s", what, WORK_EPSG)
            return rasterio.crs.CRS.from_epsg(WORK_EPSG)
        log.warning("%s: CRS не задана, по границам принята EPSG:%s", what, guess)
        return rasterio.crs.CRS.from_epsg(guess)
    if guess is not None and guess != declared and declared in (4326, 3857, WORK_EPSG):
        log.warning("%s: объявлена EPSG:%s, но координаты соответствуют EPSG:%s — исправлено", what, declared, guess)
        return rasterio.crs.CRS.from_epsg(guess)
    return obj_crs


def read_vector(path: str, layer: str | None = None, bbox_wgs84=None) -> gpd.GeoDataFrame:
    """Чтение вектора любого формата с приведением к WORK_EPSG. bbox_wgs84 — (minx,miny,maxx,maxy) в EPSG:4326."""
    kw = {}
    if layer:
        kw["layer"] = layer
    # bbox в CRS файла
    src_crs = None
    try:
        info = pyogrio.read_info(path, **kw)
        src_crs = info.get("crs")
        n = info.get("features", -1)
    except Exception:
        n = -1
    if bbox_wgs84 is not None and n > 20000:
        try:
            bbs = gpd.GeoSeries([box(*bbox_wgs84)], crs=4326)
            bb = bbs.to_crs(src_crs).total_bounds if src_crs else bbs.total_bounds
            kw["bbox"] = tuple(bb)
        except Exception:
            pass
    try:
        g = gpd.read_file(path, **kw)
    except UnicodeDecodeError:
        g = gpd.read_file(path, encoding="cp1251", **kw)
    if g.empty:
        return g
    g = fix_text_encoding(g, path, kw)
    g = g[~g.geometry.isna() & ~g.geometry.is_empty].copy()
    crs = fix_crs(g.crs, g.total_bounds, Path(path).name)
    g = g.set_crs(crs, allow_override=True)
    if g.crs.to_epsg() != WORK_EPSG:
        g = g.to_crs(WORK_EPSG)
    g["__src"] = Path(path).stem if not layer else f"{Path(path).stem}:{layer}"
    return g


_CYR_RE = re.compile(r"[\u0400-\u04FF]")
_MOJI_RE = re.compile(r"[\u00C0-\u00FF][\u0080-\u00BF]")   # UTF-8, прочитанный как latin-1


def _remojibake(s):
    """«Р›РµСЃ» / «Ð›ÐµÑ» → «Лес»: пробуем обратить неверное декодирование UTF-8 как cp1251/latin1."""
    if not isinstance(s, str) or _CYR_RE.search(s) and not re.search(r"[РС][\u0080-\u04FF]", s):
        return s
    for enc in ("cp1251", "latin1"):
        try:
            r = s.encode(enc).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
        if _CYR_RE.search(r) or r.isascii():
            return r
    return s


def fix_text_encoding(g: gpd.GeoDataFrame, path: str, kw: dict) -> gpd.GeoDataFrame:
    """Shapefile без .cpg часто читается с «кракозябрами». Пробуем latin1→utf8, иначе перечитываем как cp1251."""
    # имена полей (кириллица в shp без .cpg)
    ren = {c: _remojibake(c) for c in g.columns if isinstance(c, str) and c != g.geometry.name}
    ren = {k: v for k, v in ren.items() if v != k and v not in g.columns}
    if ren:
        g = g.rename(columns=ren)
        log.info("%s: имена полей перекодированы: %s", Path(path).name, ", ".join(ren.values()))
    cols = [c for c in g.columns if c != g.geometry.name and (g[c].dtype == object or pd.api.types.is_string_dtype(g[c]))]
    if not cols:
        return g
    sample = " ".join(str(v) for c in cols for v in g[c].dropna().astype(str).head(30))
    if not sample.strip():
        return g
    moji_cp = re.search(r"[РС][\u0402-\u045F\u2018-\u2122]", sample)   # UTF-8, прочитанный как cp1251 («Р›РµСЃ»)
    if _CYR_RE.search(sample) and not moji_cp:
        return g                                  # уже нормальная кириллица
    if _MOJI_RE.search(sample) or moji_cp:
        for c in cols:
            g[c] = g[c].map(_remojibake)
        log.info("%s: атрибуты перекодированы latin1→utf-8", Path(path).name)
        return g
    if re.search(r"[\u0080-\u00FF]", sample):      # похоже на cp1251, прочитанный как latin1
        try:
            g2 = gpd.read_file(path, encoding="cp1251", **kw)
            log.info("%s: атрибуты перечитаны в cp1251", Path(path).name)
            return g2
        except Exception:
            pass
    return g


def read_many(paths: Iterable[str], bbox_wgs84=None, clip=None) -> gpd.GeoDataFrame | None:
    parts = []
    for p in paths:
        layer = None
        if "|" in p:                      # «file.gpkg|layer»
            p, layer = p.split("|", 1)
        try:
            if p.lower().endswith(".gpkg") and layer is None:
                for ly in [l[0] for l in pyogrio.list_layers(p)]:
                    parts.append(read_vector(p, ly, bbox_wgs84))
            else:
                parts.append(read_vector(p, layer, bbox_wgs84))
            log.info("Загружен слой %s (%d объектов)", Path(p).name, len(parts[-1]))
        except Exception as e:
            log.error("Не удалось прочитать %s: %s", p, e)
    parts = [x for x in parts if x is not None and not x.empty]
    if not parts:
        return None
    g = gpd.GeoDataFrame(pd.concat(parts, ignore_index=True), crs=WORK_EPSG)
    if clip is not None:
        g = g[g.intersects(clip)].copy()
    return g


LATIN_RE = re.compile(r"^[A-Z][a-z]+(\s+(sp\.?|[a-z\-]+))?(\s+[a-z\-]+)?$")


def detect_species_column(g: gpd.GeoDataFrame) -> str | None:
    """Ищет колонку с латинскими названиями видов: сначала по имени, потом по содержимому."""
    def latin_share(c):
        s = g[c].dropna().astype(str).str.strip()
        s = s[s != ""].head(300)
        return s.str.match(LATIN_RE).mean() if not s.empty else 0.0

    prefer = ["species", "Species", "SPECIES", "taxon", "Taxon", "latin", "Latin", "lat_name", "Latin_name",
              "Вид_лат", "Вид", "вид"]
    for c in prefer:
        if c in g.columns and latin_share(c) >= 0.5:     # «Вид» может быть и русским названием — проверяем содержимое
            return c
    best, best_score = None, 0.0
    for c in g.columns:
        if c == g.geometry.name or not (g[c].dtype == object or pd.api.types.is_string_dtype(g[c])):
            continue
        score = latin_share(c)
        if score > best_score:
            best, best_score = c, score
    return best if best_score >= 0.5 else None


def species_from_filename(path: str) -> str:
    stem = Path(path).stem
    stem = re.sub(r"[_\-]+", " ", stem).strip()
    return stem[0].upper() + stem[1:] if stem else Path(path).stem


def read_raster_clipped(path: str, clip_bounds_work, resolution: float | None = None, what="raster"):
    """Читает растр, исправляет CRS, перепроецирует в WORK_EPSG, обрезает по bounds (в WORK_EPSG).
    Возвращает (array, transform, nodata)."""
    with rasterio.open(path) as src:
        crs = fix_crs(src.crs, src.bounds, Path(path).name)
        src_nodata = src.nodata
        if src_nodata is None and path.lower().endswith(".asc"):
            src_nodata = -9999.0
        minx, miny, maxx, maxy = clip_bounds_work
        if resolution is None:
            # оценка разрешения источника в метрах
            if crs.to_epsg() == 4326:
                resolution = float(abs(src.res[0])) * 111320 * np.cos(np.deg2rad((miny + maxy) / 2 / 111320))
                resolution = max(1.0, round(resolution / 2.5, 0) * 2.5) if resolution > 0 else CELL
            else:
                resolution = float(abs(src.res[0]))
        width = int(np.ceil((maxx - minx) / resolution))
        height = int(np.ceil((maxy - miny) / resolution))
        if width <= 0 or height <= 0:
            return None, None, None
        dst_transform = from_origin(minx, maxy, resolution, resolution)
        dst = np.full((height, width), np.nan, dtype="float32")
        # reproject поддерживает пересчёт из исправленной CRS (src_crs передаём явно)
        reproject(
            source=rasterio.band(src, 1), destination=dst,
            src_transform=src.transform, src_crs=crs, src_nodata=src_nodata,
            dst_transform=dst_transform, dst_crs=rasterio.crs.CRS.from_epsg(WORK_EPSG),
            dst_nodata=np.nan, resampling=Resampling.bilinear,
        )
        if src_nodata is not None:
            dst[dst == src_nodata] = np.nan
        return dst, dst_transform, np.nan


def polygonize_mask(mask: np.ndarray, transform, min_area_m2: float = 0.0) -> gpd.GeoDataFrame:
    """Булева маска → полигоны (WORK_EPSG)."""
    m = mask.astype("uint8")
    geoms = [shape(geom) for geom, val in rfeatures.shapes(m, mask=m.astype(bool), transform=transform) if val == 1]
    g = gpd.GeoDataFrame(geometry=geoms, crs=WORK_EPSG)
    if min_area_m2 > 0 and not g.empty:
        g = g[g.area >= min_area_m2].copy()
    return g


# =====================================================================================
#                              РАСТРЫ ПЛОТНОСТИ И УКЛОНА
# =====================================================================================
def make_grid(bounds, cell=CELL):
    minx, miny, maxx, maxy = bounds
    width = int(np.ceil((maxx - minx) / cell))
    height = int(np.ceil((maxy - miny) / cell))
    return from_origin(minx, maxy, cell, cell), (height, width)


def circular_kernel(radius_m: float, cell: float) -> np.ndarray:
    r = int(np.ceil(radius_m / cell))
    y, x = np.ogrid[-r:r + 1, -r:r + 1]
    return ((x * x + y * y) <= r * r).astype("float32")


def line_density(lines: gpd.GeoDataFrame, bounds, radius_m: float, cell=CELL):
    """Плотность линий (м/га) в круговом скользящем окне radius_m. Линии дискретизируются с шагом cell/2."""
    transform, shp = make_grid(bounds, cell)
    acc = np.zeros(shp, dtype="float32")
    step = cell / 2
    minx, miny, maxx, maxy = bounds
    for geom in lines.geometry:
        for part in getattr(geom, "geoms", [geom]):
            L = part.length
            if L == 0:
                continue
            n = max(2, int(L / step) + 1)
            d = np.linspace(0, L, n)
            pts = np.array([part.interpolate(t).coords[0] for t in d])
            seg = L / (n - 1)
            cols = ((pts[:, 0] - minx) / cell).astype(int)
            rows = ((maxy - pts[:, 1]) / cell).astype(int)
            ok = (cols >= 0) & (cols < shp[1]) & (rows >= 0) & (rows < shp[0])
            np.add.at(acc, (rows[ok], cols[ok]), seg)
    k = circular_kernel(radius_m, cell)
    total = ndimage.convolve(acc, k, mode="constant", cval=0.0)
    area_ha = k.sum() * cell * cell / 1e4
    return total / area_ha, transform            # м/га


def point_density(points: gpd.GeoDataFrame, bounds, radius_m: float, cell=CELL, weights=None):
    """Плотность точек (шт/га) в круговом окне radius_m."""
    transform, shp = make_grid(bounds, cell)
    acc = np.zeros(shp, dtype="float32")
    minx, miny, maxx, maxy = bounds
    xs = points.geometry.x.values
    ys = points.geometry.y.values
    w = np.ones(len(points), dtype="float32") if weights is None else np.asarray(weights, dtype="float32")
    cols = ((xs - minx) / cell).astype(int)
    rows = ((maxy - ys) / cell).astype(int)
    ok = (cols >= 0) & (cols < shp[1]) & (rows >= 0) & (rows < shp[0])
    np.add.at(acc, (rows[ok], cols[ok]), w[ok])
    k = circular_kernel(radius_m, cell)
    total = ndimage.convolve(acc, k, mode="constant", cval=0.0)
    area_ha = k.sum() * cell * cell / 1e4
    return total / area_ha, transform


def slope_percent(dem: np.ndarray, cell: float) -> np.ndarray:
    dem_f = dem.astype("float64")
    dem_f = np.where(np.isnan(dem_f), np.nanmedian(dem_f), dem_f)
    dem_f = ndimage.uniform_filter(dem_f, size=3, mode="nearest")   # лёгкое сглаживание шума
    dzdy, dzdx = np.gradient(dem_f, cell)
    slope = np.sqrt(dzdx ** 2 + dzdy ** 2) * 100.0
    slope[np.isnan(dem)] = np.nan
    return slope.astype("float32")


def upper_quantile_mask(density: np.ndarray, inside: np.ndarray, q=QUANTILE):
    """Маска ячеек ≥ q-квантиля, рассчитанного по ячейкам внутри ООЗТ с ненулевой плотностью."""
    vals = density[inside & (density > 0)]
    if vals.size == 0:
        return np.zeros_like(density, dtype=bool), np.nan
    thr = float(np.quantile(vals, q))
    return inside & (density >= thr) & (density > 0), thr


# =====================================================================================
#                                 КОНТЕКСТ ДАННЫХ
# =====================================================================================
def _nat_key(v):
    """Естественный порядок: числа — по значению (1110 < 5110 < 44504), текст — по фрагментам
    («кв2» < «кв10», «5-12» < «5-110»)."""
    s_ = "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip()
    try:
        return (0, float(s_.replace(",", ".")), ())
    except ValueError:
        pass
    parts = tuple((0, int(p), "") if p.isdigit() else (1, 0, p.lower()) for p in re.split(r"(\d+)", s_) if p)
    return (1, 0.0, parts)


def _nat_rank(values) -> np.ndarray:
    """Ранги значений в естественном порядке (для сортировки смешанных / текстовых полей)."""
    vals = ["" if pd.isna(v) else str(v).strip() for v in values]
    order = {v: i for i, v in enumerate(sorted(set(vals), key=_nat_key))}
    return np.array([order[v] for v in vals], dtype=np.int64)


def _nat_sort_key(*cols):
    """key= для DataFrame.sort_values: указанные колонки — в естественном порядке."""
    def key(sr: pd.Series):
        return pd.Series(_nat_rank(sr.values), index=sr.index) if sr.name in cols else sr
    return key


def _maybe_int(sr: pd.Series) -> pd.Series:
    """Номера кварталов/выделов/ID → Int64, если все значения целые; иначе — текст как есть."""
    try:
        num = pd.to_numeric(sr)
        if (num.dropna() % 1 == 0).all():
            return num.astype("Int64")
    except (ValueError, TypeError):
        pass
    return sr.astype(str).str.strip()


MIN_PART_HA = 0.01      # контур многочастного объекта меньше этой площади — не отдельная единица


def _split_multipart(g: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    """Многочастный полигон → отдельные контуры. Мелкие части (< MIN_PART_HA) остаются при крупнейшем контуре."""
    from shapely.geometry import MultiPolygon
    rows = []
    for i, (idx, r) in enumerate(g.iterrows()):
        geom = r.geometry
        parts = list(geom.geoms) if geom.geom_type == "MultiPolygon" else [geom]
        parts.sort(key=lambda x: -x.area)
        big = [p for p in parts if p.area / 1e4 >= MIN_PART_HA] or parts[:1]
        small = [p for p in parts if p not in big]
        if small:
            # мелкие части — к ближайшему крупному контуру
            groups = [[b] for b in big]
            for sp in small:
                k = min(range(len(big)), key=lambda j: big[j].distance(sp))
                groups[k].append(sp)
            big = [grp[0] if len(grp) == 1 else MultiPolygon(grp) for grp in groups]
        for k, p in enumerate(big, 1):
            rr = r.copy()
            rr["geometry"] = p
            rr["_part"], rr["_nparts"], rr["_src_row"] = k, len(big), i
            rows.append(rr)
    return gpd.GeoDataFrame(rows, geometry="geometry", crs=g.crs).reset_index(drop=True)


def _part_ids(ids: pd.Series, g: gpd.GeoDataFrame) -> pd.Series:
    """ID контура: исходный ID, для частей многочастного объекта — «ID-1», «ID-2», …"""
    multi = g["_nparts"] > 1
    out = ids.copy().astype(object)
    out[multi] = ids[multi].astype(str) + "-" + g.loc[multi, "_part"].astype(str)
    return out


class Data:
    """Все загруженные слои в WORK_EPSG + служебные объекты."""

    def __init__(self, cfg: Config, oozt_only: bool = False, pad_m: float = 0.0):
        global _CUR_CFG
        self.cfg = cfg
        _CUR_CFG = cfg
        self.oozt = self._load_oozt()
        self.union = unary_union(self.oozt.geometry.values)
        self.bounds = tuple(self.oozt.total_bounds)
        pad = max(TRAIL_WINDOW_M, CORVUS_WINDOW_M, RARE_BUFFER_M, SHORE_BUFFER_M, pad_m) + 2 * CELL
        b = self.bounds
        self.bounds_pad = (b[0] - pad, b[1] - pad, b[2] + pad, b[3] + pad)
        self.bbox_wgs84 = tuple(gpd.GeoSeries([box(*self.bounds_pad)], crs=WORK_EPSG).to_crs(4326).total_bounds)
        self.clip = box(*self.bounds_pad)
        self.trails = self.hydro_lines = self.hydro_polys = None
        self.sp_points = self.sp_polys = None     # находки видов: species, category, group, layer, __src
        self.sp_rasters: list[dict] = []          # модели SDM: path, category, group, species
        self.dem = None
        self._inside = None
        if oozt_only:                      # режим --report-only: остальные слои не нужны
            return

        self.trails = self._load_trails()
        self.hydro_lines = read_many(cfg.hydro_lines, self.bbox_wgs84, self.clip)
        self.hydro_polys = read_many(cfg.hydro_polys, self.bbox_wgs84, self.clip)
        self.sp_points, self.sp_polys, self.sp_rasters = self._load_species_layers(species_layers(cfg))
        self.dem = None
        if cfg.dem_path:
            try:
                arr, tr, _ = read_raster_clipped(cfg.dem_path, self.bounds_pad, resolution=CELL, what="ЦМР")
                self.dem = (arr, tr)
                log.info("ЦМР загружена: %s, ячейка %.0f м", Path(cfg.dem_path).name, CELL)
            except Exception as e:
                log.error("ЦМР не загружена: %s", e)
        self._inside = None

    # ---- ООЗТ ----
    def _load_oozt(self) -> gpd.GeoDataFrame:
        cfg = self.cfg
        g = read_vector(cfg.oozt_path, cfg.oozt_layer)
        g = g[g.geom_type.isin(["Polygon", "MultiPolygon"])].copy()
        g["_part"], g["_nparts"] = 1, 1
        if getattr(cfg, "split_multipart", True):
            g = _split_multipart(g)
        f = cfg.fields or {}
        cols = list(g.columns)

        def pick(key, candidates, required=True):
            if f.get(key) and f[key] in cols:
                return f[key]
            for c in candidates:
                if c in cols:
                    return c
            for c in cols:
                if c.lower() in [x.lower() for x in candidates]:
                    return c
            if required:
                raise KeyError(f"В слое территории не найдено поле {key} (искал {candidates}); "
                               f"задайте его в окне выбора или в cfg.fields['{key}']")
            return None

        spec = {k: (req, cands) for k, _, req, cands in FIELD_SPECS[cfg.unit_mode]}
        if cfg.unit_mode == "eco":
            self.F_ID = pick("ID", spec["ID"][1])
            self.F_TYPE = pick("TYPE", spec["TYPE"][1])
            self.F_NAME = pick("NAME", spec["NAME"][1], required=False)
            self.F_OOZT = pick("OOZT", spec["OOZT"][1], required=False)
            cfg.fields = {k: v for k, v in {"ID": self.F_ID, "TYPE": self.F_TYPE, "NAME": self.F_NAME,
                                            "OOZT": self.F_OOZT}.items() if v}
            _dup = g[[self.F_ID, "_part"]].duplicated()
            if _dup.any():
                dup = g.loc[_dup, self.F_ID].astype(str).unique()[:10]
                raise ValueError(f"Поле ID «{self.F_ID}» не уникально: повторы {', '.join(dup)}")
            gf = cfg.report_group_field
            if gf and gf not in cols:
                raise KeyError(f"Поле группировки «{gf}» не найдено в слое территории")
            grp_vals = g[gf].astype(str).str.strip().values if gf else None
            g = g.drop(columns=[c for c in ("OOZT", "KV", "VYD", "Type", "Name", "KV_SORT")
                                if c in cols and c not in (self.F_ID, self.F_TYPE, self.F_NAME, self.F_OOZT)])
            if self.F_NAME:
                g["Name"] = g[self.F_NAME]
            if not self.F_OOZT:
                g["OOZT"] = cfg.oozt_name or Path(cfg.oozt_path).stem
            else:
                g["OOZT"] = g[self.F_OOZT].astype(str).str.strip()
            g[self.F_ID] = _part_ids(g[self.F_ID], g)
            g["VYD"] = g[self.F_ID]
            g["Type"] = g[self.F_TYPE]
            if not self.F_NAME:
                g["Name"] = g["Type"]
            g["KV"], g["KV_SORT"] = self._report_groups(g, grp_vals)
        else:
            self.F_OOZT = pick("OOZT", spec["OOZT"][1])
            self.F_KV = pick("KV", spec["KV"][1])
            self.F_VYD = pick("VYD", spec["VYD"][1])
            self.F_TYPE = pick("TYPE", spec["TYPE"][1])
            self.F_NAME = pick("NAME", spec["NAME"][1], required=False)
            cfg.fields = {k: v for k, v in {"OOZT": self.F_OOZT, "KV": self.F_KV, "VYD": self.F_VYD,
                                            "TYPE": self.F_TYPE, "NAME": self.F_NAME}.items() if v}
            g = g.rename(columns={self.F_OOZT: "OOZT", self.F_KV: "KV", self.F_VYD: "VYD",
                                  self.F_TYPE: "Type", **({self.F_NAME: "Name"} if self.F_NAME else {})})
            if not self.F_NAME:
                g["Name"] = g["Type"]
            g["VYD"] = _part_ids(g["VYD"], g)
        n_split = int((g["_nparts"] > 1).sum())
        if n_split:
            log.info("Многочастные объекты разделены на отдельные контуры: %d объектов → %d контуров "
                     "(ID-1, ID-2, …; контуры < %.2f га присоединены к соседнему крупному контуру)",
                     g.loc[g["_nparts"] > 1, "_src_row"].nunique(), n_split, MIN_PART_HA)
        g = g.drop(columns=["_part", "_nparts", "_src_row"], errors="ignore")
        for c in ("KV", "VYD"):
            g[c] = _maybe_int(g[c])
        if cfg.unit_mode != "eco":
            g["KV_SORT"] = _nat_rank(g["KV"])
        g["VYD_SORT"] = _nat_rank(g["VYD"])
        g["Type"] = g["Type"].astype(str).str.strip()
        g["Name"] = g["Name"].astype(str).str.strip()
        g["uid"] = g["OOZT"].astype(str) + "|" + g["KV"].astype(str) + "|" + g["VYD"].astype(str)
        g["area_ha"] = g.geometry.area / 1e4
        g = g.reset_index(drop=True)
        L = unit_labels(cfg)
        log.info("Слой территории (%s): %d %s, %s: %d, ООЗТ: %s", UNIT_MODES[cfg.unit_mode], len(g),
                 L["units_gen"], L["groups_n"].lower(), g["KV"].nunique(),
                 ", ".join(map(str, g["OOZT"].unique())))
        return g

    def _report_groups(self, g: gpd.GeoDataFrame, grp_vals=None):
        """Режим eco: группа отчёта — значение поля группировки или лист регулярной сетки (A1, A2, …)."""
        cfg = self.cfg
        self.sheet_grid = None
        if grp_vals is not None:
            return grp_vals, _nat_rank(grp_vals)
        size = float(cfg.sheet_size_m or 1000.0)
        minx, miny, maxx, maxy = g.total_bounds
        pts = g.geometry.representative_point()
        col = np.floor((pts.x.values - minx) / size).astype(int)
        row = np.floor((maxy - pts.y.values) / size).astype(int)     # ряды сверху вниз (с севера)

        def letters(n):
            s = ""
            n += 1
            while n:
                n, r = divmod(n - 1, 26)
                s = chr(65 + r) + s
            return s
        labels = [f"{letters(r)}{c + 1}" for r, c in zip(row, col)]
        self.sheet_grid = dict(minx=minx, maxy=maxy, size=size)
        log.info("Листы сетки %.0f × %.0f м: %d листов с полигонами", size, size, len(set(labels)))
        return labels, row * 10000 + col

    # ---- тропы ----
    def _load_trails(self):
        g = read_many(self.cfg.trails, self.bbox_wgs84, self.clip)
        if g is None:
            return None
        col = next((c for c in ("HIGHWAY", "highway", "fclass", "FCLASS", "type") if c in g.columns), None)
        if col is not None:
            types = g[col].astype(str).str.lower()
            keep = types.isin(TRAIL_HIGHWAY_TYPES) | types.isna() | (types == "none") | (types == "nan")
            if keep.sum() > 0:
                log.info("Тропы: отфильтровано по %s — оставлено %d из %d (типы: %s)", col, keep.sum(), len(g),
                         ", ".join(sorted(types[keep].unique())[:12]))
                g = g[keep].copy()
        g = g[g.geom_type.isin(["LineString", "MultiLineString"])].copy()
        return g if not g.empty else None

    # ---- виды ----
    def _load_species_layers(self, layers):
        """Слои видов → общие таблицы точек и полигонов + список растров SDM (с категорией и группой)."""
        pts, polys, rasters = [], [], []
        for i, l in enumerate(layers):
            name = Path(str(l["path"]).split("|")[0]).name
            tag = f"{l['category']}, {l['group']}"
            if l["kind"] == "sdm":
                rasters.append(dict(l, species=species_from_filename(str(l["path"]).split("|")[0]), layer=i))
                log.info("Слой видов %s (%s): модель SDM, вид «%s»", name, tag, rasters[-1]["species"])
                continue
            if l["kind"] not in ("points", "polygons"):
                log.warning("Слой видов %s: не удалось определить тип данных — слой не используется", name)
                continue
            g = read_many([l["path"]], self.bbox_wgs84, self.clip)
            want = ["Point", "MultiPoint"] if l["kind"] == "points" else ["Polygon", "MultiPolygon"]
            if g is not None:
                g = g[g.geom_type.isin(want)]
                if l["kind"] == "points":
                    g = g.explode(index_parts=False)
            if g is None or g.empty:
                log.warning("Слой видов %s (%s): нет объектов в границах ООЗТ — слой не используется", name, tag)
                continue
            g = g.copy()
            col = l["field"] if l.get("field") in g.columns else detect_species_column(g)
            g["species"] = g[col].astype(str).str.strip() if col else g["__src"]
            g["category"], g["group"], g["layer"] = l["category"], l["group"], i
            if col:
                log.info("Слой видов %s (%s, %s): поле вида «%s», %d видов, %d объектов", name, tag,
                         SPECIES_KINDS[l["kind"]], col, g["species"].nunique(), len(g))
            else:
                log.warning("Слой видов %s (%s): поле с латинскими названиями не найдено — вид = имя файла",
                            name, tag)
            (pts if l["kind"] == "points" else polys).append(
                g[["species", "category", "group", "layer", "__src", "geometry"]])
        cat = lambda parts: gpd.GeoDataFrame(pd.concat(parts, ignore_index=True), crs=WORK_EPSG) if parts else None
        return cat(pts), cat(polys), rasters

    # ---- маска «внутри ООЗТ» на сетке CELL ----
    def inside_mask(self, transform, shp):
        m = rfeatures.rasterize([(self.union, 1)], out_shape=shp, transform=transform, fill=0, dtype="uint8")
        return m.astype(bool)

    def vydels_of_type(self, *keys) -> gpd.GeoDataFrame:
        m = self.oozt["Type"].map(lambda t: type_matches(t, keys)).values.astype(bool)
        return self.oozt[m]


# =====================================================================================
#                                    РЕЗУЛЬТАТЫ
# =====================================================================================
class Results:
    def __init__(self):
        self.rows: list[dict] = []          # длинная таблица
        self.evidence: list[gpd.GeoDataFrame] = []

    def add(self, vydel_row, measure: dict, basis: str, value=None):
        self.rows.append({
            "ООЗТ": vydel_row["OOZT"], "Квартал": vydel_row["KV"], "Выдел": vydel_row["VYD"],
            "_sort": vydel_row.get("KV_SORT", vydel_row["KV"]),
            "_vsort": vydel_row.get("VYD_SORT", 0),
            "Тип выдела": vydel_row["Type"], "Экосистема": vydel_row["Name"],
            "Тип мероприятий": measure.get("group", ""), "Мероприятие": measure["name"],
            "Тип экосистемы мероприятия": measure.get("eco", ""),
            "Пункт Приказа": measure.get("prikaz", ""), "Сезонность": measure.get("season", ""),
            "Основание": basis, "Значение": value,
        })

    def add_evidence(self, g: gpd.GeoDataFrame, measure_name: str, layer: str, label_col: str | None = None):
        if g is None or g.empty:
            return
        n = len(g)
        labels = g[label_col].astype(str).values if label_col and label_col in g.columns else [""] * n
        e = gpd.GeoDataFrame({"measure": [measure_name] * n, "layer": [layer] * n, "label": list(labels)},
                             geometry=g.geometry.values, crs=WORK_EPSG)
        self.evidence.append(e)

    def long_table(self) -> pd.DataFrame:
        df = pd.DataFrame(self.rows)
        if df.empty:
            return df
        # объединяем дубли (один выдел — одно мероприятие) с конкатенацией оснований
        key = ["ООЗТ", "Квартал", "Выдел", "Мероприятие"]
        agg = df.groupby(key, as_index=False, sort=False).agg({
            "Тип выдела": "first", "Экосистема": "first", "Тип мероприятий": "first", "_sort": "first", "_vsort": "first",
            "Тип экосистемы мероприятия": "first",
            "Пункт Приказа": "first", "Сезонность": "first",
            "Основание": lambda s: "; ".join(dict.fromkeys([x for x in s if x])),
            "Значение": "first",
        })
        agg = agg.sort_values(["ООЗТ", "_sort", "_vsort", "Тип мероприятий", "Мероприятие"])
        cols = ["ООЗТ", "Квартал", "Выдел", "Тип выдела", "Экосистема", "Тип мероприятий", "Мероприятие",
                "Тип экосистемы мероприятия", "Пункт Приказа", "Сезонность", "Основание", "Значение"]
        return agg[cols].reset_index(drop=True)

    def evidence_gdf(self) -> gpd.GeoDataFrame | None:
        if not self.evidence:
            return None
        return gpd.GeoDataFrame(pd.concat(self.evidence, ignore_index=True), crs=WORK_EPSG)


# =====================================================================================
#                                     ПРАВИЛА
# =====================================================================================
_CUR_CFG: Config | None = None     # текущая конфигурация (для подписей в журнале и основаниях)


def _units() -> str:
    return unit_labels(_CUR_CFG)["units_gen"] if _CUR_CFG else "выделов"


def _unit_word() -> str:
    return "полигона" if (_CUR_CFG and _CUR_CFG.unit_mode == "eco") else "выдела"


def _overlay_patches_to_vydels(data: Data, patches: gpd.GeoDataFrame, target: gpd.GeoDataFrame,
                               min_ha=MIN_PATCH_HA, share_pct: float | None = None):
    """Пересечение полигонов-пятен с выделами; возвращает DataFrame uid → площадь пятен (га), доля."""
    if patches is None or patches.empty or target.empty:
        return pd.DataFrame(columns=["uid", "patch_ha", "share"])
    pu = unary_union(patches.geometry.values)
    inter = target.geometry.intersection(pu)
    ha = inter.area / 1e4
    df = pd.DataFrame({"uid": target["uid"].values, "patch_ha": ha.values,
                       "share": (ha / target["area_ha"]).values})
    df = df[df["patch_ha"] >= min_ha]
    thr = _min_share(data) if share_pct is None else float(share_pct)
    if thr > 0 and len(df):
        n0 = len(df)
        df = df[df["share"] * 100 >= thr]
        log.info("  порог: участок ≥ %g%% площади полигона — оставлено %d из %d", thr, len(df), n0)
    return df


def _min_share(data) -> float:
    try:
        return float(getattr(data.cfg, "min_share_pct", 0) or 0)
    except (TypeError, ValueError):
        return 0.0


def _target(data: Data, p: dict) -> gpd.GeoDataFrame:
    t = p.get("types")
    return data.vydels_of_type(*t[0]) if t else data.oozt


def _qtag(q) -> str:
    return f"Q{int(round(float(q) * 100))}"


def _eco_label(p: dict) -> str:
    t = p.get("types")
    if not t:
        return "в любых экосистемах"
    if len(t[0]) == 1 and t[0][0] == "луг":
        return "в луговых экосистемах"
    return f"в полигонах типа «{t[1]}»"


def rule_type(data: Data, res: Results, measure: dict, p: dict):
    """Назначение по типу выдела."""
    keys, label = p["types"]
    sel = data.vydels_of_type(*keys)
    for _, r in sel.iterrows():
        res.add(r, measure, f"тип {_unit_word()} «{r['Type']}» ({label})")
    log.info("  → %d %s", len(sel), _units())


def _trail_hotspots(data: Data, window=TRAIL_WINDOW_M, q=QUANTILE):
    """Кэш: участки ≥ квантиля плотности троп в окне window (по ячейкам внутри ООЗТ)."""
    cache = data.__dict__.setdefault("_hot_cache", {})
    key = (float(window), float(q))
    if key in cache:
        return cache[key]
    if data.trails is None:
        cache[key] = (None, None, None)
        return cache[key]
    dens, tr = line_density(data.trails, data.bounds_pad, window, CELL)
    inside = data.inside_mask(tr, dens.shape)
    mask, thr = upper_quantile_mask(dens, inside, q)
    polys = polygonize_mask(mask, tr)
    log.info("  плотность троп (окно %.0f м): порог %s = %.0f м/га, %d пятен", window, _qtag(q), thr, len(polys))
    cache[key] = (polys, thr, (dens, tr))
    return cache[key]


def rule_trail_density_vydel(data: Data, res: Results, measure: dict, p: dict):
    """Плотность троп (м/га) по выделу; квантиль среди выделов всех ООЗТ заданного типа."""
    if data.trails is None:
        log.warning("  тропы не загружены — пропуск")
        return
    q = p["quantile"]
    target = _target(data, p)
    tu = unary_union(data.trails.geometry.values)
    dens = target.geometry.intersection(tu).length / target["area_ha"]
    dens = dens.replace([np.inf], np.nan).fillna(0)
    pos = dens[dens > 0]
    if pos.empty:
        return
    thr = float(np.quantile(pos, q))
    sel = target[dens >= thr]
    for (_, r), d in zip(sel.iterrows(), dens[dens >= thr]):
        res.add(r, measure, f"плотность троп {d:.0f} м/га ≥ {_qtag(q)} = {thr:.0f} м/га", round(d, 1))
    trails_in = None
    if not sel.empty:
        # Красным показываем не целые исходные линии, а только их фрагменты
        # внутри локальных пятен высокой плотности и выбранных выделов.
        hot_polys, _, _ = _trail_hotspots(data, TRAIL_WINDOW_M, q)
        if hot_polys is not None and not hot_polys.empty:
            clip_geom = unary_union(hot_polys.geometry.values).intersection(
                unary_union(sel.geometry.values)
            )
            trails_in = data.trails[data.trails.intersects(clip_geom)].copy()
            if not trails_in.empty:
                trails_in["geometry"] = trails_in.geometry.intersection(clip_geom)
                trails_in = trails_in.explode(index_parts=False)
                trails_in = trails_in[
                    (~trails_in.is_empty)
                    & trails_in.geom_type.isin(["LineString", "MultiLineString"])
                ]
    res.add_evidence(trails_in, measure["name"],
                     "тропы в зонах повышенной плотности")
    log.info("  → %d %s (порог %.0f м/га)", len(sel), _units(), thr)


def rule_trail_hotspot(data: Data, res: Results, measure: dict, p: dict):
    polys, thr, _ = _trail_hotspots(data, p["window"], p["quantile"])
    if polys is None:
        log.warning("  тропы не загружены — пропуск")
        return
    target = _target(data, p)
    hit = _overlay_patches_to_vydels(data, polys, target, p["min_ha"], spec_share(p))
    for _, h in hit.iterrows():
        r = data.oozt[data.oozt["uid"] == h["uid"]].iloc[0]
        res.add(r, measure, f"участки повышенной плотности троп (окно {p['window']:.0f} м, "
                            f"≥{_qtag(p['quantile'])}={thr:.0f} м/га): "
                            f"{h['patch_ha']:.2f} га ({100 * h['share']:.0f}% полигона)", round(h["patch_ha"], 2))
    if not hit.empty:
        sel = target[target["uid"].isin(hit["uid"])]
        pieces = polys[polys.intersects(unary_union(sel.geometry.values))]
        res.add_evidence(pieces, measure["name"],
                         "зоны повышенной плотности троп (верхний квартиль)" if abs(p["quantile"] - 0.75) < 1e-9
                         else f"зоны повышенной плотности троп ({_qtag(p['quantile'])})")
    log.info("  → %d %s", len(hit), _units())


def rule_shore_trails(data: Data, res: Results, measure: dict, p: dict):
    """Полоса вдоль водных объектов ∩ пятна повышенной плотности троп."""
    polys, thr, _ = _trail_hotspots(data, p["window"], p["quantile"])
    if polys is None or (data.hydro_lines is None and data.hydro_polys is None):
        log.warning("  нет троп или гидрологии — пропуск")
        return
    dist = p["dist"]
    parts = [g.geometry for g in (data.hydro_lines, data.hydro_polys) if g is not None]
    water = unary_union(pd.concat(parts).values)
    shore = water.buffer(dist).difference(water) if water.area > 0 else water.buffer(dist)
    hot_shore = polys[polys.intersects(shore)].copy()
    hot_shore["geometry"] = hot_shore.geometry.intersection(shore)
    hot_shore = hot_shore[~hot_shore.is_empty]
    hit = _overlay_patches_to_vydels(data, hot_shore, _target(data, p), p["min_ha"], spec_share(p))
    for _, h in hit.iterrows():
        r = data.oozt[data.oozt["uid"] == h["uid"]].iloc[0]
        res.add(r, measure, f"тропы повышенной плотности в {dist:.0f} м от водных объектов: "
                            f"{h['patch_ha']:.2f} га ({100 * h['share']:.0f}% полигона)", round(h["patch_ha"], 2))
    res.add_evidence(hot_shore, measure["name"],
                     f"тропы на расстоянии {dist:.0f} м от водных объектов ({_qtag(p['quantile'])})")
    # сами водные объекты рядом с найденными участками — как контекст на схеме
    if not hot_shore.empty:
        near = hot_shore.geometry.buffer(dist).union_all()
        if data.hydro_lines is not None:
            res.add_evidence(data.hydro_lines[data.hydro_lines.intersects(near)], measure["name"], "водные объекты (линии)")
        if data.hydro_polys is not None:
            res.add_evidence(data.hydro_polys[data.hydro_polys.intersects(near)], measure["name"], "водные объекты (полигоны)")
    log.info("  → %d %s", len(hit), _units())


# ---- виды ----------------------------------------------------------------------------
CATEGORY_LABEL = {"инвазивные": ("инвазивные виды", "инвазивных видов"),
                  "охраняемые": ("охраняемые виды", "охраняемых видов"),
                  "прочие": ("виды", "видов")}


def _sp_label(p: dict, gen=False) -> str:
    """Подпись выбранных видов для слоёв-обоснований: «инвазивные виды», «Corvus», «виды»."""
    if p.get("taxon"):
        return ", ".join(p["taxon"])
    cats = p.get("category") or []
    lab = CATEGORY_LABEL.get(cats[0], ("виды", "видов")) if len(cats) == 1 else ("виды", "видов")
    return lab[1] if gen else lab[0]


def _taxon_mask(species: pd.Series, taxa) -> np.ndarray:
    pat = "|".join(rf"\b{re.escape(t)}\b" for t in taxa)
    return species.astype(str).str.contains(pat, case=False, regex=True).values


def select_species(data: Data, p: dict):
    """Точки, полигоны и растры SDM, подходящие мероприятию (категория, группы, таксоны, тип данных)."""
    def filt(g):
        if g is None or g.empty:
            return None
        m = np.ones(len(g), dtype=bool)
        if p.get("category"):
            m &= g["category"].isin(p["category"]).values
        if p.get("groups"):
            m &= g["group"].map(lambda x: group_matches(x, p["groups"])).values.astype(bool)
        if p.get("taxon"):
            m &= _taxon_mask(g["species"], p["taxon"])
        g = g[m]
        return g if not g.empty else None
    pts = filt(data.sp_points) if _uses(p, "points") else None
    polys = filt(data.sp_polys) if _uses(p, "polygons") else None
    rasters = []
    if _uses(p, "sdm"):
        for r in data.sp_rasters:
            if p.get("category") and r["category"] not in p["category"]:
                continue
            if not group_matches(r["group"], p.get("groups")):
                continue
            if p.get("taxon") and not _taxon_mask(pd.Series([r["species"]]), p["taxon"])[0]:
                continue
            rasters.append(r["path"])
    return pts, polys, rasters


def _sdm_patches(data: Data, raster_paths, thr=SDM_THRESHOLD):
    """Растры SDM → полигоны ≥ thr (по каждому виду)."""
    out = []
    for p in raster_paths:
        try:
            arr, tr, _ = read_raster_clipped(p, data.bounds_pad)
            if arr is None:
                continue
            inside = data.inside_mask(tr, arr.shape)
            mask = np.nan_to_num(arr, nan=-1) >= thr
            polys = polygonize_mask(mask & inside, tr)
            polys["species"] = species_from_filename(p)
            out.append(polys)
            log.info("  SDM %s: %d пятен ≥ %.2f", Path(p).name, len(polys), thr)
        except Exception as e:
            log.error("  SDM %s не обработан: %s", p, e)
    if not out:
        return None
    return gpd.GeoDataFrame(pd.concat(out, ignore_index=True), crs=WORK_EPSG)


def _sdm_monitoring(data: Data, res: Results, measure: dict, p: dict, rasters, label: str):
    """Модели SDM: отдельная строка «… — мониторинг (SDM ≥ порог)» по участкам ≥ порога."""
    thr = p["sdm"]
    sdm = _sdm_patches(data, rasters, thr=thr)
    if sdm is None or sdm.empty:
        return
    m2 = dict(measure)
    m2["name"] = measure["name"] + " — мониторинг (SDM ≥ %.2f)" % thr
    target = _target(data, p)
    j = gpd.overlay(sdm, target[["uid", "area_ha", "geometry"]], how="intersection")
    j["ha"] = j.area / 1e4
    share = spec_share(p)
    for uid, grp in j.groupby("uid"):
        tot = grp.groupby("species")["ha"].sum()
        tot = tot[tot >= p["min_ha"]]
        if tot.empty:
            continue
        r = data.oozt[data.oozt["uid"] == uid].iloc[0]
        if 100 * tot.sum() / max(r["area_ha"], 1e-9) < share:
            continue
        res.add(r, m2, f"вероятность SDM ≥ {_fmt(thr)}: " + ", ".join(f"{s} ({a:.2f} га)" for s, a in tot.items()),
                round(float(tot.sum()), 2))
    res.add_evidence(j, m2["name"], f"{label} — SDM (≥ {thr:.2f})", "species")


def rule_species(data: Data, res: Results, measure: dict, p: dict):
    """Находки видов: полигоны с точками / пересекающиеся с полигонами находок; SDM — мониторинг."""
    pts, polys, rasters = select_species(data, p)
    lab = _sp_label(p)
    target = _target(data, p)
    if pts is not None:
        j = gpd.sjoin(pts, target[["uid", "geometry"]], how="inner", predicate="within")
        for uid, grp in j.groupby("uid"):
            r = data.oozt[data.oozt["uid"] == uid].iloc[0]
            sp = grp.groupby("species").size().sort_values(ascending=False)
            coords = []
            for s_, g in grp.groupby("species"):
                gg = gpd.GeoSeries(g.geometry, crs=WORK_EPSG).to_crs(4326)
                coords.append(f"{s_} ({gg.x.mean():.6f}° в.д., {gg.y.mean():.6f}° с.ш.)")
            res.add(r, measure, f"точечные находки: {', '.join(f'{k} ({v})' for k, v in sp.items())}; "
                                f"координаты WGS 84: {'; '.join(coords)}", int(len(grp)))
        res.add_evidence(j, measure["name"], f"{lab} — точки", "species")
        log.info("  точки: %d %s", j["uid"].nunique(), _units())
    if polys is not None:
        j = gpd.overlay(polys, target[["uid", "geometry"]], how="intersection")
        j["ha"] = j.area / 1e4
        for uid, grp in j.groupby("uid"):
            r = data.oozt[data.oozt["uid"] == uid].iloc[0]
            txt = ", ".join(f"{s_} ({a:.2f} га)" for s_, a in grp.groupby("species")["ha"].sum().items())
            res.add(r, measure, f"в пределах полигонов находок: {txt}", round(grp["ha"].sum(), 2))
        res.add_evidence(j, measure["name"], f"{lab} — полигоны", "species")
        log.info("  полигоны: %d %s", j["uid"].nunique(), _units())
    _sdm_monitoring(data, res, measure, p, rasters, lab)


def rule_species_trails(data: Data, res: Results, measure: dict, p: dict):
    """Тропы в заданном расстоянии от точек и внутри полигонов видов; SDM — мониторинг."""
    pts, polys, rasters = select_species(data, p)
    lab, lab_gen, dist = _sp_label(p), _sp_label(p, gen=True), p["dist"]
    trails = data.trails
    target = _target(data, p)
    zones = []
    if pts is not None:
        b = pts.copy()
        b["geometry"] = b.geometry.buffer(dist)
        b["src"] = "точки (буфер %d м)" % dist
        zones.append(b)
    if polys is not None:
        q_ = polys.copy()
        q_["src"] = "полигоны"
        zones.append(q_)
    if zones:
        z = gpd.GeoDataFrame(pd.concat(zones, ignore_index=True), crs=WORK_EPSG)
        if trails is not None:
            zu = unary_union(z.geometry.values)
            t = trails[trails.intersects(zu)].copy()
            t["geometry"] = t.geometry.intersection(zu)
            t = t[~t.is_empty]
            if not t.empty:
                j = gpd.overlay(t.explode(index_parts=False), target[["uid", "geometry"]], how="intersection",
                                keep_geom_type=False)
                j = j[j.geom_type.isin(["LineString", "MultiLineString"])]
                j["len"] = j.length
                for uid, grp in j.groupby("uid"):
                    r = data.oozt[data.oozt["uid"] == uid].iloc[0]
                    zz = z[z.intersects(r.geometry)]
                    sp = ", ".join(sorted(zz["species"].unique())[:8])
                    res.add(r, measure, f"тропы в зоне {lab_gen} ({sp}): {grp['len'].sum():.0f} м",
                            round(grp["len"].sum()))
                res.add_evidence(j, measure["name"], f"тропы на расстоянии {dist:.0f} м от {lab_gen}")
                log.info("  тропы на расстоянии %.0f м от %s: %d %s", dist, lab_gen, j["uid"].nunique(), _units())
            res.add_evidence(pts[pts.intersects(data.union)] if pts is not None else None,
                             measure["name"], f"{lab} — точки", "species")
            res.add_evidence(polys, measure["name"], f"{lab} — полигоны", "species")
        else:
            # без троп — назначаем по самим зонам
            j = gpd.overlay(z, target[["uid", "geometry"]], how="intersection")
            for uid, grp in j.groupby("uid"):
                r = data.oozt[data.oozt["uid"] == uid].iloc[0]
                res.add(r, measure, f"местообитания {lab_gen}: " + ", ".join(sorted(grp["species"].unique())[:8]))
            res.add_evidence(j, measure["name"], lab, "species")
    _sdm_monitoring(data, res, measure, p, rasters, lab)


def rule_slope(data: Data, res: Results, measure: dict, p: dict):
    if data.dem is None:
        log.warning("  ЦМР не загружена — пропуск")
        return
    threshold_pct = p["slope"]
    arr, tr = data.dem
    if getattr(data, "_slope", None) is None:
        data._slope = slope_percent(arr, CELL)
    slope = data._slope
    inside = data.inside_mask(tr, slope.shape)
    mask = inside & (np.nan_to_num(slope, nan=0) > threshold_pct)
    # убираем одиночные ячейки
    mask = ndimage.binary_opening(mask, structure=np.ones((2, 2)))
    polys = polygonize_mask(mask, tr)
    target = _target(data, p)
    hit = _overlay_patches_to_vydels(data, polys, target, p["min_ha"], spec_share(p))
    eco_label = _eco_label(p)
    for _, h in hit.iterrows():
        r = data.oozt[data.oozt["uid"] == h["uid"]].iloc[0]
        res.add(r, measure, f"уклон более {threshold_pct:g}% {eco_label}: "
                           f"{h['patch_ha']:.2f} га ({h['share'] * 100:.0f} % {_unit_word()})",
                round(h["patch_ha"], 2))
    if not hit.empty:
        sel = target[target["uid"].isin(hit["uid"])]
        pieces = polys[polys.intersects(unary_union(sel.geometry.values))]
        res.add_evidence(pieces, measure["name"], f"уклон более {threshold_pct:g}% {eco_label}")
    log.info("  → %d %s (%d пятен)", len(hit), _units(), len(polys))


def rule_species_density(data: Data, res: Results, measure: dict, p: dict):
    """Плотность точек выбранных видов (например, рода Corvus) в окне; участки ≥ квантиля."""
    pts, _, _ = select_species(data, dict(p, sources=["points"]))
    lab = _sp_label(p)
    if pts is None or pts.empty:
        log.warning("  точки видов (%s) не найдены — пропуск", _sel_txt(p))
        return
    w, q = p["window"], p["quantile"]
    dens, tr = point_density(pts, data.bounds_pad, w, CELL)
    inside = data.inside_mask(tr, dens.shape)
    mask, thr = upper_quantile_mask(dens, inside, q)
    polys = polygonize_mask(mask, tr)
    target = _target(data, p)
    hit = _overlay_patches_to_vydels(data, polys, target, p["min_ha"], spec_share(p))
    for _, h in hit.iterrows():
        r = data.oozt[data.oozt["uid"] == h["uid"]].iloc[0]
        res.add(r, measure, f"плотность {lab} (окно {w:.0f} м) ≥ {_qtag(q)} = {thr:.3f} шт/га: "
                            f"{h['patch_ha']:.2f} га ({100 * h['share']:.0f}% полигона)", round(h["patch_ha"], 2))
    res.add_evidence(polys[polys.intersects(data.union)], measure["name"], f"плотность {lab} ({_qtag(q)})")
    res.add_evidence(pts[pts.intersects(data.union.buffer(w))], measure["name"], f"{lab} — точки", "species")
    log.info("  → %d %s (%d точек: %s)", len(hit), _units(), len(pts), lab)


RULE_FUNCS = {
    "type": rule_type, "trail_density": rule_trail_density_vydel, "trail_hotspot": rule_trail_hotspot,
    "shore": rule_shore_trails, "slope": rule_slope, "species": rule_species,
    "species_trails": rule_species_trails, "species_density": rule_species_density,
}


def dispatch_rule(measure: dict, data: Data, res: Results) -> bool:
    spec = measure_spec(measure)
    if spec["rule"] is None:
        if spec["reason"] not in ("метод выявления не задан", "по таблице: нет данных"):
            log.warning("  %s: %s — пропуск", measure["name"], spec["reason"])
        return False
    log.info("Мероприятие: %s%s", (measure.get("code") + " ") if measure.get("code") else "", measure["name"])
    st = measure_status(measure, data.cfg)
    if st["state"] == "missing":
        log.warning("  %s — пропуск", st["text"])
        return False
    RULE_FUNCS[spec["rule"]](data, res, measure, spec["p"])
    return True


def spec_pad_m(measures) -> float:
    """Запас вокруг ООЗТ для расчётов: наибольшее окно или расстояние среди включённых мероприятий."""
    v = 0.0
    for m in measures or []:
        if m.get("on"):
            p = measure_spec(m).get("p", {})
            v = max([v] + [float(p[k]) for k in ("window", "dist") if p.get(k)])
    return v


# =====================================================================================
#                              ТАБЛИЦА МЕРОПРИЯТИЙ (Excel)
# =====================================================================================
def read_measures(path: str) -> list[dict]:
    df = pd.read_excel(path)
    df.columns = [str(c).strip() for c in df.columns]

    def col(*names):
        for n in names:
            for c in df.columns:
                if c.lower().startswith(n.lower()):
                    return c
        return None

    c_name, c_on, c_method = col("Мероприяти"), col("Вкл"), col("Метод")
    c_group, c_eco, c_prikaz, c_season = col("Тип мероприят"), col("Тип экосистем"), col("Соответствие"), col("Сезон")
    if not (c_name and c_on and c_method):
        raise ValueError(f"В Excel не найдены колонки «Мероприятия», «Вкл/выкл», «Метод выявления»: {list(df.columns)}")
    df[c_group] = df[c_group].ffill() if c_group else ""
    code_re = re.compile(r"^[IVX]+\.\d+$")
    c_code = next((c for c in df.columns
                   if df[c].dropna().astype(str).str.strip().str.match(code_re).mean() > 0.6), None)
    spec_cols = {}
    for key, title, stem, _ in SPEC_COLUMNS:
        c = next((c for c in df.columns if re.sub(r"\s+", " ", c.lower().replace("ё", "е")).startswith(stem)), None)
        if c is not None:
            spec_cols[key] = c
    if spec_cols:
        log.info("Столбцы правил в таблице: %s", ", ".join(spec_cols.values()))
    out = []
    for i, r in df.iterrows():
        if pd.isna(r[c_name]):
            continue
        on = str(r[c_on]).strip().lower() in ("1", "1.0", "true", "да", "yes", "вкл")
        out.append({
            "name": re.sub(r"\s+", " ", str(r[c_name])).strip(),
            "code": str(r[c_code]).strip() if c_code and not pd.isna(r[c_code]) else "",
            "on": on,
            "method": None if pd.isna(r[c_method]) else str(r[c_method]).strip(),
            "group": re.sub(r"\s+", " ", str(r[c_group])).strip() if c_group and not pd.isna(r[c_group]) else "",
            "eco": str(r[c_eco]).strip() if c_eco and not pd.isna(r[c_eco]) else "",
            "prikaz": str(r[c_prikaz]).strip() if c_prikaz and not pd.isna(r[c_prikaz]) else "",
            "season": str(r[c_season]).strip() if c_season and not pd.isna(r[c_season]) else "",
            "spec": {k: (None if pd.isna(r[c]) else r[c]) for k, c in spec_cols.items()},
            "row": int(i) + 2,                 # номер строки в Excel (для записи описаний)
        })
    log.info("Таблица мероприятий: %d строк, включено %d", len(out), sum(m["on"] for m in out))
    return out


def _cell_value(key: str, v):
    """Параметр правила → значение ячейки Excel."""
    if key == "types":
        return v[1] if v else "любые"
    if key in ("category", "groups"):
        return ", ".join(v) if v else "любые"
    if key == "sources":
        return ", ".join({"points": "точки", "polygons": "полигоны", "sdm": "SDM"}[x] for x in v) if v else "любые"
    if key == "taxon":
        return ", ".join(v) if v else None
    return None if v is None else float(v)


def update_measures_xlsx(path: str, cfg: Config | None = None, fill_from_method: bool = True,
                         backup: bool = True) -> str:
    """Дополняет таблицу мероприятий: столбцы правил (заполняются из текста метода там, где правило не задано),
    выпадающие списки, лист «Справка по правилам», столбец с полным описанием расчёта. Возвращает сводку."""
    import shutil
    import openpyxl
    from copy import copy as _cp
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    if cfg is not None:
        _MIN_SHARE[0] = float(getattr(cfg, "min_share_pct", 0) or 0)
    path = str(path)
    measures = read_measures(path)
    if backup:
        bak = Path(path).with_name(Path(path).stem + "_до_правил" + Path(path).suffix)
        if not bak.exists():
            shutil.copy2(path, bak)
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    norm = lambda v: re.sub(r"\s+", " ", str(v or "").strip().lower().replace("ё", "е"))
    hdr = {c: norm(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)}
    last = max([c for c, h in hdr.items() if h] or [1])
    proto = ws.cell(1, last)

    def ensure(title, stem, width):
        nonlocal last
        c = next((c for c, h in hdr.items() if h.startswith(stem)), None)
        if c is None:
            last += 1
            c = last
            cell = ws.cell(1, c, title)
            cell.font, cell.border = _cp(proto.font), _cp(proto.border)
            cell.fill = PatternFill("solid", fgColor="E8F0E0")
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            hdr[c] = norm(title)
            ws.column_dimensions[get_column_letter(c)].width = width
        return c

    widths = dict(rule=24, types=16, category=14, groups=18, taxon=14, sources=14)
    cols = {k: ensure(t, stem, widths.get(k, 11)) for k, t, stem, _ in SPEC_COLUMNS}
    c_desc = ensure(DESC_COLUMN, "процедура назначения", 95)
    n_fill = n_desc = 0
    wrap = Alignment(wrap_text=True, vertical="top")
    for m in measures:
        r = m["row"]
        spec = measure_spec(m)
        if fill_from_method and m["on"] and spec["rule"] and spec["origin"] == "текст метода":
            ws.cell(r, cols["rule"], RULES[spec["rule"]]["title"])
            for k, v in spec["p"].items():
                val = _cell_value(k, v)
                if val is not None and k != "min_share":
                    ws.cell(r, cols[k], val)
            m = dict(m, spec={k: ws.cell(r, c).value for k, c in cols.items()})
            spec = measure_spec(m)
            n_fill += 1
        for k, c in cols.items():
            ws.cell(r, c).alignment = wrap
        if spec["rule"] or m["on"]:
            txt = describe_spec(spec, full=True)
            ws.cell(r, c_desc, txt).alignment = wrap
            n_desc += 1
            need = (len(txt) // int(95 * 1.15) + 2) * 15
            if (ws.row_dimensions[r].height or 15) < need:
                ws.row_dimensions[r].height = need
    # выпадающие списки (подсказки: можно ввести и своё значение / несколько через запятую)
    ours = set(cols.values())
    ws.data_validations.dataValidation = [
        dv for dv in ws.data_validations.dataValidation
        if not any(rng.min_col in ours for rng in dv.sqref.ranges)]
    nrow = max(ws.max_row, 2)
    lists = {"rule": RULE_TITLES, "types": ["лесные", "луговые", "водные", "околоводные", "болотные",
                                             "водные, околоводные, болотные", "любые"],
             "category": SPECIES_CATEGORIES + ["любые"], "groups": SPECIES_GROUPS[:-1] + ["любые"],
             "sources": ["точки", "полигоны", "SDM", "точки, полигоны", "любые"]}
    for k, vals in lists.items():
        f = '"' + ",".join(v.replace(", ", "; ") for v in vals) + '"'
        dv = DataValidation(type="list", formula1=f, allow_blank=True, showErrorMessage=(k == "rule"),
                            errorTitle="Правило назначения",
                            error="Выберите правило из списка (см. лист «Справка по правилам»).")
        dv.add(f"{get_column_letter(cols[k])}2:{get_column_letter(cols[k])}{nrow}")
        ws.add_data_validation(dv)
    # лист-справка
    if "Справка по правилам" in wb.sheetnames:
        del wb["Справка по правилам"]
    h = wb.create_sheet("Справка по правилам")
    bold = Font(bold=True)
    rows = [("Как скрипт читает способ назначения мероприятия", None, None, None), (None,) * 4,
            ("Строка таблицы рассчитывается, если «Вкл/выкл» = 1. Правило берётся из столбца «Правило назначения»; "
             "параметры — из столбцов справа. Пустой параметр — значение по умолчанию (указано ниже). Если "
             "«Правило назначения» пусто, правило и параметры определяются по ключевым словам «Метода выявления» "
             "(как в прежних версиях). Столбец «" + DESC_COLUMN + "» заполняется скриптом и только описывает расчёт.",
             None, None, None), (None,) * 4,
            ("Правило назначения", "Что делает", "Нужные данные", "Параметры (по умолчанию)")]
    data_t = {"trails": "тропы", "hydro": "гидрография", "dem": "ЦМР", "species": "слои видов"}
    for k, v in RULES.items():
        ps = []
        for pk, d in v["params"].items():
            dv_ = ("обязательно" if d == REQ else "общий порог из окна скрипта" if d == SHARED else
                   "любые" if d is None else _fmt(d))
            ps.append(f"{SPEC_TITLE[pk]}: {dv_}")
        rows.append((v["title"], v["help"],
                     ", ".join(data_t[d] for d in v["data"]) + (" (+ тропы, если есть)" if v["opt"] else "")
                     or "только слой территории", "; ".join(ps)))
    rows += [(None,) * 4, ("Столбец", "Что указывать", None, None)]
    rows += [(t, help_, None, None) for _, t, _, help_ in SPEC_COLUMNS]
    rows += [(None,) * 4, ("Слои видов (задаются в окне скрипта)", "Значения", None, None),
             ("Категория", "; ".join(f"{k} — {v}" for k, v in SPECIES_CATEGORY_HELP.items()), None, None),
             ("Группа организмов", ", ".join(SPECIES_GROUPS[:-1]) + "; можно ввести новую группу", None, None),
             ("Тип данных", "точки, полигоны — векторные слои (поле с латинскими названиями находится само); "
                            "SDM — растр вероятности 0–1, название вида = имя файла", None, None)]
    for row in rows:
        h.append(list(row))
    for rr in h.iter_rows():
        for cell in rr:
            cell.alignment = wrap
    for rr in (1, 5):
        for cell in h[rr]:
            cell.font = bold
    for rr in range(1, h.max_row + 1):
        if h.cell(rr, 1).value in ("Столбец", "Слои видов (задаются в окне скрипта)"):
            for cell in h[rr]:
                cell.font = bold
    h["A1"].font = Font(bold=True, size=13)
    h["A1"].alignment = Alignment(wrap_text=False)
    h.page_setup.orientation = "landscape"
    h.sheet_properties.pageSetUpPr.fitToPage = True
    h.page_setup.fitToWidth, h.page_setup.fitToHeight = 1, 0
    h.freeze_panes = "A6"
    for col, w in zip("ABCD", (30, 70, 26, 60)):
        h.column_dimensions[col].width = w
    h.merge_cells("A3:D3")
    h.row_dimensions[3].height = 60
    wb.save(path)
    return (f"Таблица обновлена: {Path(path).name}. Правила записаны в {n_fill} строк(и), описания — в {n_desc}; "
            f"добавлен лист «Справка по правилам».")


# =====================================================================================
#                                       ОТЧЁТ
# =====================================================================================
LAYER_STYLE = {
    # layer → (цвет, тип: point/line/poly)
    "тропы в зонах повышенной плотности": ("#d62728", "line"),
    "зоны повышенной плотности троп (верхний квартиль)": ("#ff7f0e", "poly"),
    f"тропы на расстоянии {SHORE_BUFFER_M:.0f} м от водных объектов (Q75)": ("#1f77b4", "poly"),
    "водные объекты (линии)": ("#3182bd", "line"),
    "водные объекты (полигоны)": ("#9ecae1", "poly"),
    f"тропы на расстоянии {RARE_BUFFER_M:.0f} м от редких видов": ("#9467bd", "line"),
    "редкие виды — точки": ("#2ca02c", "point"),
    "редкие виды — полигоны": ("#2ca02c", "poly"),
    "редкие виды — SDM": ("#98df8a", "poly"),
    "редкие виды": ("#2ca02c", "poly"),
    "инвазионные виды — точки": ("#e377c2", "point"),
    "инвазионные виды — полигоны": ("#e377c2", "poly"),
    f"инвазионные виды — SDM (≥ {SDM_THRESHOLD_INVASIVE:.2f})": ("#f7b6d2", "poly"),
    "Corvus — точки": ("#000000", "point"),
    "плотность Corvus (Q75)": ("#7f7f7f", "poly"),
}
TYPE_COLORS = {"лес": "#b7dfb0", "луг": "#f5e9a6", "околовод": "#bfe3ea", "водн": "#9ec9ff"}


_TYPE_FALLBACK = ["#e6d3f5", "#f7d6c4", "#d9d9d9", "#cfe8cf", "#ffe0e0"]


def _type_color(t: str):
    t = str(t).lower()
    for k, c in TYPE_COLORS.items():
        if k in t:
            return c
    # неизвестный тип выдела: устойчивый цвет из резервной палитры (чтобы попал в легенду)
    return _TYPE_FALLBACK[hash(t) % len(_TYPE_FALLBACK)]


def _slope_style(layer):
    if layer.startswith("уклон"):
        # 50% и 70% чёрного соответственно.
        return ("#808080" if "5%" in layer else "#4D4D4D", "slope")
    if layer in LAYER_STYLE:
        return LAYER_STYLE[layer]
    return _pattern_style(layer)


def _pattern_style(layer: str):
    """Стиль слоёв-обоснований, подписи которых зависят от параметров мероприятия."""
    t = str(layer).lower()
    inv = "инваз" in t
    if t.startswith("тропы на расстоянии"):
        return ("#1f77b4", "poly") if "водн" in t else ("#9467bd", "line")
    if "— sdm" in t:
        return ("#f7b6d2" if inv else "#98df8a", "poly")
    if t.startswith("плотность"):
        return ("#7f7f7f", "poly")
    if t.startswith("зоны повышенной плотности троп"):
        return ("#ff7f0e", "poly")
    if "— точки" in t:
        return ("#e377c2" if inv else "#2ca02c" if ("охран" in t or "редк" in t) else "#000000", "point")
    if "— полигоны" in t or "виды" in t or "видов" in t:
        return ("#e377c2" if inv else "#2ca02c", "poly")
    return ("#17becf", "poly")


_BASEMAP_PROVIDER: list = []


def _basemap_candidates():
    import contextily as cx
    # Любой загруженный растр ниже принудительно переводится в серый.
    # OSM ставим первым: CartoDB Positron в некоторых средах показывает водяной знак API KEY REQUIRED.
    return [cx.providers.OpenStreetMap.Mapnik, cx.providers.Esri.WorldGrayCanvas]


def pick_basemap_provider(skip=0):
    """Выбор доступного тайлового источника; отображение затем переводится в оттенки серого."""
    if _BASEMAP_PROVIDER and not skip:
        return _BASEMAP_PROVIDER[0]
    import contextily as cx
    import requests
    cx.tile.USER_AGENT = "oozt_recs/1.0 (natural-area management report generator)"
    try:                                   # кэш тайлов: повторные прогоны не тянут их заново
        cx.set_cache_dir(str(Path(tempfile.gettempdir()) / "oozt_tiles"))
    except Exception:
        pass
    candidates = _basemap_candidates()[skip:]
    if skip:
        _BASEMAP_PROVIDER.clear()
    for prov in candidates:
        try:
            url = prov.build_url(x=39530, y=20465, z=16)   # тайл над Москвой
            r = requests.get(url, headers={"user-agent": cx.tile.USER_AGENT}, timeout=15)
            if r.status_code == 200 and r.headers.get("content-type", "").startswith("image"):
                _BASEMAP_PROVIDER.append(prov)
                log.info("Подложка: %s", prov.get("name"))
                return prov
            log.warning("Тайлы %s недоступны (HTTP %s)", prov.get("name"), r.status_code)
        except Exception as e:
            log.warning("Тайлы %s недоступны: %s", prov.get("name"), e)
    _BASEMAP_PROVIDER.append(None)
    return None


def desaturate_last_basemap(ax):
    """Перевод последнего растрового слоя осей в оттенки серого."""
    if not ax.images:
        return
    im = ax.images[-1]
    arr = np.asarray(im.get_array())
    if arr.ndim != 3 or arr.shape[2] < 3:
        return
    rgb = arr[..., :3].astype(float)
    gray = 0.2126 * rgb[..., 0] + 0.7152 * rgb[..., 1] + 0.0722 * rgb[..., 2]
    out = np.repeat(gray[..., None], 3, axis=2)
    if arr.shape[2] == 4:
        out = np.dstack([out, arr[..., 3]])
    im.set_data(out.astype(arr.dtype))


def _deg_label(v: float, step: float) -> str:
    """Подпись координатной сетки в десятичных градусах WGS 84."""
    nd = max(3, int(np.ceil(-np.log10(max(step, 1e-12)))) + 1)
    return f"{v:.{nd}f}°"


def set_degree_ticks(ax, ext):
    """Оси подписываются градусами широты/долготы (WGS 84), сама карта остаётся в UTM."""
    from pyproj import Transformer
    fwd = Transformer.from_crs(f"EPSG:{WORK_EPSG}", "EPSG:4326", always_xy=True)
    inv = Transformer.from_crs("EPSG:4326", f"EPSG:{WORK_EPSG}", always_xy=True)
    xc, yc = (ext[0] + ext[1]) / 2, (ext[2] + ext[3]) / 2
    lon0, lat0 = fwd.transform(ext[0], yc)
    lon1, lat1 = fwd.transform(ext[1], yc)
    _, latmin = fwd.transform(xc, ext[2])
    _, latmax = fwd.transform(xc, ext[3])
    def ticks(v0, v1, to_axis):
        span = max(v1 - v0, 1e-12)
        target = span / 4
        magnitude = 10 ** np.floor(np.log10(target))
        step = next(m * magnitude for m in (1, 2, 2.5, 5, 10)
                    if m * magnitude >= target)
        first = np.ceil(v0 / step) * step
        vals = np.arange(first, v1 + step / 2, step)
        return [to_axis(v) for v in vals], [_deg_label(v, step) for v in vals]

    xt, xl = ticks(lon0, lon1, lambda v: inv.transform(v, (lat0 + lat1) / 2)[0])
    yt, yl = ticks(latmin, latmax, lambda v: inv.transform((lon0 + lon1) / 2, v)[1])
    ax.set_xticks(xt); ax.set_xticklabels(xl)
    ax.set_yticks(yt); ax.set_yticklabels(yl)
    ax.set_xlabel("Долгота, в.д. (WGS 84)", fontsize=12)
    ax.set_ylabel("Широта, с.ш.", fontsize=12)


def _legend_label(layer: str) -> str:
    """Перенос длинных подписей, чтобы легенда не растягивала схему."""
    if len(layer) > 45 and " м от " in layer:
        return layer.replace(" м от ", " м\nот ", 1)
    return layer


def _place_labels(ax, items, sizes=(11, 10, 9, 8, 7), leaders=True, max_labels=150):
    """Подписи полигонов без наложений.
    items: [(текст, геометрия)] в порядке приоритета. Подпись ставится внутрь полигона, если помещается
    (перебор кеглей sizes); иначе — рядом с полигоном с выноской; если и так не помещается — пропускается.
    Возвращает число размещённых подписей."""
    import matplotlib.patheffects as pe
    from shapely.geometry import box as _box
    fig = ax.figure
    ax.apply_aspect()
    pos = ax.get_position()
    x0, x1 = ax.get_xlim()
    y0, y1 = ax.get_ylim()
    m_per_pt = (x1 - x0) / (pos.width * fig.get_figwidth() * 72)
    frame = _box(x0, y0, x1, y1)
    placed = []
    halo = [pe.withStroke(linewidth=2.6, foreground="white")]

    def lbl_box(cx, cy, t, fs):
        hw = (len(t) * fs * 0.62 + 3) * m_per_pt / 2
        hh = (fs * 1.05 + 2) * m_per_pt / 2
        return _box(cx - hw, cy - hh, cx + hw, cy + hh)

    def free(b):
        return frame.contains(b) and not any(b.intersects(q) for q in placed)
    n = 0
    for t, geom in items:
        if n >= max_labels:
            break
        if geom is None or geom.is_empty:
            continue
        p = geom.representative_point()
        done = False
        for fs in sizes:                                   # 1) внутри полигона
            b = lbl_box(p.x, p.y, t, fs)
            if geom.buffer(fs * 0.25 * m_per_pt).contains(b) and free(b):
                ax.annotate(t, (p.x, p.y), ha="center", va="center", fontsize=fs, fontweight="bold",
                            color="#111111", zorder=13, path_effects=halo)
                placed.append(b); done = True
                break
        if done or not leaders:
            n += done
            continue
        fs = sizes[-1] + 1                                 # 2) рядом с выноской
        gx0, gy0, gx1, gy1 = geom.bounds
        b0 = lbl_box(0, 0, t, fs)
        bw, bh = b0.bounds[2] - b0.bounds[0], b0.bounds[3] - b0.bounds[1]
        off = 6 * m_per_pt
        cands = [(gx1 + off + bw / 2, (gy0 + gy1) / 2), (gx0 - off - bw / 2, (gy0 + gy1) / 2),
                 ((gx0 + gx1) / 2, gy1 + off + bh / 2), ((gx0 + gx1) / 2, gy0 - off - bh / 2),
                 (gx1 + off + bw / 2, gy1 + off + bh / 2), (gx0 - off - bw / 2, gy1 + off + bh / 2),
                 (gx1 + off + bw / 2, gy0 - off - bh / 2), (gx0 - off - bw / 2, gy0 - off - bh / 2),
                 (gx1 + 3 * off + bw / 2, (gy0 + gy1) / 2), (gx0 - 3 * off - bw / 2, (gy0 + gy1) / 2)]
        for cx, cy in cands:
            b = lbl_box(cx, cy, t, fs)
            if free(b):
                ax.annotate(t, xy=(p.x, p.y), xytext=(cx, cy), textcoords="data", ha="center", va="center",
                            fontsize=fs, fontweight="bold", color="#111111", zorder=13, path_effects=halo,
                            arrowprops=dict(arrowstyle="-", color="#333333", lw=0.6, shrinkA=0, shrinkB=0))
                placed.append(b); n += 1
                break
    return n


def draw_kvartal(data: Data, evidence: gpd.GeoDataFrame | None, kv, oozt_name, out_png: Path, basemap=True):
    vyd = data.oozt[(data.oozt["KV"] == kv) & (data.oozt["OOZT"] == oozt_name)]
    if vyd.empty:
        return None
    focus_geom = unary_union(vyd.geometry.values)
    minx, miny, maxx, maxy = vyd.total_bounds
    pad = max((maxx - minx), (maxy - miny)) * 0.12 + 30
    ext = (minx - pad, maxx + pad, miny - pad, maxy + pad)
    w, h = ext[1] - ext[0], ext[3] - ext[2]
    fig_w = 11
    fig_h = max(6, min(11, fig_w * h / w))
    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=200)
    plt.rcParams.update({"font.size": 12})

    # Лесные выделы имеют зелёную заливку, луговые прозрачны.
    # Границы всех выделов повторно рисуются поверх тематических слоёв ниже.
    cfg = data.cfg
    group_mode = cfg.unit_mode == "eco" and bool(cfg.report_group_field)   # индекс = тип, полигоны разбросаны
    map_box0 = box(ext[0], ext[2], ext[1], ext[3])
    same_oozt0 = data.oozt[data.oozt["OOZT"] == oozt_name]
    others = same_oozt0[(same_oozt0["KV"] != kv) & same_oozt0.intersects(map_box0)]
    if cfg.unit_mode == "eco" and not others.empty:      # прочие полигоны — тонкий контур без подписей
        others.boundary.plot(ax=ax, color="#9AA0A6", linewidth=0.45, zorder=3)
    is_forest = vyd["Type"].astype(str).str.lower().str.contains("лес", regex=False)
    forest = vyd[is_forest]
    if not forest.empty:
        forest.plot(ax=ax, color="#b7dfb0", edgecolor="none", alpha=0.55, zorder=1)
    if group_mode and (~is_forest).any():
        vyd[~is_forest].plot(ax=ax, color="#F6E3A1", edgecolor="none", alpha=0.6, zorder=1)
    # все тропы в окне карты — тонкий серый контекстный слой
    if data.trails is not None and not data.trails.empty:
        map_box = box(ext[0], ext[2], ext[1], ext[3])
        all_trails = data.trails[data.trails.intersects(map_box)]
        if not all_trails.empty:
            all_trails.plot(ax=ax, color="#707070", linewidth=0.7, alpha=0.8, zorder=2)
    # основания
    handles = []
    if data.trails is not None and not data.trails.empty:
        handles.append(Line2D([], [], color="#707070", linewidth=1.0, label="все тропы"))
    if cfg.unit_mode == "eco":
        grp_name = f"полигоны {unit_labels(cfg)['group'].lower()} {kv}" if group_mode else f"полигоны листа {kv}"
        if is_forest.any():
            handles.append(Patch(facecolor="#b7dfb0", alpha=0.7, edgecolor="#333333", label=grp_name + " — лесные"))
        if (~is_forest).any():
            handles.append(Patch(facecolor="#F6E3A1" if group_mode else "none", alpha=0.7, edgecolor="#333333",
                                 label=grp_name + (" — прочие" if is_forest.any() else "")))
        if not others.empty:
            handles.append(Patch(facecolor="none", edgecolor="#9AA0A6", label="прочие полигоны"))
    used_layers = []
    if evidence is not None and not evidence.empty:
        kv_geom = unary_union(vyd.geometry.values)
        ev = evidence[evidence.intersects(kv_geom.buffer(5))]
        for layer, grp in ev.groupby("layer"):
            color, kind = _slope_style(layer)
            if kind == "point":
                species_point = layer in ("редкие виды — точки", "инвазионные виды — точки")
                point_size = 60 if species_point else 40
                legend_size = 13.5 if species_point else 9
                grp.plot(ax=ax, color=color, markersize=point_size,
                         edgecolor="white", linewidth=0.6, zorder=6)
                handles.append(Line2D([], [], marker="o", color="w", markerfacecolor=color,
                                      markersize=legend_size, label=_legend_label(layer)))
            elif kind == "line":
                grp.plot(ax=ax, color=color, linewidth=2.2, zorder=5)
                handles.append(Line2D([], [], color=color, linewidth=2.5,
                                      label=_legend_label(layer)))
            elif kind == "slope":
                grp.plot(ax=ax, facecolor=color, edgecolor=color,
                         alpha=1.0, linewidth=0.8, zorder=4)
                handles.append(Patch(facecolor=color, edgecolor=color,
                                     label=_legend_label(layer)))
            elif "SDM" in layer:      # модельные ареалы занимают большие площади — только штриховка
                grp.plot(ax=ax, facecolor="none", edgecolor=color, alpha=0.9,
                         linewidth=0.9, hatch="////", zorder=4)
                handles.append(Patch(facecolor="none", edgecolor=color, hatch="////",
                                     label=_legend_label(layer)))
            else:
                grp.plot(ax=ax, facecolor=color, edgecolor=color, alpha=0.45, linewidth=0.8, zorder=4)
                handles.append(Patch(facecolor=color, alpha=0.5, edgecolor=color,
                                     label=_legend_label(layer)))
            used_layers.append(layer)
    # границы выделов поверх всех тематических слоёв
    vyd.boundary.plot(ax=ax, color="#222222", linewidth=1.6 if group_mode else 1.15, zorder=8)
    # подписи выделов
    ax.set_xlim(ext[0], ext[1]); ax.set_ylim(ext[2], ext[3])
    lab_items = [(str(r["VYD"]).removesuffix(".0"), r.geometry)
                 for _, r in vyd.assign(_a=vyd.geometry.area).sort_values("_a", ascending=False).iterrows()]
    # соседние кварталы / листы: границы и номера в пределах окна карты (в режиме индексов — не нужны)
    map_box = box(ext[0], ext[2], ext[1], ext[3])
    same_oozt = data.oozt[data.oozt["OOZT"] == oozt_name]
    neighbors = same_oozt[(same_oozt["KV"] != kv) & same_oozt.intersects(map_box)]
    if group_mode:
        neighbors = neighbors.iloc[0:0]
    import matplotlib.patheffects as pe
    for n_kv, ng in neighbors.groupby("KV"):
        ngeom = unary_union(ng.geometry.values)
        clipped = ngeom.intersection(map_box)
        if clipped.is_empty:
            continue
        gpd.GeoSeries([ngeom], crs=WORK_EPSG).boundary.plot(
            ax=ax, color="#444444", linewidth=1.5, linestyle="--", zorder=9)
        p = clipped.representative_point()
        px = float(np.clip(p.x, ext[0] + w * 0.07, ext[1] - w * 0.07))
        py = float(np.clip(p.y, ext[2] + h * (0.15 if px < ext[0] + w * 0.35 else 0.09), ext[3] - h * 0.07))
        n_label = str(n_kv).removesuffix(".0")
        ax.annotate(n_label, (px, py), ha="center", va="center", fontsize=14, fontweight="bold",
                    color="#555555", zorder=12, path_effects=[pe.withStroke(linewidth=3, foreground="white")])
    # граница квартала / листа в фокусе (у разбросанных полигонов индекса совпадает с их контурами)
    if not group_mode:
        gpd.GeoSeries([focus_geom], crs=WORK_EPSG).boundary.plot(
            ax=ax, color="black", linewidth=2.4, zorder=10)
    ax.set_xlim(ext[0], ext[1]); ax.set_ylim(ext[2], ext[3])
    if basemap:
        try:
            import contextily as cx
            for _attempt in range(3):
                prov = pick_basemap_provider(skip=_attempt)
                if prov is None:
                    break
                try:
                    cx.add_basemap(ax, crs=f"EPSG:{WORK_EPSG}", source=prov,
                                   attribution=prov.get("attribution", "© OpenStreetMap"),
                                   attribution_size=10, zorder=0)
                    desaturate_last_basemap(ax)
                    break
                except Exception as e:     # таймаут/лимит запросов — пробуем следующий источник
                    log.warning("подложка %s недоступна (%s) — пробую другой источник", prov.get("name"), e)
        except Exception as e:
            log.warning("подложка OSM недоступна: %s", e)
    ax.legend(handles=handles, loc="upper left", bbox_to_anchor=(1.01, 1.0), fontsize=12, frameon=True,
              title="Легенда", title_fontsize=12)
    ax.set_title(f"{oozt_name}. {unit_labels(data.cfg)['group']} {kv}", fontsize=15, fontweight="bold")
    ax.tick_params(labelsize=12)
    try:
        set_degree_ticks(ax, ext)
    except Exception as e:                # запасной вариант — метры UTM
        log.warning("не удалось построить сетку в градусах (%s) — оси в метрах", e)
        ax.set_xlabel("E, м (UTM 37N)", fontsize=12); ax.set_ylabel("N, м", fontsize=12)
        ax.ticklabel_format(style="plain", useOffset=False)
    # масштабная линейка
    target = w / 4                        # «круглая» длина линейки: 1/2/5 × 10^k
    dec = 10 ** np.floor(np.log10(target))
    L = next(m * dec for m in (5, 2, 1) if m * dec <= target) if target >= dec else dec
    x0, y0 = ext[0] + w * 0.04, ext[2] + h * 0.04
    ax.plot([x0, x0 + L], [y0, y0], color="black", linewidth=4, zorder=14,
            path_effects=[pe.withStroke(linewidth=6.5, foreground="white")])
    ax.text(x0 + L / 2, y0 + h * 0.022, f"{int(L)} м", ha="center", va="bottom", fontsize=12, zorder=14,
            path_effects=[pe.withStroke(linewidth=3, foreground="white")])
    fig.tight_layout()
    ax.set_xlim(ext[0], ext[1]); ax.set_ylim(ext[2], ext[3])
    # подписи полигонов — последними, когда известен масштаб и оформление осей
    n_lab = _place_labels(ax, lab_items, sizes=(13, 12, 11, 10, 9, 8) if cfg.unit_mode != "eco" else (10, 9, 8, 7),
                          leaders=True)
    if n_lab < len(lab_items):
        ax.text(0.995, 0.005, f"подписаны {n_lab} из {len(lab_items)} полигонов; полный перечень — в таблице",
                transform=ax.transAxes, ha="right", va="bottom", fontsize=9, color="#444444", zorder=14,
                path_effects=[pe.withStroke(linewidth=3, foreground="white")])
    fig.savefig(out_png, bbox_inches="tight")
    plt.close(fig)
    return used_layers


def md_table(df: pd.DataFrame, cols: list[str]) -> str:
    def esc(x):
        s = "" if pd.isna(x) else str(x)
        return s.replace("|", "／").replace("\n", " ")
    lines = ["| " + " | ".join(cols) + " |", "|" + "|".join(["---"] * len(cols)) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(esc(r[c]) for c in cols) + " |")
    return "\n".join(lines)


def _report_groups_table(data: Data, kv_list) -> pd.DataFrame:
    """Список групп отчёта (кварталов / листов / групп) в порядке KV_SORT с учётом фильтра."""
    kv_all = (data.oozt[["OOZT", "KV", "KV_SORT"]].drop_duplicates(["OOZT", "KV"])
              .sort_values(["OOZT", "KV_SORT"]))
    if kv_list:
        want = {str(x).strip() for x in kv_list}
        kv_all = kv_all[kv_all["KV"].astype(str).isin(want)]
        miss = want - set(kv_all["KV"].astype(str))
        if miss:
            log.warning("В слое нет запрошенных для отчёта значений (%s): %s",
                        unit_labels(data.cfg)["groups_n"].lower(), ", ".join(sorted(miss)))
    return kv_all


def _report_rename(cfg: Config) -> dict:
    L = unit_labels(cfg)
    return {"Экосистема": L["unit_name"], "Мероприятие": "Мероприятия",
            "Тип экосистемы мероприятия": "Тип экосистемы", "Пункт Приказа": "Соответствие Пункту Приказа",
            "Выдел": L["unit"]}


def _fig_name(oozt, kv) -> str:
    safe = lambda x: re.sub(r"[^\w]+", "_", str(x))
    return "fig_" + safe(oozt) + "_kv_" + safe(kv) + ".png"


def _params_line() -> str:
    return ("Параметры каждого мероприятия (правило, расстояния, окна, квантили, пороги) заданы в таблице "
            "мероприятий и перечислены в разделе «Как рассчитаны мероприятия»; верхний квартиль считается по всем "
            "ООЗТ; общий порог доли полигона — %g%% (если в таблице не указан свой)." % _MIN_SHARE[0])


def _summary_table(table: pd.DataFrame, L: dict) -> pd.DataFrame:
    return table.groupby(["Тип мероприятий", "Мероприятие"]).agg(
        **{L["units_n"]: ("Выдел", "size"), L["groups_n"]: ("Квартал", "nunique")}).reset_index()


REPORT_FORMATS = {"docx": "Word (DOCX)", "pdf": "PDF", "doc": "Word 97–2003 (DOC)"}
PAGE_W_CM, PAGE_H_CM, MARGIN_CM = 29.7, 21.0, 1.5          # А4 альбомная
TEXT_W_CM = PAGE_W_CM - 2 * MARGIN_CM


def rules_table(measures, table: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    """Как рассчитано каждое включённое мероприятие: метод из таблицы → правило скрипта → число полигонов.
    Позволяет сверить, что название мероприятия и основание назначения соответствуют друг другу."""
    _MIN_SHARE[0] = float(getattr(cfg, "min_share_pct", 0) or 0)
    rows = []
    cnt = table["Мероприятие"].value_counts() if not table.empty else pd.Series(dtype=int)
    for m in measures or []:
        if not m.get("on"):
            continue
        n = int(cnt.get(m["name"], 0))
        extra = [k for k in cnt.index if str(k).startswith(m["name"] + " — ")]
        note = ""
        if extra:
            note = "; " + "; ".join(f"{str(k)[len(m['name']) + 3:]}: {int(cnt[k])}" for k in extra)
        st = measure_status(m, cfg)
        rows.append({
            "Код": m.get("code", ""),
            "Мероприятие": m["name"],
            "Пункт Приказа": m.get("prikaz", ""),
            "Тип экосистемы (по таблице)": m.get("eco", ""),
            "Метод выявления (из таблицы)": m.get("method") or "",
            "Правило": (RULES[st["rule"]]["title"] if st["rule"] else "—") + (
                " (по тексту метода)" if st["spec"].get("origin") == "текст метода" and st["rule"] else ""),
            "Как рассчитано скриптом": describe_rule(m),
            "Назначено полигонов": f"{n}{note}" if st["state"] != "missing" else f"0 ({st['text']})",
        })
    return pd.DataFrame(rows)


def execution_list(table: pd.DataFrame, cfg: Config, data: "Data | None" = None) -> tuple[pd.DataFrame, str]:
    """Перечень мероприятий к выполнению: одна строка на полигон, мероприятия с пунктами Приказа списком.
    Порядок — по возрастанию поля группировки (индекса), затем ID; для листов сетки и кварталов — см. ниже.
    Возвращает (таблица с подписями колонок, описание порядка)."""
    L = unit_labels(cfg)
    if table is None or table.empty:
        return pd.DataFrame(), ""
    t = table.copy()
    def _with_clause(m, pk):
        pk = "" if pd.isna(pk) else str(pk).strip().rstrip(".")
        if pk.lower() in ("", "nan", "нет", "-", "—"):
            return str(m)
        return f"{m} (п. {pk})" if re.match(r"^\d", pk) else f"{m} ({pk})"
    t["_m"] = [_with_clause(m, pk) for pk, m in zip(t["Пункт Приказа"], t["Мероприятие"])]
    t["_tk"] = t["Тип мероприятий"].astype(str)
    t = t.sort_values(["_tk", "Пункт Приказа", "Мероприятие"], key=_nat_sort_key("Пункт Приказа"))
    keys = ["ООЗТ", "Квартал", "Выдел"]
    lst = t.groupby(keys, as_index=False, sort=False, dropna=False).agg(
        **{"Экосистема": ("Экосистема", "first"), "Мероприятия": ("_m", "\n".join),
           "Число": ("Мероприятие", "size")})
    grid = cfg.unit_mode == "eco" and not cfg.report_group_field
    same = bool(cfg.unit_mode == "eco" and cfg.report_group_field and data is not None
                and cfg.report_group_field == getattr(data, "F_ID", None))
    order = ["Выдел", "Квартал"] if grid else ["Квартал", "Выдел"]
    lst = lst.sort_values(["ООЗТ"] + order, key=_nat_sort_key("Квартал", "Выдел")).reset_index(drop=True)
    cols = (["Выдел", "Квартал"] if grid else ["Квартал"] + ([] if same else ["Выдел"])) + \
           ["Экосистема", "Мероприятия", "Число"]
    if lst["ООЗТ"].nunique() > 1:
        cols = ["ООЗТ"] + cols
    out = lst[cols].rename(columns={"Квартал": L["group"], "Выдел": L["unit"], "Экосистема": L["unit_name"],
                                    "Мероприятия": "Мероприятия (пункт Приказа)", "Число": "Число мероприятий"})
    if grid:
        how = f"по возрастанию «{L['unit']}»"
    elif same:
        how = f"по возрастанию «{L['group']}»"
    else:
        how = f"по возрастанию «{L['group']}», внутри — по «{L['unit']}»"
    return out, how


def _report_blocks(data: Data, table: pd.DataFrame, evidence, rep: Path, kv_all: pd.DataFrame,
                   basemap=True, draw=True) -> list:
    """Содержание отчёта в нейтральном виде; из него собираются MD, DOCX и PDF.
    Блоки: ("title", t) ("h", level, t) ("p", t) ("note", t) ("img", path) ("break",)
           ("table", header, rows, widths_cm, font_pt)"""
    cfg = data.cfg
    L = unit_labels(cfg)
    mode_txt = UNIT_MODES[cfg.unit_mode]
    if cfg.unit_mode == "eco" and not cfg.report_group_field:
        mode_txt += f"; листы сетки {cfg.sheet_size_m:.0f} × {cfg.sheet_size_m:.0f} м"
    n_units = table[["ООЗТ", "Квартал", "Выдел"]].drop_duplicates().shape[0] if not table.empty else 0
    B = [("title", f"Природоохранные мероприятия по {L['group_pl']} ООЗТ"),
         ("p", f"Слой территории: {Path(cfg.oozt_path).name} ({mode_txt}); "
               f"таблица мероприятий: {Path(cfg.measures_xlsx).name}."),
         ("p", f"Всего назначений: {len(table)}; {L['units_gen']} с мероприятиями: {n_units}; "
               f"{L['groups_n'].lower()} в отчёте: {len(kv_all)}."),
         ("p", _params_line())]
    s = lambda v: "" if pd.isna(v) else str(v)
    if not table.empty:
        summ = _summary_table(table, L)
        hdr = ["Тип мероприятий", "Мероприятие", L["units_n"], L["groups_n"]]
        B += [("h", 1, "Сводка по мероприятиям"),
              ("table", hdr, [[s(r[c]) for c in hdr] for _, r in summ.iterrows()], [5.0, 15.7, 3.0, 3.0], 9)]
    rt = rules_table(getattr(data, "measures", None), table, cfg)
    if not rt.empty:
        hdr2 = ["Код", "Мероприятие", "Метод выявления (из таблицы)", "Правило", "Как рассчитано скриптом",
                "Назначено полигонов"]
        B += [("h", 1, "Как рассчитаны мероприятия"),
              ("p", "Для каждого включённого мероприятия указаны метод выявления из таблицы мероприятий, "
                    "правило назначения с параметрами (из столбцов «Правило назначения», «Типы полигонов», "
                    "«Виды: …», «Расстояние, м» и др.; если правило не задано — по тексту метода) и число "
                    "полигонов, которым назначено мероприятие."),
              ("table", hdr2, [[s(r[c]) for c in hdr2] for _, r in rt.iterrows()],
               [1.3, 4.6, 5.0, 2.8, 7.6, 2.4], 8)]
    report_table = table.rename(columns=_report_rename(cfg))
    group_cols = ["Тип мероприятий", "Тип экосистемы", "Соответствие Пункту Приказа"]
    cols = [L["unit"], L["unit_name"], "Мероприятия", "Основание"]
    w0 = 1.8 if cfg.unit_mode == "eco" else 1.3
    widths = [w0, 4.6, 7.4, TEXT_W_CM - w0 - 4.6 - 7.4]
    for _, kr in kv_all.iterrows():
        oozt, kv = kr["OOZT"], kr["KV"]
        png = rep / _fig_name(oozt, kv)
        if draw:
            log.info("Схема: %s, %s %s", oozt, L["group"].lower(), kv)
            draw_kvartal(data, evidence, kv, oozt, png, basemap=basemap)
        sub = report_table[
            (report_table["ООЗТ"] == oozt) & (report_table["Квартал"].astype(str) == str(kv))
        ] if not report_table.empty else report_table
        vyd = data.oozt[(data.oozt["KV"].astype(str) == str(kv)) & (data.oozt["OOZT"] == oozt)]
        B += [("break",), ("h", 1, f"{oozt}. {L['group']} {kv}")]
        if png.exists():
            B.append(("img", png))
        B.append(("p", f"{L['units_n']}: {len(vyd)}; площадь {vyd['area_ha'].sum():.1f} га; "
                       f"назначено мероприятий: {len(sub)}."))
        if sub.empty:
            B.append(("note", "Мероприятия по загруженным данным не назначены."))
            continue
        B.append(("h", 2, "Мероприятия по приказу"))
        grouped = sub.sort_values(group_cols + [L["unit"]], key=_nat_sort_key(L["unit"], "Соответствие Пункту Приказа")
                                  ).groupby(group_cols, sort=False, dropna=False)
        for (measure_type, ecosystem_type, order_clause), group in grouped:
            B += [("h", 3, f"Тип мероприятий: {s(measure_type)}"),
                  ("h", 4, f"Тип экосистемы: {s(ecosystem_type)}"),
                  ("h", 5, f"Соответствие Пункту Приказа: {s(order_clause)}"),
                  ("table", cols, [[s(r[c]) for c in cols] for _, r in group.sort_values([L["unit"]], key=_nat_sort_key(L["unit"])).iterrows()],
                   widths, 9)]
    # перечень мероприятий к выполнению (по выбранным в отчёт кварталам / листам / индексам)
    if not table.empty and len(kv_all):
        sel = table.merge(kv_all[["OOZT", "KV"]].rename(columns={"OOZT": "ООЗТ", "KV": "Квартал"}),
                          on=["ООЗТ", "Квартал"], how="inner") if len(kv_all) else table
        ex, how = execution_list(sel, cfg, data)
        if not ex.empty:
            hdr = list(ex.columns)
            ex = ex.rename(columns={"Число мероприятий": "Число"})
            fixed = {L["group"]: 2.4, L["unit"]: 2.2, L["unit_name"]: 5.6, "Число": 1.5, "ООЗТ": 2.6}
            wd = [fixed.get(c, 0) for c in hdr]
            rest = TEXT_W_CM - sum(wd)
            wd = [w if w else rest for w in wd]
            B += [("break",), ("h", 1, "Перечень мероприятий к выполнению"),
                  ("p", f"Одна строка — один полигон; порядок — {how}. Полигонов: {len(ex)}; "
                        f"мероприятий: {int(ex['Число'].sum())}."),
                  ("table", hdr, [[s(v) for v in r] for r in ex.itertuples(index=False)], wd, 9)]
    return B


# ------------------------------------------------------------------ Markdown
def _write_md(blocks: list, path: Path):
    def esc(x):
        return str(x).replace("|", "／").replace("\n", " ")
    out = []
    for b in blocks:
        k = b[0]
        if k == "title":
            out.append(f"# {b[1]}\n")
        elif k == "h":
            out.append(f"{'#' * (b[1] + 1)} {b[2]}\n")
        elif k == "p":
            out.append(b[1] + "\n")
        elif k == "note":
            out.append(f"_{b[1]}_\n")
        elif k == "img":
            try:
                rel = b[1].resolve().relative_to(path.parent.resolve()).as_posix()
            except ValueError:
                rel = Path(os.path.relpath(b[1].resolve(), path.parent.resolve())).as_posix()
            out.append(f"![{b[1].stem}]({rel})\n")
        elif k == "table":
            hdr, rows = b[1], b[2]
            t = ["| " + " | ".join(hdr) + " |", "|" + "|".join(["---"] * len(hdr)) + "|"]
            t += ["| " + " | ".join(esc(v) for v in r) + " |" for r in rows]
            out.append("\n".join(t) + "\n")
    path.write_text("\n".join(out), encoding="utf-8")


# ------------------------------------------------------------------ DOCX (без внешних библиотек)
_W_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
         'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
         'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
         'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"')
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _x(t) -> str:
    from xml.sax.saxutils import escape
    return escape(_BAD_XML.sub("", str(t)), {'"': "&quot;"})


def _png_size(path: Path) -> tuple[int, int]:
    import struct
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        return struct.unpack(">II", head[16:24])
    from PIL import Image
    with Image.open(path) as im:
        return im.size


def _docx_styles() -> str:
    def para(sid, name, sz, bold=False, ital=False, before=0, after=120, lvl=None, color=None, keep=False):
        ppr = ("<w:keepNext/>" if keep else "") + f'<w:spacing w:before="{before}" w:after="{after}"/>'
        if lvl is not None:
            ppr += f'<w:outlineLvl w:val="{lvl}"/>'
        rpr = ("<w:b/>" if bold else "") + ("<w:i/>" if ital else "") + \
              (f'<w:color w:val="{color}"/>' if color else "") + f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/>'
        return (f'<w:style w:type="paragraph" w:styleId="{sid}"><w:name w:val="{name}"/>'
                f'<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/>'
                f'<w:pPr>{ppr}</w:pPr><w:rPr>{rpr}</w:rPr></w:style>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:styles {_W_NS}>'
        '<w:docDefaults><w:rPrDefault><w:rPr>'
        '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman" w:cs="Times New Roman"/>'
        '<w:sz w:val="22"/><w:szCs w:val="22"/><w:lang w:val="ru-RU"/></w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr><w:spacing w:after="120" w:line="252" w:lineRule="auto"/></w:pPr></w:pPrDefault>'
        '</w:docDefaults>'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>'
        + para("Title", "Title", 34, bold=True, after=200, color="1F3B4D")
        + para("Heading1", "heading 1", 28, bold=True, before=120, after=120, lvl=0, color="1F3B4D", keep=True)
        + para("Heading2", "heading 2", 25, bold=True, before=160, after=80, lvl=1, color="1F3B4D", keep=True)
        + para("Heading3", "heading 3", 23, bold=True, before=160, after=40, lvl=2, keep=True)
        + para("Heading4", "heading 4", 22, bold=False, ital=True, before=0, after=40, lvl=3, keep=True)
        + para("Heading5", "heading 5", 22, bold=False, ital=True, before=0, after=80, lvl=4, color="44546A", keep=True)
        + para("Note", "Note", 22, ital=True, color="5F6368")
        + '<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/>'
          '<w:tblPr><w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/>'
          '<w:left w:w="85" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="85" w:type="dxa"/>'
          '</w:tblCellMar></w:tblPr></w:style>'
          '<w:style w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:basedOn w:val="TableNormal"/>'
          '<w:tblPr><w:tblBorders>'
        + "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="8C9AA5"/>'
                  for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
        + '</w:tblBorders></w:tblPr></w:style></w:styles>')


def _write_docx(blocks: list, path: Path, title: str):
    import zipfile, datetime
    tw = lambda cm: int(round(cm * 567))          # твипы
    emu = lambda cm: int(round(cm * 360000))
    body, rels, media = [], [], []

    def run(t, bold=False, sz=None, ital=False):
        rpr = ("<w:b/>" if bold else "") + ("<w:i/>" if ital else "") + \
              (f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/>' if sz else "")
        parts = str(t).split("\n")
        inner = "<w:br/>".join(f'<w:t xml:space="preserve">{_x(p)}</w:t>' for p in parts)
        return f"<w:r>{'<w:rPr>' + rpr + '</w:rPr>' if rpr else ''}{inner}</w:r>"

    def p(t, style=None, extra_ppr=""):
        ppr = (f'<w:pStyle w:val="{style}"/>' if style else "") + extra_ppr
        return f"<w:p>{'<w:pPr>' + ppr + '</w:pPr>' if ppr else ''}{run(t)}</w:p>"

    pending_break = False
    for b in blocks:
        k = b[0]
        if k == "break":
            pending_break = True
            continue
        brk = "<w:pageBreakBefore/>" if pending_break else ""
        pending_break = False
        if k == "title":
            body.append(p(b[1], "Title"))
        elif k == "h":
            body.append(p(b[2], f"Heading{min(b[1], 5)}", brk))
        elif k == "p":
            body.append(p(b[1], None, brk))
        elif k == "note":
            body.append(p(b[1], "Note", brk))
        elif k == "img":
            src: Path = b[1]
            n = len(media) + 1
            rid = f"rIdImg{n}"
            media.append((f"media/image{n}.png", src))
            rels.append(f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                        f'relationships/image" Target="media/image{n}.png"/>')
            wpx, hpx = _png_size(src)
            w_cm = TEXT_W_CM - 0.7
            h_cm = w_cm * hpx / max(wpx, 1)
            max_h = PAGE_H_CM - 2 * MARGIN_CM - 3.2
            if h_cm > max_h:
                w_cm, h_cm = w_cm * max_h / h_cm, max_h
            cx, cy = emu(w_cm), emu(h_cm)
            body.append(
                f'<w:p><w:pPr><w:keepNext/>{brk}<w:spacing w:after="60"/><w:jc w:val="center"/></w:pPr><w:r><w:drawing>'
                f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cx}" cy="{cy}"/>'
                f'<wp:docPr id="{n}" name="Схема {n}"/><wp:cNvGraphicFramePr>'
                f'<a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
                f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
                f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{n}" name="{_x(src.name)}"/><pic:cNvPicPr/></pic:nvPicPr>'
                f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
                f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
                f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
                f'</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>')
        elif k == "table":
            hdr, rows, widths, fpt = b[1], b[2], b[3], b[4]
            sz = int(fpt * 2)
            ws = [tw(w) for w in widths]

            def cell(t, w, head):
                shd = '<w:shd w:val="clear" w:color="auto" w:fill="E3EAEE"/>' if head else ""
                return (f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/>{shd}</w:tcPr>'
                        f'<w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
                        f'{run(t, bold=head, sz=sz)}</w:p></w:tc>')
            t = [f'<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="{sum(ws)}" w:type="dxa"/>'
                 f'<w:tblLayout w:type="fixed"/><w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" '
                 f'w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/></w:tblPr><w:tblGrid>'
                 + "".join(f'<w:gridCol w:w="{w}"/>' for w in ws) + "</w:tblGrid>",
                 '<w:tr><w:trPr><w:cantSplit/><w:tblHeader/></w:trPr>'
                 + "".join(cell(h, w, True) for h, w in zip(hdr, ws)) + "</w:tr>"]
            for r in rows:
                t.append("<w:tr><w:trPr><w:cantSplit/></w:trPr>"
                         + "".join(cell(v, w, False) for v, w in zip(r, ws)) + "</w:tr>")
            t.append("</w:tbl>")
            if brk:
                body.append(f"<w:p><w:pPr>{brk}<w:spacing w:after='0'/></w:pPr></w:p>".replace("'", '"'))
            body.append("".join(t))
            body.append('<w:p><w:pPr><w:spacing w:after="60"/></w:pPr></w:p>')
    m = tw(MARGIN_CM)
    sect = (f'<w:sectPr><w:footerReference w:type="default" r:id="rIdFtr1"/>'
            f'<w:pgSz w:w="{tw(PAGE_W_CM)}" w:h="{tw(PAGE_H_CM)}" w:orient="landscape"/>'
            f'<w:pgMar w:top="{m}" w:right="{m}" w:bottom="{m}" w:left="{m}" w:header="425" w:footer="425" '
            f'w:gutter="0"/></w:sectPr>')
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f'<w:document {_W_NS}><w:body>{"".join(body)}{sect}</w:body></w:document>')
    footer = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              f'<w:ftr {_W_NS}><w:p><w:pPr><w:jc w:val="center"/></w:pPr>'
              '<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r>'
              '<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
              '<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:fldChar w:fldCharType="separate"/></w:r>'
              '<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:t>1</w:t></w:r>'
              '<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:fldChar w:fldCharType="end"/></w:r></w:p></w:ftr>')
    rels_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/styles" Target="styles.xml"/>'
                '<Relationship Id="rIdFtr1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/footer" Target="footer1.xml"/>' + "".join(rels) + "</Relationships>")
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{_x(title)}</dc:title><dc:creator>oozt_recs.py</dc:creator><dc:language>ru-RU</dc:language>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified></cp:coreProperties>')
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
              '<Default Extension="xml" ContentType="application/xml"/>'
              '<Default Extension="png" ContentType="image/png"/>'
              '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
              'wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.'
              'wordprocessingml.styles+xml"/>'
              '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.'
              'wordprocessingml.footer+xml"/>'
              '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.'
              'core-properties+xml"/></Types>')
    pkg_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/officeDocument" Target="word/document.xml"/>'
                '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/'
                'metadata/core-properties" Target="docProps/core.xml"/></Relationships>')
    tmp = path.with_name(path.stem + ".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ctypes)
        z.writestr("_rels/.rels", pkg_rels)
        z.writestr("docProps/core.xml", core)
        z.writestr("word/document.xml", document)
        z.writestr("word/styles.xml", _docx_styles())
        z.writestr("word/footer1.xml", footer)
        z.writestr("word/_rels/document.xml.rels", rels_xml)
        for name, src in media:
            z.write(src, "word/" + name, compress_type=zipfile.ZIP_STORED)
    os.replace(tmp, path)


# ------------------------------------------------------------------ PDF / DOC через Word или LibreOffice
def _convert_with_word(docx: Path, targets: dict) -> list[str]:
    """targets: {"pdf": Path, "doc": Path}. Возвращает список созданных форматов."""
    import win32com.client  # noqa: F401  (pywin32; есть в conda-окружениях Windows)
    import pythoncom
    fmt = {"pdf": 17, "doc": 0}           # wdFormatPDF, wdFormatDocument (Word 97–2003)
    done = []
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        d = word.Documents.Open(str(docx.resolve()), ConfirmConversions=False, ReadOnly=True,
                                AddToRecentFiles=False, Visible=False)
        try:
            d.Fields.Update()
            for k, p in targets.items():
                d.SaveAs2(str(p.resolve()), FileFormat=fmt[k], AddToRecentFiles=False)
                done.append(k)
        finally:
            d.Close(False)
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()
    return done


def _find_soffice() -> str | None:
    import shutil
    for n in ("soffice", "libreoffice"):
        p = shutil.which(n)
        if p:
            return p
    for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                 os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        p = Path(base) / "LibreOffice" / "program" / "soffice.exe"
        if p.exists():
            return str(p)
    return None


def _convert_with_soffice(soffice: str, docx: Path, targets: dict) -> list[str]:
    import subprocess, tempfile
    done = []
    for k, p in targets.items():
        with tempfile.TemporaryDirectory() as td:
            subprocess.run([soffice, "--headless", f"-env:UserInstallation=file:///{Path(td).as_posix()}/u",
                            "--convert-to", k, "--outdir", td, str(docx)],
                           check=True, timeout=900, capture_output=True)
            res = Path(td) / (docx.stem + "." + k)
            if not res.exists():
                raise RuntimeError(f"LibreOffice не создал {res.name}")
            shutil_move(res, p)
            done.append(k)
    return done


def shutil_move(src: Path, dst: Path):
    import shutil
    if dst.exists():
        dst.unlink()
    shutil.move(str(src), str(dst))


# ------------------------------------------------------------------ PDF собственной вёрсткой (matplotlib)
def _write_pdf_mpl(blocks: list, path: Path, title: str):
    """Запасной путь, если нет Word и LibreOffice: та же структура, страницы А4 альбомные."""
    enable_win_dll_path()
    import matplotlib.image as mpimg
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.patches import Rectangle
    PT = 72 / 2.54
    W, H, M = PAGE_W_CM * PT, PAGE_H_CM * PT, MARGIN_CM * PT
    TW = W - 2 * M
    from matplotlib.font_manager import FontProperties, findfont, get_font
    fonts = {b: get_font(findfont(FontProperties(family="DejaVu Sans", weight="bold" if b else "normal")))
             for b in (False, True)}
    _wcache = {}

    def measure(t, size, bold):
        key = (t, size, bold)
        if key not in _wcache:
            f = fonts[bold]
            f.set_size(size, 72)
            f.set_text(t, 0.0)
            _wcache[key] = f.get_width_height()[0] / 64.0 * 1.02
        return _wcache[key]
    st = {"title": (17, True), 1: (14, True), 2: (12.5, True), 3: (11.5, True), 4: (11, False), 5: (11, False),
          "p": (10.5, False), "note": (10.5, False)}
    matplotlib.rcParams["pdf.fonttype"] = 42          # шрифты TrueType: текст копируется и ищется
    pdf = PdfPages(path, metadata={"Title": title, "Creator": "oozt_recs.py"})
    state = {"fig": None, "y": 0.0, "page": 0}

    def wrap(t, size, bold, width_pt):
        out = []
        for para in str(t).split("\n"):
            words = re.split(r"(?<=-)|\s+", para.strip())
            line = ""
            for w in [w for w in words if w]:
                cand = w if not line else (line + w if line.endswith("-") else line + " " + w)
                if measure(cand, size, bold) <= width_pt:
                    line = cand
                    continue
                if line:
                    out.append(line)
                while measure(w, size, bold) > width_pt and len(w) > 1:   # слово длиннее колонки
                    k = len(w) - 1
                    while k > 1 and measure(w[:k] + "-", size, bold) > width_pt:
                        k -= 1
                    out.append(w[:k] + "-")
                    w = w[k:]
                line = w
            out.append(line)
        return out

    def new_page():
        if state["fig"] is not None:
            finish()
        fig = plt.figure(figsize=(W / 72, H / 72))
        state.update(fig=fig, y=H - M, page=state["page"] + 1)

    def finish():
        fig = state["fig"]
        fig.text(0.5, (M * 0.45) / H, str(state["page"]), ha="center", va="center", fontsize=8, color="#5F6368")
        pdf.savefig(fig)
        plt.close(fig)
        state["fig"] = None

    def txt(x, y, t, size, bold=False, color="#202124", style="normal"):
        state["fig"].text(x / W, y / H, t, fontsize=size, fontweight="bold" if bold else "normal",
                          color=color, va="top", ha="left", fontstyle=style)

    def ensure(h):
        if state["fig"] is None or state["y"] - h < M:
            new_page()

    def para(t, size, bold, color="#202124", style="normal", before=0, after=4, keep=0):
        lines = wrap(t, size, bold, TW)
        lh = size * 1.3
        ensure(before + lh * min(len(lines), 2) + keep)
        state["y"] -= before
        for ln in lines:
            ensure(lh)
            txt(M, state["y"], ln, size, bold, color, style)
            state["y"] -= lh
        state["y"] -= after

    def table(hdr, rows, widths, fpt):
        ws = [w / sum(widths) * TW for w in widths]
        lh, pad = fpt * 1.25, 2.5

        def row_lines(r, bold):
            return [wrap(v, fpt, bold, w - 2 * pad - 1) for v, w in zip(r, ws)]

        def draw_row(cells, bold, fill=None):
            n = max(len(c) for c in cells)
            h = n * lh + 2 * pad
            x, y = M, state["y"]
            for c, w in zip(cells, ws):
                state["fig"].add_artist(Rectangle((x / W, (y - h) / H), w / W, h / H, transform=state["fig"].transFigure,
                                                  facecolor=fill or "none", edgecolor="#8C9AA5", linewidth=0.5))
                for i, ln in enumerate(c):
                    txt(x + pad, y - pad - i * lh, ln, fpt, bold)
                x += w
            state["y"] -= h

        head = row_lines(hdr, True)
        hh = max(len(c) for c in head) * lh + 2 * pad
        ensure(hh + 3 * lh)
        draw_row(head, True, "#E3EAEE")
        for r in rows:
            cells = row_lines(r, False)
            while cells:
                h = max(len(c) for c in cells) * lh + 2 * pad
                if state["y"] - h < M:
                    room = int((state["y"] - M - 2 * pad) / lh)
                    if room >= 3:                       # длинную строку делим между страницами
                        draw_row([c[:room] for c in cells], False)
                        cells = [c[room:] for c in cells]
                        if not any(cells):
                            break
                    new_page()
                    draw_row(head, True, "#E3EAEE")
                    continue
                draw_row(cells, False)
                break
        state["y"] -= 6

    def img_size(path):
        img = mpimg.imread(str(path))
        ih, iw = img.shape[0], img.shape[1]
        w = TW - 20
        h = w * ih / iw
        max_h = H - 2 * M - 90
        if h > max_h:
            w, h = w * max_h / h, max_h
        return img, w, h

    pending_break = False
    for i, b in enumerate(blocks):
        k = b[0]
        if k == "break":
            pending_break = True
            continue
        if pending_break:
            new_page()
            pending_break = False
        if k == "title":
            para(b[1], *st["title"], color="#1F3B4D", after=10)
        elif k == "h":
            size, bold = st[min(b[1], 5)]
            keep = {1: 60, 2: 140, 3: 110, 4: 85, 5: 60}[min(b[1], 5)]
            if i + 1 < len(blocks) and blocks[i + 1][0] == "img":      # заголовок не отрываем от схемы
                keep = img_size(blocks[i + 1][1])[2] + 8
            para(b[2], size, bold, color="#1F3B4D" if b[1] <= 2 else "#202124",
                 style="italic" if b[1] >= 4 else "normal", before=8 if b[1] <= 3 else 0, after=3, keep=keep)
        elif k == "p":
            para(b[1], *st["p"])
        elif k == "note":
            para(b[1], *st["note"], color="#5F6368", style="italic")
        elif k == "img":
            img, w, h = img_size(b[1])
            ensure(h + 6)
            x0 = M + (TW - w) / 2
            ax = state["fig"].add_axes([x0 / W, (state["y"] - h) / H, w / W, h / H])
            ax.imshow(img, interpolation="lanczos")
            ax.set_axis_off()
            state["y"] -= h + 6
        elif k == "table":
            table(b[1], b[2], b[3], b[4])
    if state["fig"] is None:
        new_page()
    finish()
    pdf.close()


def _write_report_files(blocks: list, rep: Path, title: str, formats) -> dict:
    """Записывает report.md и выбранные форматы; возвращает {формат: путь}."""
    formats = [f for f in (formats or []) if f in REPORT_FORMATS]
    made = {}
    _write_md(blocks, rep / "report.md")
    made["md"] = rep / "report.md"
    if not formats:
        return made
    docx = rep / "report.docx"
    try:
        _write_docx(blocks, docx, title)
        if "docx" in formats:
            made["docx"] = docx
            log.info("Отчёт DOCX: %s", docx)
    except PermissionError:
        log.warning("report.docx открыт в другой программе — закройте его и пересоберите отчёт (--report-only)")
        return made
    targets = {k: rep / f"report.{k}" for k in ("pdf", "doc") if k in formats}
    methods = []
    if os.name == "nt":
        methods.append(("Microsoft Word", lambda t: _convert_with_word(docx, t)))
    so = _find_soffice()
    if so:
        methods.append(("LibreOffice", lambda t: _convert_with_soffice(so, docx, t)))
    for name, fn in methods:
        if not targets:
            break
        try:
            for k in fn(dict(targets)):
                made[k] = targets.pop(k)
                log.info("Отчёт %s (%s): %s", k.upper(), name, made[k])
        except ImportError:
            log.info("%s: нет модуля pywin32 — пропускаю", name)
        except Exception as e:
            log.info("%s недоступен для конвертации (%s)", name, str(e).splitlines()[0][:200] if str(e) else type(e).__name__)
    if "pdf" in targets:
        try:
            _write_pdf_mpl(blocks, targets["pdf"], title)
            made["pdf"] = targets.pop("pdf")
            log.info("Отчёт PDF (встроенная вёрстка): %s", made["pdf"])
        except PermissionError:
            log.warning("report.pdf открыт в другой программе — закройте его и пересоберите отчёт (--report-only)")
        except Exception as e:
            log.warning("PDF не создан: %s", e)
    if "doc" in targets:
        log.warning("Формат DOC (Word 97–2003) требует установленного Microsoft Word или LibreOffice — не создан")
    if "docx" not in formats and docx.exists() and "docx" not in made:
        try:
            docx.unlink()                       # промежуточный файл, если DOCX не заказан
        except OSError:
            pass
    return made


def build_report(data: Data, table: pd.DataFrame, evidence, out_dir: Path, kv_list, basemap=True,
                 formats=("docx", "pdf")):
    enable_win_dll_path()
    cfg = data.cfg
    rep = out_dir / "report"
    rep.mkdir(parents=True, exist_ok=True)
    kv_all = _report_groups_table(data, kv_list)
    blocks = _report_blocks(data, table, evidence, rep, kv_all, basemap=basemap)
    title = blocks[0][1]
    made = _write_report_files(blocks, rep, title, formats)
    return made.get("pdf") or made.get("docx") or made["md"]


# =====================================================================================
#                                 ПРОЕКТ QGIS
# =====================================================================================
QGIS_PROJECT_NAME = "oozt_measures.qgs"
_MEASURE_COLORS = ["#E6194B", "#3CB44B", "#4363D8", "#F58231", "#911EB4", "#42D4F4", "#F032E6", "#BFEF45",
                   "#469990", "#9A6324", "#800000", "#808000", "#000075", "#DCBEFF", "#FFE119", "#AAFFC3",
                   "#FABED4", "#A9A9A9"]


def _rgba(hex_color: str, alpha: float = 1.0) -> str:
    h = hex_color.lstrip("#")
    return f"{int(h[0:2], 16)},{int(h[2:4], 16)},{int(h[4:6], 16)},{int(round(alpha * 255))}"


def _q_opts(d: dict) -> str:
    """Свойства слоя символа: формат <Option> (QGIS ≥ 3.26) и <prop> (более ранние версии)."""
    opt = "".join(f'<Option name="{_x(k)}" type="QString" value="{_x(v)}"/>' for k, v in d.items())
    prop = "".join(f'<prop k="{_x(k)}" v="{_x(v)}"/>' for k, v in d.items())
    return f'<Option type="Map">{opt}</Option>{prop}'


def _q_symbol(kind: str, name: str, color: str, alpha=1.0, outline="#333333", width=None, size=None,
              outline_width=0.26) -> str:
    if kind == "fill":
        lay = ("SimpleFill", {"color": _rgba(color, alpha), "outline_color": _rgba(outline, 1.0),
                              "outline_style": "solid", "outline_width": str(outline_width),
                              "outline_width_unit": "MM", "style": "solid" if alpha > 0 else "no",
                              "joinstyle": "bevel"})
    elif kind == "line":
        lay = ("SimpleLine", {"line_color": _rgba(color, alpha), "line_width": str(width or 0.6),
                              "line_width_unit": "MM", "line_style": "solid", "capstyle": "round",
                              "joinstyle": "round"})
    else:
        lay = ("SimpleMarker", {"name": "circle", "color": _rgba(color, alpha), "outline_color": "255,255,255,255",
                                "outline_style": "solid", "outline_width": "0.3", "outline_width_unit": "MM",
                                "size": str(size or 2.6), "size_unit": "MM"})
    return (f'<symbol type="{kind}" name="{_x(name)}" alpha="1" clip_to_extent="1" force_rhr="0">'
            f'<layer class="{lay[0]}" enabled="1" locked="0" pass="0">{_q_opts(lay[1])}</layer></symbol>')


def _q_single(sym: str) -> str:
    return (f'<renderer-v2 type="singleSymbol" forceraster="0" symbollevels="0" enableorderby="0">'
            f'<symbols>{sym}</symbols></renderer-v2>')


def _q_categorized(attr: str, cats: list) -> str:
    """cats: [(значение, подпись, символ_xml_с_именем_i)]"""
    c = "".join(f'<category render="true" symbol="{i}" value="{_x(v)}" label="{_x(lbl)}"/>'
                for i, (v, lbl, _) in enumerate(cats))
    s_ = "".join(sym for _, _, sym in cats)
    return (f'<renderer-v2 type="categorizedSymbol" attr="{_x(attr)}" forceraster="0" symbollevels="0" '
            f'enableorderby="0"><categories>{c}</categories><symbols>{s_}</symbols></renderer-v2>')


def _q_labels(field_name: str, max_scale=20000, size=8) -> str:
    return (f'<labeling type="simple"><settings calloutType="simple">'
            f'<text-style fieldName="{_x(field_name)}" isExpression="0" fontFamily="Arial" fontSize="{size}" '
            f'fontSizeUnit="Point" fontWeight="75" fontItalic="0" textColor="20,20,20,255" textOpacity="1" '
            f'multilineHeight="1" blendMode="0" previewBkgrdColor="255,255,255,255">'
            f'<text-buffer bufferDraw="1" bufferSize="0.8" bufferSizeUnits="MM" bufferColor="255,255,255,255" '
            f'bufferOpacity="1" bufferJoinStyle="128" bufferNoFill="1" bufferBlendMode="0"/>'
            f'</text-style><text-format wrapChar="" multilineAlign="1" plussign="0" decimals="3" formatNumbers="0"/>'
            f'<placement placement="1" centroidInside="1" centroidWhole="0" fitInPolygonOnly="0" priority="5" '
            f'offsetType="0" dist="0" overrunDistance="0"/>'
            f'<rendering scaleVisibility="1" scaleMin="1" scaleMax="{max_scale}" obstacle="1" '
            f'displayAll="0" upsidedownLabels="0" labelPerPart="0" mergeLines="0" minFeatureSize="0" '
            f'limitNumLabels="0" maxNumLabels="2000" drawLabels="1" fontLimitPixelSize="0"/>'
            f'</settings></labeling>')


def _q_srs(epsg: int) -> str:
    from pyproj import CRS
    c = CRS.from_epsg(epsg)
    return (f'<spatialrefsys nativeFormat="Wkt"><wkt>{_x(c.to_wkt())}</wkt><proj4>{_x(c.to_proj4())}</proj4>'
            f'<srid>{epsg}</srid><authid>EPSG:{epsg}</authid><description>{_x(c.name)}</description>'
            f'<projectionacronym></projectionacronym><ellipsoidacronym></ellipsoidacronym>'
            f'<geographicflag>{"true" if c.is_geographic else "false"}</geographicflag></spatialrefsys>')


def write_qgis_project(out: Path, cfg: Config, data: "Data", gpkg: Path, terr_layer: str, label_field: str,
                       measures_layer: str | None, measure_types: list, evidence_layers: list, nmax: int = 0,
                       zones_layer: str | None = None, zone_measures=()) -> Path:
    """Файл проекта QGIS (.qgs, относительные пути) с результатами из evidence.gpkg.
    measure_types: [(тип мероприятий, [мероприятие, …])]; evidence_layers: [(имя слоя gpkg, геометрия, [значения layer])]"""
    import uuid
    L = unit_labels(cfg)
    rel = "./" + gpkg.name
    layers_xml, tree = [], []
    wgs = _q_srs(4326)
    b = data.oozt.to_crs(3857).total_bounds
    pad = max(b[2] - b[0], b[3] - b[1]) * 0.05

    def add_vec(name, layer, geom, renderer, subset=None, labels=None, visible=True):
        lid = f"{re.sub(r'[^0-9A-Za-z]+', '_', layer)}_{uuid.uuid4().hex[:10]}"
        src = f"{rel}|layername={layer}" + (f"|subset={subset}" if subset else "")
        gtype = {"poly": "Polygon", "line": "Line", "point": "Point"}[geom]
        lab_attr = ' labelsEnabled="1"' if labels else ""
        layers_xml.append(
            f'<maplayer type="vector" geometry="{gtype}" autoRefreshEnabled="0" refreshOnNotifyEnabled="0" '
            f'hasScaleBasedVisibilityFlag="0" readOnly="0" simplifyDrawingHints="1"{lab_attr}>'
            f'<id>{lid}</id><datasource>{_x(src)}</datasource><layername>{_x(name)}</layername>'
            f'<srs>{wgs}</srs><provider encoding="UTF-8">ogr</provider>{renderer}{labels or ""}'
            f'<blendMode>0</blendMode><featureBlendMode>0</featureBlendMode><layerOpacity>1</layerOpacity>'
            f'</maplayer>')
        return (f'<layer-tree-layer id="{lid}" name="{_x(name)}" checked="{"Qt::Checked" if visible else "Qt::Unchecked"}" '
                f'expanded="0" providerKey="ogr" source="{_x(src)}"><customproperties/></layer-tree-layer>')

    def group(name, children, checked=True, expanded=True):
        return (f'<layer-tree-group name="{_x(name)}" checked="{"Qt::Checked" if checked else "Qt::Unchecked"}" '
                f'expanded="{1 if expanded else 0}"><customproperties/>{"".join(children)}</layer-tree-group>')

    # 1) основания — цвета как на схемах
    ev_nodes = []
    for lyr, geom, values in sorted(evidence_layers, key=lambda t: {"point": 0, "line": 1, "poly": 2}[t[1]]):
        cats = []
        for i, v in enumerate(sorted(values)):
            color, kind = _slope_style(v)
            if geom == "point":
                sym = _q_symbol("marker", str(i), color, size=3.0 if "виды" in v else 2.2)
            elif geom == "line":
                sym = _q_symbol("line", str(i), color, width=0.9)
            else:
                alpha = 0.0 if "SDM" in v else (0.85 if kind == "slope" else 0.45)
                sym = _q_symbol("fill", str(i), color, alpha=alpha, outline=color, outline_width=0.3)
            cats.append((v, v, sym))
        title = {"point": "Основания — точки", "line": "Основания — линии", "poly": "Основания — участки"}[geom]
        ev_nodes.append(add_vec(title, lyr, geom, _q_categorized("layer", cats)))
    # 2) мероприятия: по типам, слой на каждое мероприятие (фильтр по слою measures)
    m_nodes = []
    if measures_layer:
        ci = 0
        for mtype, names in measure_types:
            kids = []
            for nm in names:
                col = _MEASURE_COLORS[ci % len(_MEASURE_COLORS)]
                ci += 1
                subset = '"Мероприятие" = ' + "'" + str(nm).replace("'", "''") + "'"
                if zones_layer and nm in zone_measures:
                    # участки работ — заливкой, полигоны с мероприятием — только контуром
                    kids.append(group(nm, [
                        add_vec("участки работ", zones_layer, "poly",
                                _q_single(_q_symbol("fill", "0", col, alpha=0.75, outline=col, outline_width=0.3)),
                                subset=subset, visible=True),
                        add_vec("полигоны с мероприятием (контур)", measures_layer, "poly",
                                _q_single(_q_symbol("fill", "0", col, alpha=0.0, outline=col, outline_width=0.6)),
                                subset=subset, visible=True),
                    ], checked=False, expanded=False))
                    continue
                kids.append(add_vec(nm, measures_layer, "poly",
                                    _q_single(_q_symbol("fill", "0", col, alpha=0.55, outline=col, outline_width=0.4)),
                                    subset=subset, visible=False))
            m_nodes.append(group(mtype or "Без типа", kids, checked=True, expanded=False))
    # 3) территория: число мероприятий на полигон
    reds = ["#FFFFFF", "#FEE8C8", "#FDBB84", "#FC8D59", "#E34A33", "#B30000", "#7F0000"]
    cats = [(str(k), (f"{k} мероприят." if k < 6 or nmax <= 6 else f"{k}"), _q_symbol(
        "fill", str(k), reds[min(k, len(reds) - 1)], alpha=0.0 if k == 0 else 0.75, outline="#555555",
        outline_width=0.2)) for k in range(0, nmax + 1)]
    terr = add_vec(f"{L['units_gen'].capitalize()}: число мероприятий", terr_layer, "poly",
                   _q_categorized("n_measures", cats), labels=_q_labels(label_field))
    # 4) исходные векторные слои (выключены)
    src_nodes = []
    for title, key, geom, color in (("Тропы", "trails", "line", "#707070"),
                                    ("Гидрография — линии", "hydro_lines", "line", "#1F78B4"),
                                    ("Гидрография — полигоны", "hydro_polys", "poly", "#A6CEE3")):
        for pth in getattr(cfg, key) or []:
            pth = Path(pth)
            try:
                rp = os.path.relpath(pth.resolve(), out.resolve())
                ds = "./" + Path(rp).as_posix() if not os.path.isabs(rp) else pth.as_posix()
            except ValueError:                  # другой диск — абсолютный путь
                ds = pth.as_posix()
            lid = f"src_{key}_{uuid.uuid4().hex[:10]}"
            sym = _q_symbol("line", "0", color, width=0.3) if geom == "line" else \
                _q_symbol("fill", "0", color, alpha=0.6, outline="#1F78B4", outline_width=0.2)
            layers_xml.append(
                f'<maplayer type="vector" geometry="{"Line" if geom == "line" else "Polygon"}">'
                f'<id>{lid}</id><datasource>{_x(ds)}</datasource><layername>{_x(title + " — " + pth.stem)}</layername>'
                f'<provider encoding="UTF-8">ogr</provider>{_q_single(sym)}</maplayer>')
            src_nodes.append(f'<layer-tree-layer id="{lid}" name="{_x(title + " — " + pth.stem)}" '
                             f'checked="Qt::Unchecked" expanded="0" providerKey="ogr" source="{_x(ds)}">'
                             f'<customproperties/></layer-tree-layer>')
    # 5) подложка OSM
    osm_id = f"osm_{uuid.uuid4().hex[:10]}"
    osm_src = "type=xyz&url=https://tile.openstreetmap.org/%7Bz%7D/%7Bx%7D/%7By%7D.png&zmax=19&zmin=0"
    layers_xml.append(
        f'<maplayer type="raster" autoRefreshEnabled="0"><id>{osm_id}</id><datasource>{_x(osm_src)}</datasource>'
        f'<layername>OpenStreetMap</layername><srs>{_q_srs(3857)}</srs><provider>wms</provider>'
        f'<pipe><rasterrenderer type="singlebandcolordata" opacity="0.8" alphaBand="-1" band="1">'
        f'<rasterTransparency/></rasterrenderer><brightnesscontrast brightness="0" contrast="0"/>'
        f'<huesaturation saturation="-70" grayscaleMode="0" colorizeOn="0" colorizeRed="255" colorizeGreen="128" '
        f'colorizeBlue="128" colorizeStrength="100"/><rasterresampler maxOversampling="2"/></pipe>'
        f'<blendMode>0</blendMode></maplayer>')
    osm_node = (f'<layer-tree-layer id="{osm_id}" name="OpenStreetMap" checked="Qt::Checked" expanded="0" '
                f'providerKey="wms" source="{_x(osm_src)}"><customproperties/></layer-tree-layer>')

    tree = []
    if ev_nodes:
        tree.append(group("Основания назначения", ev_nodes))
    if m_nodes:
        tree.append(group("Мероприятия (включайте нужные)", m_nodes, expanded=True))
    tree.append(terr)
    if src_nodes:
        tree.append(group("Исходные слои", src_nodes, checked=False, expanded=False))
    tree.append(osm_node)
    title = f"Природоохранные мероприятия: {', '.join(map(str, data.oozt['OOZT'].unique()[:3]))}"
    xml = ("<!DOCTYPE qgis PUBLIC 'http://mrcc.com/qgis.dtd' 'SYSTEM'>\n"
           f'<qgis projectname="{_x(title)}" version="3.28.0-Firenze">'
           f'<homePath path=""/><title>{_x(title)}</title>'
           f'<projectCrs>{_q_srs(3857)}</projectCrs>'
           f'<layer-tree-group><customproperties/>{"".join(tree)}</layer-tree-group>'
           f'<mapcanvas annotationsVisible="1" name="theMapCanvas"><units>meters</units>'
           f'<extent><xmin>{b[0] - pad}</xmin><ymin>{b[1] - pad}</ymin><xmax>{b[2] + pad}</xmax>'
           f'<ymax>{b[3] + pad}</ymax></extent><rotation>0</rotation>'
           f'<destinationsrs>{_q_srs(3857)}</destinationsrs><rendermaptile>0</rendermaptile></mapcanvas>'
           f'<projectlayers>{"".join(layers_xml)}</projectlayers>'
           f'<properties><Paths><Absolute type="bool">false</Absolute></Paths>'
           f'<Measurement><DistanceUnits type="QString">meters</DistanceUnits>'
           f'<AreaUnits type="QString">ha</AreaUnits></Measurement></properties>'
           f'</qgis>')
    path = out / QGIS_PROJECT_NAME
    path.write_text(xml, encoding="utf-8")
    return path


# =====================================================================================
#                                       MAIN
# =====================================================================================
def _setup_logging(out: Path, name="run.log"):
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s  %(message)s", datefmt="%H:%M:%S",
                        handlers=[logging.StreamHandler(sys.stdout),
                                  logging.FileHandler(out / name, encoding="utf-8", mode="w")])


def report_only(cfg: Config):
    """Перестроить отчёт по уже рассчитанным measures_long.csv и evidence.gpkg (без пересчёта правил)."""
    out = Path(cfg.out_dir)
    _setup_logging(out, "report.log")
    data = Data(cfg, oozt_only=True)
    try:
        data.measures = read_measures(cfg.measures_xlsx)
    except Exception as e:
        log.warning("Таблица мероприятий не прочитана (%s) — раздел «Как рассчитаны мероприятия» пропущен", e)
    # Для схем нужен полный контекст троп, а не только фрагменты,
    # послужившие основанием для назначения мероприятий.
    data.trails = data._load_trails()
    disp = display_columns(cfg)
    table = pd.read_csv(out / "measures_long.csv", encoding="utf-8-sig",
                        dtype={disp["Квартал"]: str, disp["Выдел"]: str})
    table = table.rename(columns={v: k for k, v in disp.items()})
    for c in ("Квартал", "Выдел"):
        table[c] = _maybe_int(table[c])
    gpkg = out / "evidence.gpkg"
    parts = []
    for ly, _ in pyogrio.list_layers(gpkg):
        if ly.startswith("evidence_"):
            g = gpd.read_file(gpkg, layer=ly)
            parts.append(g.to_crs(WORK_EPSG))
    ev = gpd.GeoDataFrame(pd.concat(parts, ignore_index=True), crs=WORK_EPSG) if parts else None
    rep = build_report(data, table, ev, out, cfg.report_kv, basemap=cfg.basemap,
                       formats=cfg.report_formats)
    log.info("Готово. Отчёт: %s", rep)


def run(cfg: Config):
    out = Path(cfg.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    _setup_logging(out)
    _MIN_SHARE[0] = float(getattr(cfg, "min_share_pct", 0) or 0)
    measures = read_measures(cfg.measures_xlsx)
    data = Data(cfg, pad_m=spec_pad_m(measures))
    data.measures = measures
    cfg.save(out / "config_used.json")
    res = Results()
    log.info("Мероприятия при выбранных данных:")
    for m in measures:
        stt = measure_status(m, cfg)
        log.info("  %-6s %-60s %s", m.get("code", ""), m["name"][:60], stt["text"])
        if m.get("on"):
            sp_ = stt["spec"]
            if sp_["rule"]:
                log.info("         правило: %s (%s)", RULES[sp_["rule"]]["title"],
                         "из столбцов таблицы" if sp_["origin"] == "таблица" else "по тексту «Метода выявления»")
            log.info("         как рассчитывается: %s", describe_rule(m))
            for k in stt.get("layers") or []:
                log.info("         слой видов: %s", species_layer_label(k))
    for m in measures:
        if not m["on"]:
            continue
        try:
            dispatch_rule(m, data, res)
        except Exception as e:
            log.exception("Ошибка в мероприятии «%s»: %s", m["name"], e)

    table = res.long_table()
    if table.empty:
        log.warning("Ни одно мероприятие не назначено")
        table = pd.DataFrame(columns=["ООЗТ", "Квартал", "Выдел", "Тип выдела", "Экосистема", "Тип мероприятий",
                                      "Мероприятие", "Тип экосистемы мероприятия", "Пункт Приказа",
                                      "Сезонность", "Основание", "Значение"])
    disp = display_columns(cfg)
    table.rename(columns=disp).to_csv(out / "measures_long.csv", index=False, encoding="utf-8-sig")
    try:
        with pd.ExcelWriter(out / "measures_long.xlsx", engine="openpyxl") as xw:
            table.rename(columns=disp).to_excel(xw, index=False, sheet_name="long")
            rt = rules_table(measures, table, cfg)
            if not rt.empty:
                rt.to_excel(xw, index=False, sheet_name="правила")
                from openpyxl.styles import Alignment as _Al2, Font as _Ft2
                wr = xw.book["правила"]
                wr.freeze_panes = "A2"
                for i, w in enumerate([8, 40, 10, 18, 50, 24, 60, 22], 1):
                    wr.column_dimensions[chr(64 + i)].width = w
                for cell in wr[1]:
                    cell.font = _Ft2(bold=True)
                for row in wr.iter_rows():
                    for cell in row:
                        cell.alignment = _Al2(wrap_text=True, vertical="top")
            if not table.empty:
                order_cols = ["ООЗТ", "Квартал", "Выдел", "Тип выдела", "Экосистема",
                              "Тип мероприятий", "Мероприятие", "Тип экосистемы мероприятия",
                              "Пункт Приказа"]
                by_order = table[order_cols].rename(columns={
                    "Мероприятие": "Мероприятия",
                    "Тип экосистемы мероприятия": "Тип экосистемы",
                    "Пункт Приказа": "Соответствие Пункту Приказа",
                })
                by_order.rename(columns=disp).to_excel(xw, index=False, sheet_name="мероприятия по приказу")

                # Значения не дедуплицируем: каждая строка во всех четырёх колонках
                # должна соответствовать одному и тому же мероприятию.
                join_lines = lambda s: "\n".join(
                    "" if pd.isna(x) else str(x) for x in s
                )
                wide = table.groupby(
                    ["ООЗТ", "Квартал", "Выдел", "Тип выдела", "Экосистема"],
                    as_index=False, dropna=False
                ).agg(**{
                    "Тип мероприятий": ("Тип мероприятий", join_lines),
                    "Мероприятия": ("Мероприятие", join_lines),
                    "Тип экосистемы": ("Тип экосистемы мероприятия", join_lines),
                    "Соответствие Пункту Приказа": ("Пункт Приказа", join_lines),
                    "Число": ("Мероприятие", "size"),
                })
                wide.rename(columns=disp).to_excel(xw, index=False, sheet_name="by_vydel")
                ex, how = execution_list(table, cfg, data)
                ex.to_excel(xw, index=False, sheet_name="перечень")
                ws = xw.book["перечень"]
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for i, c in enumerate(ex.columns, 1):
                    ws.column_dimensions[chr(64 + i)].width = 70 if c.startswith("Мероприятия") else (
                        30 if i == len(ex.columns) - 2 else 14)
                from openpyxl.styles import Alignment as _Al, Font as _Ft, PatternFill as _Pf
                for cell in ws[1]:
                    cell.font = _Ft(bold=True); cell.fill = _Pf("solid", fgColor="D9EAD3")
                    cell.alignment = _Al(wrap_text=True, vertical="center", horizontal="center")
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = _Al(wrap_text=True, vertical="top")
                from openpyxl.styles import Alignment, Font, PatternFill
                for ws in (xw.book["мероприятия по приказу"], xw.book["by_vydel"]):
                    # Двухуровневый заголовок: четыре колонки объединены
                    # общей шапкой «Мероприятия по приказу».
                    ws.insert_rows(1)
                    ws.merge_cells("A1:E1")
                    ws["A1"] = "Территориальная привязка"
                    ws.merge_cells("F1:I1")
                    ws["F1"] = "Мероприятия по приказу"
                    if ws.max_column >= 10:
                        ws["J1"] = "Число мероприятий"
                    ws.freeze_panes = "A3"
                    ws.auto_filter.ref = f"A2:{chr(64 + ws.max_column)}{ws.max_row}"
                    for cell in ws[1] + ws[2]:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="D9EAD3")
                        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                                   horizontal="center")
                    widths = [16, 11, 10, 16, 28, 34, 55, 22, 25, 10]
                    for i, width in enumerate(widths[:ws.max_column], 1):
                        ws.column_dimensions[chr(64 + i)].width = width
                    for row in ws.iter_rows(min_row=3):
                        for cell in row:
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
    except Exception as e:
        log.warning("Excel не записан (%s) — есть CSV", e)

    ev = res.evidence_gdf()
    try:
        gpkg = out / "evidence.gpkg"
        if gpkg.exists():
            gpkg.unlink()
        vm = data.oozt.copy()
        if not table.empty:
            agg = table.groupby(["ООЗТ", "Квартал", "Выдел"])["Мероприятие"].agg(lambda s: "; ".join(dict.fromkeys(s)))
            vm = vm.merge(agg.rename("measures").reset_index().rename(
                columns={"ООЗТ": "OOZT", "Квартал": "KV", "Выдел": "VYD"}), on=["OOZT", "KV", "VYD"], how="left")
        vm["n_measures"] = (vm["measures"].fillna("").map(lambda x: len([y for y in x.split("; ") if y]))
                            if "measures" in vm.columns else 0)
        vm = vm.drop(columns=["__src", "KV_SORT", "VYD_SORT"], errors="ignore")
        if cfg.unit_mode == "eco":
            # исходное поле ID уже есть в слое; служебный VYD не дублируем, KV — лист / группа отчёта
            vm = vm.drop(columns=["VYD"] if data.F_ID in vm.columns and data.F_ID != "VYD" else [])
            vm = vm.rename(columns={"KV": "report_sheet" if not cfg.report_group_field else "report_group"})
        terr_layer = "vydel_measures" if cfg.unit_mode != "eco" else "eco_measures"
        vm.to_crs(4326).to_file(gpkg, layer=terr_layer, driver="GPKG")
        # мероприятия: одна запись на пару «полигон — мероприятие», с основанием (для фильтров и стилей)
        measures_layer = None
        if not table.empty:
            geo = data.oozt[["OOZT", "KV", "VYD", "geometry"]].rename(
                columns={"OOZT": "ООЗТ", "KV": "Квартал", "VYD": "Выдел"})
            ml = geo.merge(table, on=["ООЗТ", "Квартал", "Выдел"], how="inner")
            ml = gpd.GeoDataFrame(ml.rename(columns=disp), geometry="geometry", crs=data.oozt.crs)
            ml.to_crs(4326).to_file(gpkg, layer="measures", driver="GPKG")
            measures_layer = "measures"
        # участки работ: часть полигона, где выполнено условие мероприятия (пятна-основания ∩ полигон).
        # Мероприятие назначается полигону целиком, но выполнять его нужно только здесь —
        # например, укрепление берегов — в полосе у воды, а не по всему крупному полигону.
        zones_layer, zone_measures = None, set()
        if not table.empty and ev is not None and not ev.empty:
            evp = ev[ev.geom_type.isin(["Polygon", "MultiPolygon"])
                     & ~ev["layer"].astype(str).str.startswith("водные объекты")]
            mlw = geo.merge(table, on=["ООЗТ", "Квартал", "Выдел"], how="inner")
            mlw = gpd.GeoDataFrame(mlw, geometry="geometry", crs=data.oozt.crs).to_crs(WORK_EPSG)
            zparts = []
            for nm, g in evp.groupby("measure"):
                sub = mlw[mlw["Мероприятие"] == nm]
                if sub.empty:
                    continue
                u = unary_union(g.geometry.buffer(0).values)
                z = sub.copy()
                z["geometry"] = sub.geometry.intersection(u)
                z = z[~z.is_empty & (z.area > 1)]
                if z.empty:
                    continue
                z["Площадь участка, га"] = (z.area / 1e4).round(3)
                z["Доля полигона, %"] = (100 * z.area / sub.loc[z.index].area).round(1)
                zparts.append(z)
                zone_measures.add(nm)
            if zparts:
                zz = gpd.GeoDataFrame(pd.concat(zparts), geometry="geometry", crs=WORK_EPSG)
                keep = ["ООЗТ", "Квартал", "Выдел", "Тип выдела", "Экосистема", "Тип мероприятий", "Мероприятие",
                        "Пункт Приказа", "Основание", "Площадь участка, га", "Доля полигона, %", "geometry"]
                zz = zz[[c for c in keep if c in zz.columns]].rename(columns=disp)
                zz.to_crs(4326).to_file(gpkg, layer="measure_zones", driver="GPKG")
                zones_layer = "measure_zones"
                for nm in sorted(zone_measures):
                    za = sum(z.area.sum() for z in zparts if z["Мероприятие"].iloc[0] == nm) / 1e4
                    pa = mlw[mlw["Мероприятие"] == nm].area.sum() / 1e4
                    log.info("  участки работ «%s»: %.1f га из %.1f га полигонов (%.0f%%)", nm, za, pa,
                             100 * za / pa if pa else 0)
        ev_layers = []
        if ev is not None and not ev.empty:
            for gt, grp in ev.groupby(ev.geom_type.str.replace("Multi", "")):
                grp.to_crs(4326).to_file(
                    gpkg, layer=f"evidence_{gt.lower()}", driver="GPKG")
                ev_layers.append((f"evidence_{gt.lower()}",
                                  {"Point": "point", "LineString": "line"}.get(gt, "poly"),
                                  sorted(grp["layer"].astype(str).unique())))
        log.info("Результаты записаны в GeoPackage: %s (слои %s%s%s)", gpkg, terr_layer,
                 ", measures" if measures_layer else "", "".join(", " + x[0] for x in ev_layers))
        try:
            mt = []
            if not table.empty:
                for t_, g_ in table.groupby("Тип мероприятий", sort=True):
                    mt.append((t_, list(dict.fromkeys(g_.sort_values("Пункт Приказа", key=_nat_sort_key(
                        "Пункт Приказа"))["Мероприятие"]))))
            qp = write_qgis_project(out, cfg, data, gpkg, terr_layer,
                                    data.F_ID if cfg.unit_mode == "eco" else "VYD", measures_layer, mt, ev_layers,
                                    nmax=int(vm["n_measures"].max()) if len(vm) else 0,
                                    zones_layer=zones_layer, zone_measures=zone_measures)
            log.info("Проект QGIS: %s", qp)
        except Exception as e:
            log.warning("Проект QGIS не создан: %s", e)
    except Exception as e:
        log.warning("GPKG не записан: %s", e)

    rep = build_report(data, table, ev, out, cfg.report_kv, basemap=cfg.basemap,
                       formats=cfg.report_formats)
    log.info("Готово. Таблица: %s; отчёт: %s", out / "measures_long.xlsx", rep)
    return table, rep


def main():
    ap = argparse.ArgumentParser(description="Природоохранные мероприятия для ООЗТ: кварталы/выделы "
                                             "или слой экосистем с уникальным ID")
    ap.add_argument("--config", help="JSON-конфигурация (без диалогов)")
    ap.add_argument("--kv", help="кварталы / листы / группы для отчёта, через запятую (переопределяет конфиг)")
    ap.add_argument("--gui", action="store_true",
                    help="открыть окно выбора данных, предзаполненное из --config")
    ap.add_argument("--no-basemap", action="store_true", help="не добавлять подложку OSM")
    ap.add_argument("--formats", help="форматы отчёта через запятую: docx,pdf,doc (report.md — всегда); "
                                      "пусто — только Markdown")
    ap.add_argument("--no-split", action="store_true",
                    help="не делить многочастные объекты слоя территории на отдельные контуры")
    ap.add_argument("--save-config-only", action="store_true", help="только собрать конфиг диалогами и сохранить")
    ap.add_argument("--update-table", metavar="XLSX",
                    help="записать в таблицу мероприятий столбцы правил, справку и описания расчёта и выйти")
    ap.add_argument("--report-only", action="store_true",
                    help="не пересчитывать: построить отчёт по measures_long.csv и evidence.gpkg из out_dir")
    args = ap.parse_args()
    if args.update_table:
        logging.basicConfig(level=logging.INFO, format="%(message)s")
        print(update_measures_xlsx(args.update_table, Config.load(args.config) if args.config else None))
        return

    if args.config and not args.gui:
        cfg = Config.load(args.config)
    else:
        cfg = gui_collect_config(Config.load(args.config) if args.config else None)
    if args.kv:
        cfg.report_kv = [int(x) if x.isdigit() else x for x in re.split(r"[,;\s]+", args.kv) if x]
    if args.no_basemap:
        cfg.basemap = False
    if args.no_split:
        cfg.split_multipart = False
    if args.formats is not None:
        cfg.report_formats = [x.strip().lower().lstrip(".") for x in re.split(r"[,;\s]+", args.formats) if x.strip()]
        bad = [x for x in cfg.report_formats if x not in REPORT_FORMATS]
        if bad:
            ap.error("неизвестные форматы: " + ", ".join(bad) + " (допустимы: docx, pdf, doc)")
    if args.save_config_only:
        Path(cfg.out_dir).mkdir(parents=True, exist_ok=True)
        cfg.save(Path(cfg.out_dir) / "config.json")
        print("Конфигурация сохранена:", Path(cfg.out_dir) / "config.json")
        return
    if args.report_only:
        report_only(cfg)
        return
    run(cfg)


if __name__ == "__main__":
    main()
